# -*- coding: utf-8 -*-
"""
==============================================================================
 KİTAPÇIK EŞLEŞTİRİCİ (Uygulama Sürümü) — Limit Yayınları
==============================================================================
Pencereli sürüm: A/B kitapçık PDF'lerini dosya seçiciyle seçin, EŞLEŞTİR'e
basın; kazanım tablosu Excel'i otomatik oluşsun. Terminal bilgisi gerekmez.

Çalıştırma: "Kitapcik Eslestirici.app" uygulamasına çift tıklayın
            (veya terminalden: python3 kitapcik_eslestirici_app.py)
==============================================================================
"""

import difflib
import importlib
import importlib.util
import os
import queue
import re
import subprocess
import sys
import threading
import traceback
import unicodedata
from pathlib import Path

SURUM = "2.31"
GEREKLI = ["pymupdf", "numpy", "scipy", "openpyxl", "tkinterdnd2",
           "python-docx"]
LOG_DOSYASI = Path.home() / "Library" / "Logs" / "KitapcikEslestirici.log"


def logla(mesaj):
    try:
        LOG_DOSYASI.parent.mkdir(parents=True, exist_ok=True)
        with open(LOG_DOSYASI, "a", encoding="utf-8") as f:
            f.write(mesaj + "\n")
    except OSError:
        pass


def hata_penceresi(mesaj):
    try:
        subprocess.run(["osascript", "-e",
                        'display dialog "' + mesaj.replace('"', "'") +
                        '" buttons {"Tamam"} default button 1 '
                        'with title "Kitapçık Eşleştirici"'], timeout=60)
    except Exception:
        pass

# ----------------------------------------------------------------------------
# SINAV YAPILARI
# ----------------------------------------------------------------------------
YAPILAR = {
    "TYT": [
        {"test": "Türkçe", "dersler": [["Türkçe", 40]]},
        {"test": "Sosyal Bilimler", "dersler": [["Tarih", 5], ["Coğrafya", 5],
                                                ["Felsefe", 5], ["Din Kültürü", 5],
                                                ["Seçmeli Felsefe", 5]]},
        {"test": "Temel Matematik", "dersler": [["Temel Matematik", 40]]},
        {"test": "Fen Bilimleri", "dersler": [["Fizik", 7], ["Kimya", 7],
                                              ["Biyoloji", 6]]},
    ],
    "AYT": [
        {"test": "Türk Dili ve Edebiyatı-Sosyal Bilimler-1",
         "dersler": [["Edebiyat", 24], ["Tarih", 10], ["Coğrafya", 6]]},
        {"test": "Sosyal Bilimler-2",
         "dersler": [["Tarih-2", 11], ["Coğrafya-2", 11], ["Felsefe Grubu", 12],
                     ["Din Kültürü", 6], ["Seçmeli Felsefe", 6]]},
        {"test": "Matematik", "dersler": [["Matematik", 40]]},
        {"test": "Fen Bilimleri", "dersler": [["Fizik", 14], ["Kimya", 13],
                                              ["Biyoloji", 13]]},
    ],
}

SORU_BASI_RE = re.compile(r"^(\d{1,3})\s*[\.\)]\s*(.*)$")
SECENEK_A_RE = re.compile(r"(?m)^\s*A\s*[\)\.]")
CEVAP_SATIRI_RE = re.compile(r"^\s*([A-E])\s+([A-Z0-9ÇĞİÖŞÜ]{5,})\s*$")
# Test başı yönerge satırları ("1. Bu testte 40 soru vardır." vb.) soru DEĞİLDİR
YONERGE_RE = re.compile(r"^(Bu\s+test|Cevaplar)", re.IGNORECASE)


def renkli_mi(renk):
    """Span rengi siyah/gri değilse True (magenta cevap işaretleri için)."""
    r, g, b = (renk >> 16) & 255, (renk >> 8) & 255, renk & 255
    return max(r, g, b) - min(r, g, b) > 60


def cevap_kodu_bul(satir, sadece_renk=False):
    """Satırdaki renkli (magenta) cevap harfi ile soru kodunu döndürür.
    Parçalara bölünmüş kodları birleştirir; sadece_renk=True ise yalnızca
    renkli span'lara bakar (soru numarası satırı için güvenli mod)."""
    cevap, kod_parcalari = None, []
    for metin, renk in satir.get("spanlar", []):
        t = metin.strip()
        if not t or not renkli_mi(renk):
            continue
        if re.fullmatch(r"[A-E]", t):
            cevap = t
        else:
            temiz = t.replace(" ", "")
            if re.fullmatch(r"[A-Za-z0-9ÇĞİÖŞÜçğıöşü]{2,}", temiz):
                kod_parcalari.append(temiz)
    kod = "".join(kod_parcalari) if kod_parcalari else None
    if kod and (len(kod) < 5 or not re.search(r"\d", kod)):
        kod = None  # rakamsız/kısa renkli metin kod değildir (vurgu vb. olabilir)
    if cevap is None and not sadece_renk:  # renk bilgisi yoksa satır deseni
        m = CEVAP_SATIRI_RE.match(satir["metin"])
        if m:
            cevap, kod = m.group(1), m.group(2)
    return cevap, kod


def cevap_konumdan_bul(satirlar, sat_i, baslar):
    """Cevap harfinin konumdan bulunması: aynı sayfada, soru numarasının sol
    hizasında (±8 pt), numaranın altında ve aynı sütundaki bir sonraki soru
    numarasından önce duran ilk renkli A–E harfi. Satır sırasından bağımsızdır:
    tam genişlik sayfada numara satırı sağ yarıya düşse de, sorunun şeklinde
    renkli harf etiketleri olsa da doğru harfi verir."""
    bas = satirlar[sat_i]
    pno, x0, y0 = bas["sayfa"], bas["x0"], bas["y0"]
    alt = min([satirlar[b]["y0"] for b in baslar
               if satirlar[b]["sayfa"] == pno and satirlar[b]["y0"] > y0 + 2
               and abs(satirlar[b]["x0"] - x0) < 30] + [1e9])
    adaylar = []
    for r in satirlar:
        if r["sayfa"] != pno or not y0 < r["y0"] < alt or abs(r["x0"] - x0) > 8:
            continue
        c, _k = cevap_kodu_bul(r, sadece_renk=True)
        if c:
            adaylar.append((r["y0"], c))
    return min(adaylar)[1] if adaylar else None


def bagimliliklari_yukle():
    """Ağır kütüphaneleri yükler; eksikse ImportError fırlatır."""
    global fitz, np, linear_sum_assignment
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # eski sürüm adı
    import numpy as np
    from scipy.optimize import linear_sum_assignment


def eksikleri_bul():
    importlib.invalidate_caches()
    return [p for p in GEREKLI if importlib.util.find_spec({"pymupdf": "fitz", "python-docx": "docx"}.get(p, p)) is None
        and not (p == "pymupdf" and importlib.util.find_spec("fitz"))]


_XL_YASAK_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xl(v):
    """Excel hücresine yazılamayan kontrol karakterlerini atar (PDF metninden
    gelen açıklamalarda bulunabilir; openpyxl bunlarla dosyayı yazamaz)."""
    return _XL_YASAK_RE.sub("", v) if isinstance(v, str) else v


def normalize(s):
    s = unicodedata.normalize("NFC", s).casefold()
    return re.sub(r"[^a-zçğıöşü0-9]+", "", s)


# Deneme ders adı <-> ana sözlük ders adı eşlemesi (özel durumlar dahil)
_DERS_ESLEME_EK = {
    "felsefegrubu": ["felsefe", "psikoloji", "sosyoloji", "mantık"],
    "edebiyat": ["edebiyat", "türk dili ve edebiyatı"],
}


def _ders_uyusur(yapi_ders, sozluk_ders):
    a = normalize(re.sub(r"-\d+$", "", str(yapi_ders)))
    b = normalize(str(sozluk_ders))
    if not a or not b:
        return False
    if a in b or b in a:
        return True
    for hedef in _DERS_ESLEME_EK.get(a, []):
        if normalize(hedef) in b or b in normalize(hedef):
            return True
    return False


# ----------------------------------------------------------------------------
# MOTOR (PDF okuma → soru ayıklama → eşleştirme → Excel)
# ----------------------------------------------------------------------------
def satirlari_al(pdf_yolu):
    doc = fitz.open(pdf_yolu)
    satirlar = []
    kolonlar = {0: [1e9, 0], 1: [1e9, 0]}
    for pno, page in enumerate(doc):
        d = page.get_text("dict")
        orta = page.rect.width / 2
        sol, sag = [], []
        for blk in d.get("blocks", []):
            if blk.get("type", 0) != 0:
                continue
            for ln in blk.get("lines", []):
                spanlar = [(sp.get("text", ""), sp.get("color", 0))
                           for sp in ln.get("spans", [])]
                kutular = [tuple(sp.get("bbox", ln["bbox"]))
                           for sp in ln.get("spans", [])]
                metin = "".join(t for t, _ in spanlar).strip()
                if not metin:
                    continue
                x0, y0, x1, y1 = ln["bbox"]
                sutun = 0 if (x0 + x1) / 2 < orta else 1
                kayit = {"metin": metin, "x0": x0, "y0": y0, "x1": x1, "y1": y1,
                         "sayfa": pno, "sutun": sutun, "spanlar": spanlar,
                         "kutular": kutular}
                (sol if sutun == 0 else sag).append(kayit)
                kolonlar[sutun][0] = min(kolonlar[sutun][0], x0)
                kolonlar[sutun][1] = max(kolonlar[sutun][1], x1)
        for sutun in (sol, sag):
            sutun.sort(key=lambda r: (round(r["y0"], 1), r["x0"]))
        satirlar.extend(sol)
        satirlar.extend(sag)
    doc.close()
    if not satirlar:
        raise RuntimeError(f"'{Path(pdf_yolu).name}' içinden metin okunamadı. "
                           f"Taranmış/görsel PDF ise önce OCR gerekir.")
    return satirlar, kolonlar



def yapi_cikar(a_pdf):
    """Yapıyı A kitapçığının kendisinden türetir: numaraların 1'e döndüğü
    yerler test bloklarını, blok üstündeki başlık satırı test adını verir."""
    satirlar, _kolonlar = satirlari_al(a_pdf)
    adaylar, _yedek = _adaylari_bul(satirlar)
    if not adaylar:
        raise RuntimeError("A kitapçığında soru numarası bulunamadı; yapı "
                           "çıkarılamadı.")
    bloklar = []
    bas_idx, en_buyuk = None, 0
    for p, (i, no) in enumerate(adaylar):
        yeni = (no == 1 and (bas_idx is None or en_buyuk >= 3)
                and any(a[1] == 2 for a in adaylar[p + 1:p + 4]))
        if yeni:
            if bas_idx is not None:
                bloklar.append((bas_idx, en_buyuk))
            bas_idx, en_buyuk = i, 1
        elif bas_idx is not None:
            en_buyuk = max(en_buyuk, no)
    if bas_idx is not None:
        bloklar.append((bas_idx, en_buyuk))
    if not bloklar:
        raise RuntimeError("A kitapçığında test bloğu bulunamadı.")

    def blok_bilgisi(bas_satir, sira, adet, onceki_bas):
        test_adi, dersler, yonerge_adet = None, None, None
        for k in range(bas_satir - 1, max(-1, bas_satir - 10), -1):
            m = re.sub(r"[\x00-\x1f]+", " ", satirlar[k]["metin"]).strip()
            if not m:
                continue
            if test_adi is None:
                eb = re.search(r"BÖLÜM\s*[-–—]\s*(.{2,30})$", m)
                if eb:
                    test_adi = eb.group(1).strip()
            if test_adi is None:
                e = re.search(r"cevap kâğıdının\s+(.+?)\s+Testi",
                              m, re.IGNORECASE)
                if e:
                    test_adi = e.group(1).strip()
                else:
                    e2 = re.match(r"^(.+?)\s+TESTİNE GEÇİNİZ", m)
                    if e2:
                        test_adi = e2.group(1).strip()
            if yonerge_adet is None and "Bu testte" in m:
                ya = re.search(r"Bu testte[^0-9]*(\d+)\s*soru", m)
                if ya:
                    yonerge_adet = int(ya.group(1))
            if dersler is None and "Bu testte" in m:
                bul = re.findall(
                    r"([A-Za-zÇĞİÖŞÜçğıöşü][\w ÇĞİÖŞÜçğıöşü.\-]*?)"
                    r"\s*\((\d+)\s*[-–]\s*(\d+)\)", m)
                cift = [(re.sub(r"^(ve|ile)\s+", "",
                                ad.strip(" ,.")), int(b2) - int(b1) + 1)
                        for ad, b1, b2 in bul]
                if cift and sum(a for _n, a in cift) == adet:
                    dersler = [[n, a] for n, a in cift]
        if test_adi is None:      # daha geriye, bir önceki bloğa kadar bak
            for k in range(bas_satir - 1, max(onceki_bas - 1, -1), -1):
                mm = re.sub(r"[\x00-\x1f]+", " ", satirlar[k]["metin"]).strip()
                e3 = re.match(r"^(.+?)\s+TEST[İI]NE GEÇ[İI]N[İI]Z", mm)
                if e3:
                    test_adi = e3.group(1).strip()
                    break
                eb2 = re.search(r"BÖLÜM\s*[-–—]\s*(.{2,30})$", mm)
                if eb2:
                    test_adi = eb2.group(1).strip()
                    break
                if mm == "TEST ADI":
                    for kk3 in range(k + 1, min(k + 4, len(satirlar))):
                        ad3 = satirlar[kk3]["metin"].strip()
                        if ad3 and ad3 not in ("SORU SAYISI", "SINIFI",
                                               "ADI", "SOYADI"):
                            test_adi = ad3
                            break
                    if test_adi:
                        break
        if test_adi is None:
            for k in range(bas_satir - 1, max(-1, bas_satir - 7), -1):
                m = satirlar[k]["metin"].strip()
                if not m or SORU_BASI_RE.match(m) or "Bu test" in m:
                    continue
                if any(len(t) >= 8 and t.isupper()
                       and any(c.isdigit() for c in t) for t in m.split()):
                    continue
                if not any(len(t) >= 4 and t.isalpha() for t in m.split()):
                    continue
                buyukce = sum(1 for ch in m if ch.isalpha() and ch == ch.upper())
                harf = sum(1 for ch in m if ch.isalpha())
                if harf >= 4 and ("TEST" in m.upper()
                                  or (harf and buyukce / harf > 0.8)):
                    test_adi = m[:40]
                    break
        if test_adi is None:
            test_adi = f"Bölüm {sira}"
        if yonerge_adet and yonerge_adet > adet:
            adet = yonerge_adet          # yönergedeki sayı esas alınır
        if not dersler:
            dersler = [[test_adi, int(adet)]]
        return {"test": test_adi, "dersler": dersler}

    yapi = []
    for s, (bas_satir, adet) in enumerate(bloklar, 1):
        onceki = bloklar[s - 2][0] if s >= 2 else 0
        yapi.append(blok_bilgisi(bas_satir, s, int(adet), onceki))

    # Test (ders) adı: testin ilk sorusunun sayfasındaki üst başlık esas alınır
    # ("SÖZEL BÖLÜM - TÜRKÇE" → TÜRKÇE, "TÜRKÇE TESTİ" → TÜRKÇE). Başlık sağ
    # yarıya düşüp satır sırasında geride kaldığında ya da önceki testin
    # "… TESTİ BİTTİ" yazısı okunduğunda ad kayıyordu (ör. "Bölüm 1").
    doc = fitz.open(a_pdf)
    yukseklikler = [p.rect.height for p in doc]
    doc.close()
    kullanilan = set()
    for t, (bas_satir, _adet) in zip(yapi, bloklar):
        ad = _sayfa_basligi(satirlar, satirlar[bas_satir]["sayfa"],
                            yukseklikler[satirlar[bas_satir]["sayfa"]], kullanilan)
        if ad:
            kullanilan.add(ad)
        if ad and ad != t["test"]:
            eski, t["test"] = t["test"], ad
            for d in t["dersler"]:
                if d[0] == eski:
                    d[0] = ad
    return yapi


def _sayfa_basligi(satirlar, pno, sayfa_yuk, kullanilan=()):
    """Sayfanın üst kısmındaki test başlığı: "SÖZEL BÖLÜM - TÜRKÇE" → "TÜRKÇE",
    "SOSYAL BİLİMLER TESTİ" → "SOSYAL BİLİMLER". Yoksa None. Sayfada birden
    çok başlık varsa (ör. kalıp sayfadan kalıp üstü örtülmüş eski başlık)
    önceki testlerde kullanılmamış olan seçilir."""
    adaylar = []
    for r in satirlar:
        if r["sayfa"] != pno or r["y1"] > 0.2 * sayfa_yuk:
            continue
        m = r["metin"].strip()
        e = (re.search(r"BÖLÜM\s*[-–—]\s*(.{2,60})$", m)
             or re.match(r"^(.{2,60}?)\s+TEST[İI]$", m))
        if e:
            adaylar.append(e.group(1).strip())
    yeni = [a for a in adaylar if a not in kullanilan]
    return (yeni or adaylar or [None])[0]


def beklenen_dizi(yapi):
    dizi = []
    for ti, test in enumerate(yapi):
        n = 0
        for ders, adet in test["dersler"]:
            for _ in range(int(adet)):
                n += 1
                dizi.append((ti, ders, n))
    return dizi


def _adaylari_bul(satirlar):
    ham = []
    for i, r in enumerate(satirlar):
        m = SORU_BASI_RE.match(r["metin"])
        if not m or not (1 <= int(m.group(1)) <= 60):
            continue
        kalan = m.group(2)
        if YONERGE_RE.match(kalan):
            continue                      # "1. Bu testte ..." aynı satırda
        if re.match(r"^S[Iİı]N[Iİı]F\b", kalan, re.IGNORECASE) and \
                re.search(r"DENEME|OKUL|FINAL|TEST|YAYIN",
                          r["metin"], re.IGNORECASE):
            continue                      # "6. SINIF FINAL ..." koşan başlığı
        if not kalan.strip() and i + 1 < len(satirlar) \
                and YONERGE_RE.match(satirlar[i + 1]["metin"]):
            continue                      # "1." ayrı satır, yönerge metni altında
        ham.append((i, int(m.group(1)), r["sutun"], r["x0"]))

    # Kenar hizası filtresi: gerçek soru numaraları sütunun sol kenarına dizilir;
    # soru İÇİNDEKİ numaralı maddeler (öncüller) içeriden başlar ve elenir.
    from collections import Counter as _Sayac
    kenar = {}
    for _i, _no, _sutun, _x0 in ham:
        kenar.setdefault(_sutun, _Sayac())[round(_x0)] += 1
    kenar = {s: c.most_common(1)[0][0] for s, c in kenar.items()}
    adaylar, yedekler = [], []
    kenarlar = list(kenar.values())
    for i, no, sutun, x0 in ham:
        uyan = abs(x0 - kenar.get(sutun, x0)) <= 6
        if not uyan:
            genis = (satirlar[i].get("x1", x0) - x0) > 250  # tam genişlik satırı
            uyan = genis and min((abs(x0 - m) for m in kenarlar),
                                 default=99) <= 6
        if uyan:
            adaylar.append((i, no))
        else:
            yedekler.append((i, no))
    return adaylar, yedekler


def sorulari_ayikla(pdf_yolu, yapi, etiket, log, dokum=None):
    satirlar, kolonlar = satirlari_al(pdf_yolu)
    adaylar, yedekler = _adaylari_bul(satirlar)

    beklenen = beklenen_dizi(yapi)
    # Deterministik hizalama: beklenen numaralar sırayla yürünür; blok
    # karışması imkânsızdır. Aday uymazsa ya gürültüdür (yakında doğrusu
    # geliyorsa atlanır) ya da o soru gerçekten eksiktir.
    eslesen = {}
    p = i = 0
    while i < len(adaylar) and p < len(beklenen):
        if adaylar[i][1] == beklenen[p][2]:
            eslesen[p] = adaylar[i][0]
            i += 1
            p += 1
        elif any(adaylar[j][1] == beklenen[p][2]
                 for j in range(i + 1, min(i + 4, len(adaylar)))):
            i += 1      # gürültü aday: beklenen numara hemen arkasından geliyor
        else:
            p += 1      # bu beklenen soru adaylar arasında yok (eksik)

    # --- Kurtarma 1: sırası kaymış (çapraz kolona dizilmiş) soruları kazan ---
    kullanilan = set(eslesen.values())
    for bek_i in range(len(beklenen)):
        if bek_i in eslesen:
            continue
        ti_b, _d, no_b = beklenen[bek_i]
        komsular = [eslesen[b] for b in eslesen
                    if beklenen[b][0] == ti_b and abs(b - bek_i) <= 4]
        if not komsular:
            continue
        alt, ust = min(komsular), max(komsular)
        bulundu = False
        for sat_i, aday_no in adaylar:
            if (aday_no == no_b and sat_i not in kullanilan
                    and alt <= sat_i <= ust):
                eslesen[bek_i] = sat_i
                kullanilan.add(sat_i)
                bulundu = True
                break
        if not bulundu:          # blok-geneli arama: önce asıl, sonra yedek
            test_satirlari = [eslesen[b] for b in eslesen
                              if beklenen[b][0] == ti_b]
            sonraki_bas = min([eslesen[b] for b in eslesen
                               if beklenen[b][0] > ti_b] or [len(satirlar)])
            g_alt = min(test_satirlari) if test_satirlari else 0
            for kaynak in (adaylar, yedekler):
                for sat_i, aday_no in kaynak:
                    if (aday_no == no_b and sat_i not in kullanilan
                            and g_alt <= sat_i < sonraki_bas):
                        eslesen[bek_i] = sat_i
                        kullanilan.add(sat_i)
                        bulundu = True
                        break
                if bulundu:
                    break

    # --- Kurtarma 2: şıksız (sahte) çapaları aynı numaralı gerçek satıra taşı ---
    def _secenekli(bas, son):
        return any(re.match(r"^\s*A\s*[\)\.]", satirlar[k]["metin"])
                   for k in range(bas, min(son, len(satirlar))))

    for _tur in range(3):
        degisti = False
        baslar = sorted(eslesen.values())
        for bek_i, sat_i in list(eslesen.items()):
            son = next((b for b in baslar if b > sat_i), len(satirlar))
            if _secenekli(sat_i, son):
                continue
            ti_b, _d, no_b = beklenen[bek_i]
            ayni_test = [eslesen[b] for b in eslesen if beklenen[b][0] == ti_b]
            sinir = max(ayni_test) if ayni_test else len(satirlar)
            for aday_sat, aday_no in adaylar:
                if (aday_no != no_b or aday_sat in kullanilan
                        or aday_sat <= sat_i or aday_sat > sinir):
                    continue
                son2 = next((b for b in baslar if b > aday_sat), len(satirlar))
                if _secenekli(aday_sat, son2):
                    kullanilan.discard(sat_i)
                    kullanilan.add(aday_sat)
                    eslesen[bek_i] = aday_sat
                    degisti = True
                    break
        if not degisti:
            break

    sorular, uyarilar = {}, []
    sirali = sorted(eslesen.items())
    tum_baslar = sorted(eslesen.values())
    for bek_i, sat_i in sirali:
        ti, ders, no = beklenen[bek_i]
        son = next((b for b in tum_baslar if b > sat_i), len(satirlar))
        bolum = satirlar[sat_i:son]
        ilk_secenek = next((i for i, r in enumerate(bolum)
                            if re.match(r"^\s*A\s*[\)\.]", r["metin"])), len(bolum))
        cevap = kod = None
        atla = set()
        c0, k0 = cevap_kodu_bul(bolum[0], sadece_renk=True)
        if c0 or k0:
            cevap, kod = c0, k0
        for i in range(1, ilk_secenek):  # cevap işareti seçeneklerden önce durur
            c, k = cevap_kodu_bul(bolum[i])
            if c or k:
                atla.add(i)
                cevap = cevap or c
                kod = kod or k
        # Cevap işareti soru numarasının hizasında (sol boşlukta) duruyorsa o
        # alınır. Sıra kuralı yalnızca bu yoksa geçerlidir: sorunun içindeki
        # ilk renkli harfi aldığı için şekildeki renkli etiketi (ör. kırmızı
        # "A", "B") ya da tam genişlik sayfada sonraki sayfadaki başka bir
        # sorunun cevabını alabiliyordu.
        konum = cevap_konumdan_bul(satirlar, sat_i, tum_baslar)
        if konum:
            cevap = konum
        parcalar = []
        for i, r in enumerate(bolum):
            if i in atla:
                continue
            if i == 0:
                ham = ("".join(t for t, rk in r["spanlar"]
                               if not renkli_mi(rk)).strip()
                       if (c0 or k0) else r["metin"])
                m0 = SORU_BASI_RE.match(ham)
                parcalar.append(m0.group(2) if m0 else ham)
            else:
                parcalar.append(r["metin"])
        metin = "\n".join(parcalar)
        m = SECENEK_A_RE.search(metin)
        govde = metin[:m.start()] if m else metin
        bas = bolum[0]
        ayni = [r for r in bolum
                if r["sayfa"] == bas["sayfa"] and r["sutun"] == bas["sutun"]]
        e_idx = next((i for i, r in enumerate(ayni)
                      if re.match(r"^\s*E\s*[\)\.]", r["metin"])), None)
        if e_idx is None:
            kesit = ayni
        else:
            kesit = ayni[:e_idx + 1]
            adimlar = [ayni[j + 1]["y0"] - ayni[j]["y0"]
                       for j in range(max(0, e_idx - 4), e_idx)]
            adimlar = [d for d in adimlar if d > 0]
            tipik = sorted(adimlar)[len(adimlar) // 2] if adimlar else 14
            son_y0 = ayni[e_idx]["y0"]
            j = e_idx + 1
            while j < len(ayni) and (ayni[j]["y0"] - son_y0) < 1.8 * tipik:
                kesit.append(ayni[j])       # E şıkkının devam satırları
                son_y0 = ayni[j]["y0"]
                j += 1
        ks = kolonlar.get(bas["sutun"])
        if ks and ks[0] < ks[1]:
            gx0, gx1 = ks[0] - 2, ks[1] + 2
        else:
            gx0 = min(r["x0"] for r in ayni) - 2
            gx1 = max(r["x1"] for r in ayni) + 2
        bolge = (bas["sayfa"], gx0, bas["y0"] - 3, gx1,
                 max(r["y1"] for r in kesit) + 3)
        sorular[(ti, no)] = {"ders": ders, "no": no, "cevap": cevap, "kod": kod,
                             "bolge": bolge, "tam": metin,
                             "govde": govde.strip(), "norm": normalize(metin)}
    for bek_i, (ti, ders, no) in enumerate(beklenen):
        if bek_i not in eslesen:
            uyarilar.append(f"{etiket} kitapçığı: {yapi[ti]['test']} / {ders} "
                            f"{no}. soru PDF'te bulunamadı.")

    # Erken uyarı teli: magenta kodları numaralardan BAĞIMSIZ say ve karşılaştır.
    # Henüz görülmemiş dizgi desenleri bile bu sayımda kendini ele verir.
    kod_sayim = {}
    for r in satirlar:
        _c, _k = cevap_kodu_bul(r, sadece_renk=True)
        if _k:
            mt = re.search(r"T(\d{1,2})S\d{1,3}$", _k)
            if mt:
                t_i = int(mt.group(1)) - 1
                kod_sayim[t_i] = kod_sayim.get(t_i, 0) + 1
    bulunan_test = {}
    for (ti2, _no2) in sorular:
        bulunan_test[ti2] = bulunan_test.get(ti2, 0) + 1
    for ti2 in sorted(kod_sayim):
        if 0 <= ti2 < len(yapi) and kod_sayim[ti2] > bulunan_test.get(ti2, 0):
            uyarilar.append(
                f"{etiket} kitapçığı: {yapi[ti2]['test']} bölümünde "
                f"{kod_sayim[ti2]} soru kodu sayıldı ama "
                f"{bulunan_test.get(ti2, 0)} soru bulundu — yeni bir dizgi "
                f"deseni olabilir; log dosyasını iletin.")
    if uyarilar:
        # Tanı dökümü: sorun bölgesinde parser'ın tam olarak ne gördüğü
        logla(f"--- TANI ({etiket} kitapçığı, {Path(pdf_yolu).name}) ---")
        logla("Aday numara dizisi: " + " ".join(str(n) for _, n in adaylar))
        eksikler = [bp for bp in range(len(beklenen)) if bp not in eslesen]
        onceki = (max((bp for bp in eslesen if bp < eksikler[0]), default=None)
                  if eksikler else None)
        bas = eslesen[onceki] if onceki is not None else 0
        logla("İlk eksik bölge çevresindeki ham satırlar:")
        for k in range(bas, min(bas + 30, len(satirlar))):
            logla(f"  [{k}] {satirlar[k]['metin'][:70]}")
    if dokum is not None:   # yalnızca A–B KONTROL kullanır; EŞLEŞTİR'i etkilemez
        dokum.update(satirlar=satirlar, adaylar=adaylar, yedekler=yedekler,
                     eslesen=eslesen, beklenen=beklenen)
    log(f"  {etiket} kitapçığı: {len(sorular)} soru bulundu.")
    return sorular, uyarilar


def eslestir(a_sorular, b_sorular, yapi, esik=0.60):
    """Önce magenta soru kodlarıyla birebir eşler; kod yoksa/konumsal ise
    şıklar dahil tam metin benzerliğiyle Macar algoritmasına düşer."""
    from collections import Counter
    sonuc, uyarilar = {}, []

    def benzerlik(sa, sb):
        if sa["norm"] and sb["norm"]:
            return difflib.SequenceMatcher(None, sa["norm"], sb["norm"]).ratio()
        return 0.0

    for ti, test in enumerate(yapi):
        n = 0
        for ders, adet in test["dersler"]:
            aralik = list(range(n + 1, n + int(adet) + 1))
            n += int(adet)
            ga = [(no, a_sorular[(ti, no)]) for no in aralik if (ti, no) in a_sorular]
            gb = [(no, b_sorular[(ti, no)]) for no in aralik if (ti, no) in b_sorular]
            if not ga or not gb:
                continue

            # --- 1) Soru kodu tabanlı eşleme (kod her iki tarafta da tekilse) ---
            ka = Counter(s["kod"] for _, s in ga if s.get("kod"))
            kb = Counter(s["kod"] for _, s in gb if s.get("kod"))
            a_by = {s["kod"]: (no, s) for no, s in ga
                    if s.get("kod") and ka[s["kod"]] == 1}
            b_by = {s["kod"]: (no, s) for no, s in gb
                    if s.get("kod") and kb[s["kod"]] == 1}
            ortak = set(a_by) & set(b_by)
            kod_es = {}
            if ortak:
                kullan = True
                ayni_yerde = all(a_by[k][0] == b_by[k][0] for k in ortak)
                if ayni_yerde and len(ortak) > 1:
                    # Kodlar konumsal (kitapçığa göre yeniden numaralı) olabilir;
                    # aynı konumdaki metinler gerçekten benziyorsa güvenilir say.
                    ort = sum(benzerlik(a_by[k][1], b_by[k][1])
                              for k in ortak) / len(ortak)
                    kullan = ort >= 0.75
                if kullan:
                    for k in ortak:
                        a_no, sa = a_by[k]
                        b_no, sb = b_by[k]
                        kod_es[a_no] = b_no
                        sonuc[(ti, a_no)] = {"b_no": b_no,
                                             "benzerlik": benzerlik(sa, sb),
                                             "yontem": "KOD"}

            # --- 2) Kalanlar için tam metin benzerliğiyle en iyi birebir atama ---
            kalan_a = [(no, s) for no, s in ga if no not in kod_es]
            alinan_b = set(kod_es.values())
            kalan_b = [(no, s) for no, s in gb if no not in alinan_b]
            if kalan_a and kalan_b:
                M = np.zeros((len(kalan_a), len(kalan_b)))
                for i, (_, sa) in enumerate(kalan_a):
                    for j, (_, sb) in enumerate(kalan_b):
                        s_ = benzerlik(sa, sb)
                        ca_, cb_ = sa.get("cevap"), sb.get("cevap")
                        if ca_ and cb_:
                            if ca_ == cb_:      # magenta cevaplar uyuşuyor
                                s_ = min(1.0, s_ + 0.25)
                            else:               # farklı cevap: güçlü ceza
                                s_ *= 0.35
                        M[i, j] = s_
                sat, sut = linear_sum_assignment(1 - M)
                for i, j in zip(sat, sut):
                    a_no, b_no, oran = kalan_a[i][0], kalan_b[j][0], M[i, j]
                    sonuc[(ti, a_no)] = {"b_no": b_no, "benzerlik": oran,
                                         "yontem": "METİN"}
                    if oran < esik:
                        uyarilar.append(
                            f"DÜŞÜK BENZERLİK: {test['test']} / {ders} A-{a_no} → "
                            f"B-{b_no} (%{oran*100:.0f}).")
    return sonuc, uyarilar


# ----------------------------------------------------------------------------
# GÖRSEL DOĞRULAMA: her eşleşen çift sayfa görüntüsü olarak üst üste bindirilir
# ----------------------------------------------------------------------------
def _kupur(doc, bolge, onbellek, anahtar, genislik=180, azami_yukseklik=300):
    """Sorunun sayfadaki bölgesini küçük gri tonlamalı görüntü olarak keser."""
    if anahtar in onbellek:
        return onbellek[anahtar]
    sayfa_no, x0, y0, x1, y1 = bolge
    y1 = min(y1, y0 + azami_yukseklik)
    sayfa = doc[sayfa_no]
    r = fitz.Rect(x0, y0, x1, y1) & sayfa.rect
    sonuc = None
    if not r.is_empty and r.width >= 20 and r.height >= 20:
        olcek = genislik / r.width
        pix = sayfa.get_pixmap(clip=r, matrix=fitz.Matrix(olcek, olcek),
                               colorspace=fitz.csGRAY, alpha=False)
        a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width).astype(np.float32)
        if a.shape[0] > 4 and a.shape[1] > 4:  # 3x3 yumuşatma (kayma toleransı)
            b = a.copy()
            b[1:-1, 1:-1] = (a[:-2, :-2] + a[:-2, 1:-1] + a[:-2, 2:] +
                             a[1:-1, :-2] + a[1:-1, 1:-1] + a[1:-1, 2:] +
                             a[2:, :-2] + a[2:, 1:-1] + a[2:, 2:]) / 9.0
            a = b
        sonuc = a
    onbellek[anahtar] = sonuc
    return sonuc


def _gorsel_benzerlik(a, b):
    """İki soru görüntüsünün üst üste bindirme puanı; küçük kaymaları ve
    kitapçıklar arası satır aralığı (leading) farkını tolere eder."""
    if a is None or b is None:
        return None
    ha, hb = a.shape[0], b.shape[0]
    h = min(ha, hb)
    w = min(a.shape[1], b.shape[1])
    if h < 16 or w < 16:
        return None
    oran = h / max(ha, hb)

    def esitle(x):
        # Satır aralığı kısılmış kitapçık için dikey yeniden örnekleme
        if x.shape[0] == h:
            return x[:, :w]
        idx = np.round(np.linspace(0, x.shape[0] - 1, h)).astype(int)
        return x[idx][:, :w]

    if oran >= 0.7:          # makul sıkıştırma: ölçek eşitle
        a2, b2 = esitle(a), esitle(b)
    else:                    # çok farklı yükseklik: üstten ortak kesit
        a2, b2 = a[:h, :w], b[:h, :w]

    en_iyi = None
    for dy in range(-2, 3):
        for dx in (-1, 0, 1):
            aa = a2[max(0, dy):, max(0, dx):]
            bb = b2[max(0, -dy):, max(0, -dx):]
            hh = min(aa.shape[0], bb.shape[0])
            ww = min(aa.shape[1], bb.shape[1])
            if hh < 16 or ww < 16:
                continue
            x, y = aa[:hh, :ww].ravel(), bb[:hh, :ww].ravel()
            if x.std() < 1 or y.std() < 1:
                continue
            r = float(np.corrcoef(x, y)[0, 1])
            en_iyi = r if en_iyi is None else max(en_iyi, r)
    if en_iyi is None:
        return None
    return max(0.0, min(1.0, en_iyi * (0.85 + 0.15 * oran)))


def gorsel_dogrula(a_pdf, b_pdf, a_sorular, b_sorular, eslesme, yapi, log,
                   esik=0.78, cevap_a=None, cevap_b=None, onar=False):
    """Tüm çiftleri görüntü olarak karşılaştırıp raporlar. Varsayılan olarak
    SALT DOĞRULAMA yapar: eşleşmeleri asla değiştirmez, yalnızca şüpheli
    çiftleri işaretler. (onar=True yalnızca test amaçlıdır.)"""
    from collections import defaultdict
    cevap_a, cevap_b = cevap_a or {}, cevap_b or {}
    uyarilar, puanlar = [], {}
    docA, docB = fitz.open(a_pdf), fitz.open(b_pdf)
    onbellek = {}

    def puan(ak, bk):
        ka = _kupur(docA, a_sorular[ak]["bolge"], onbellek, ("A",) + ak)
        kb = _kupur(docB, b_sorular[bk]["bolge"], onbellek, ("B",) + bk)
        return _gorsel_benzerlik(ka, kb)

    for anahtar, es in eslesme.items():
        if (anahtar[0], es["b_no"]) in b_sorular:
            puanlar[anahtar] = puan(anahtar, (anahtar[0], es["b_no"]))

    # (Yalnızca onar=True iken) zayıf METİN çiftlerini görselle yeniden ata
    gruplar = defaultdict(list)
    if onar:
        for k, p in puanlar.items():
            es = eslesme[k]
            if es.get("yontem") != "METİN":
                continue
            ca_ = cevap_a.get(k, "")
            cb_ = cevap_b.get((k[0], es["b_no"]), "")
            celiski = bool(ca_ and cb_ and ca_ != cb_)
            if p is None or p < esik or celiski:
                gruplar[(k[0], a_sorular[k]["ders"])].append(k[1])
    for (ti, ders), a_lar in gruplar.items():
        b_ler = [eslesme[(ti, a)]["b_no"] for a in a_lar]
        if len(a_lar) < 2:
            continue
        M = np.zeros((len(a_lar), len(b_ler)))
        for i, a in enumerate(a_lar):
            for j, b in enumerate(b_ler):
                p = puan((ti, a), (ti, b))
                M[i, j] = p if p is not None else 0.0
        sat, sut = linear_sum_assignment(1 - M)
        for i, j in zip(sat, sut):
            a, b = a_lar[i], b_ler[j]
            if eslesme[(ti, a)]["b_no"] != b:
                eslesme[(ti, a)] = {"b_no": b,
                                    "benzerlik": eslesme[(ti, a)]["benzerlik"],
                                    "yontem": "GÖRSEL"}
                log(f"  ↺ Görsel doğrulama düzeltti: {ders} A-{a} → B-{b}")
            puanlar[(ti, a)] = M[i, j]

    # METİN çiftleri için çapraz kıyas: görsel olarak daha benzer aday var mı?
    ders_b = defaultdict(list)
    for (ti2, b_no2), s in b_sorular.items():
        ders_b[(ti2, s["ders"])].append(b_no2)
    for k, es in list(eslesme.items()):
        if es.get("yontem") != "METİN":
            continue
        p_mevcut = puanlar.get(k)
        taban = p_mevcut if p_mevcut is not None else (esik - 0.05)
        en_iyi_b, en_iyi_p = es["b_no"], taban
        for b in ders_b.get((k[0], a_sorular[k]["ders"]), []):
            if b == es["b_no"]:
                continue
            p2 = puan(k, (k[0], b))
            if p2 is not None and p2 > en_iyi_p + 0.05:
                en_iyi_b, en_iyi_p = b, p2
        if en_iyi_b != es["b_no"]:
            uyarilar.append(
                f"GÖRSEL KONTROL: {a_sorular[k]['ders']} A-{k[1]} şu an "
                f"B-{es['b_no']} ile eşli (görsel %{(p_mevcut or 0)*100:.0f}); "
                f"B-{en_iyi_b} görsel olarak daha benzer (%{en_iyi_p*100:.0f}). "
                f"Elle doğrulayın.")

    tam = 0
    for k in sorted(puanlar):
        p = puanlar[k]
        if p is None:
            continue
        if p >= esik:
            tam += 1
        else:
            es = eslesme[k]
            uyarilar.append(f"GÖRSEL KONTROL: {a_sorular[k]['ders']} A-{k[1]} → "
                            f"B-{es['b_no']} görsel örtüşme %{p*100:.0f} — "
                            f"elle doğrulayın.")
    docA.close()
    docB.close()
    log(f"Görsel doğrulama: {tam}/{len(puanlar)} çift birebir örtüştü.")
    return puanlar, uyarilar


def anahtar_oku(yol, yapi):
    if not yol:
        return {}
    metin = Path(yol).read_text(encoding="utf-8", errors="ignore")
    parcalar = re.split(r"\[([^\]]+)\]", metin)
    cevaplar = {}
    test_adlari = [normalize(t["test"]) for t in yapi]
    for k in range(1, len(parcalar), 2):
        na, ti = normalize(parcalar[k]), None
        for i, t in enumerate(test_adlari):
            if na and (na in t or t in na):
                ti = i
                break
        if ti is None:
            continue
        icerik = parcalar[k + 1]
        ciftler = re.findall(r"(\d{1,3})\s*[-\.\):]?\s*([A-Ea-e])\b", icerik)
        if ciftler:
            for no, harf in ciftler:
                cevaplar[(ti, int(no))] = harf.upper()
        else:
            for no, harf in enumerate(re.findall(r"[A-Ea-e]", icerik), 1):
                cevaplar[(ti, no)] = harf.upper()
    return cevaplar


def _konular_kaynagi(yol):
    """Şablon dosyasından Konular satırlarını okur.
    Dönen: (satirlar, tur)  tur: 'sablon' (Konular sayfası var, dokunma),
    'liste' (yalın kazanım listesi) ya da (None, None)."""
    import openpyxl
    u = str(yol).lower()
    if u.endswith(".csv"):
        import csv
        for kodlama in ("utf-8-sig", "utf-8", "cp1254"):
            try:
                with open(yol, encoding=kodlama, newline="") as f:
                    satirlar = [row for row in csv.reader(f)]
                if satirlar:
                    return satirlar, "liste"
            except UnicodeDecodeError:
                continue
        return None, None
    if u.endswith(".docx"):
        try:
            import docx
        except Exception:
            return None, ("hata: Word şablonu için 'python-docx' kütüphanesi "
                          "kurulamadı. Terminale şunu yapıştırın: "
                          "python3 -m pip install --user python-docx")
        belge = docx.Document(yol)
        satirlar = []
        for tablo in belge.tables:
            for row in tablo.rows:
                hucreler = [h.text.strip() for h in row.cells]
                if any(hucreler):
                    satirlar.append(hucreler)
        if not satirlar:                    # tablo yoksa sekmeli paragraflar
            for p in belge.paragraphs:
                m = p.text.strip()
                if m:
                    satirlar.append(re.split(r"\t+|\s{2,}", m))
        return (satirlar, "liste") if satirlar else (None, None)
    if u.endswith(".pdf"):
        doc = fitz.open(yol)
        satirlar = []
        for sayfa in doc:
            try:                            # önce çizgili tabloları dene
                for tablo in sayfa.find_tables():
                    for row in tablo.extract():
                        hucreler = [str(h).strip() if h is not None else ""
                                    for h in row]
                        if any(hucreler):
                            satirlar.append(hucreler)
            except Exception:
                pass
        if not satirlar:            # çizgisiz: sütunları koordinattan ayır
            for sayfa in doc:
                gruplar = {}
                for x0, y0, x1, y1, w, *_k in sayfa.get_text("words"):
                    anahtar = round((y0 + y1) / 6)
                    gruplar.setdefault(anahtar, []).append((x0, x1, w))
                for _y in sorted(gruplar):
                    parcalar = sorted(gruplar[_y])
                    hucreler, aktif, son_x1 = [], "", None
                    for x0, x1, w in parcalar:
                        if son_x1 is not None and x0 - son_x1 > 6:
                            hucreler.append(aktif)
                            aktif = w
                        else:
                            aktif = (aktif + " " + w).strip()
                        son_x1 = x1
                    if aktif:
                        hucreler.append(aktif)
                    if any(hucreler):
                        satirlar.append(hucreler)
        doc.close()
        return (satirlar, "liste") if satirlar else (None, None)
    wb = openpyxl.load_workbook(yol, data_only=True)
    if "Konular" in wb.sheetnames:
        return None, "sablon"          # şablonun kendi Konular'ı aynen taşınır
    ws = wb.worksheets[0]
    satirlar = [list(r) for r in ws.iter_rows(values_only=True)
                if r and any(c not in (None, "") for c in r)]
    return (satirlar, "liste") if satirlar else (None, None)


def excel_yaz(sablon, sayfa_adi, yapi, eslesme, a_sorular,
              cevap_a, cevap_b, uyarilar):
    sablon_verildi = bool(sablon)
    import openpyxl
    from openpyxl.styles import Font
    BASLIKLAR = ["DersAdi", "Kitapcik", "SoruNo", "Cevap",
                 "B Kitapçığındaki Soru", "K12 Kodu", "Konu Adı",
                 "Çözüm Link", "Çift Soru Var mı?"]
    liste_satirlari = None
    if sablon:
        _s, _tur = _konular_kaynagi(sablon)
        if _tur == "liste":
            liste_satirlari = _s
            sablon = None              # yalın liste: taze tabanda ilerle
        elif isinstance(_tur, str) and _tur.startswith("hata"):
            uyarilar.append(_tur.split(":", 1)[1].strip())
            sablon = None
            sablon_verildi = False
    if sablon:
        wb = openpyxl.load_workbook(sablon)
        ws = wb[sayfa_adi] if sayfa_adi in wb.sheetnames else wb.create_sheet(sayfa_adi)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sayfa_adi
        kws = wb.create_sheet("Konular")
        if liste_satirlari:
            for satir_ in liste_satirlari:   # yüklenen dosya birebir işlenir
                kws.append(list(satir_))
        else:
            kws.append(["K12 Kodu", "Ders", "Sınıf", "Ünite", "Konu",
                        "Alt Kazanım"])
    for c, b in enumerate(BASLIKLAR, 1):
        h = ws.cell(row=1, column=c, value=b)
        h.font = Font(name="Tahoma", bold=True)

    farklar, r = [], 2
    for ti, test in enumerate(yapi):
        n = 0
        for ders, adet in test["dersler"]:
            for _ in range(int(adet)):
                n += 1
                es = eslesme.get((ti, n), {})
                b_no = es.get("b_no", "")
                ca = cevap_a.get((ti, n), "")
                cb = cevap_b.get((ti, b_no), "") if b_no != "" else ""
                if ca and cb and ca != cb:
                    farklar.append(f"CEVAP UYUŞMAZLIĞI: {test['test']} "
                                   f"A-{n} ({ca}) ↔ B-{b_no} ({cb}).")
                ws.cell(row=r, column=1, value=_xl(ders))
                ws.cell(row=r, column=2, value="A")
                ws.cell(row=r, column=3, value=n)
                ws.cell(row=r, column=4, value=ca)
                ws.cell(row=r, column=5, value=b_no)
                ws.cell(row=r, column=6, value=None)
                ws.cell(row=r, column=7,
                        value=f'=IF(F{r}="","",VLOOKUP(F{r},Konular!A:F,5,0))')
                ws.cell(row=r, column=8, value=None)
                ws.cell(row=r, column=9, value=f"=CONCATENATE(A{r},E{r})")
                r += 1
    for eski in range(r, ws.max_row + 1):
        for c in range(1, 10):
            ws.cell(row=eski, column=c, value=None)

    if "Kontrol" in wb.sheetnames:
        del wb["Kontrol"]
    kw = wb.create_sheet("Kontrol")
    kw.append(["Test", "Ders", "A Soru", "B Soru", "Benzerlik %", "Yöntem",
               "A Cevap", "B Cevap", "Soru Kodu (A)", "Durum"])
    for ti, test in enumerate(yapi):
        n = 0
        for ders, adet in test["dersler"]:
            for _ in range(int(adet)):
                n += 1
                es = eslesme.get((ti, n))
                ca_ = cevap_a.get((ti, n), "")
                kod = (a_sorular.get((ti, n)) or {}).get("kod") or ""
                if es:
                    yuzde = round(es["benzerlik"] * 100, 1)
                    yontem = es.get("yontem", "METİN")
                    cb_ = cevap_b.get((ti, es["b_no"]), "")
                    if ca_ and cb_ and ca_ != cb_:
                        durum = "CEVAP FARKLI"
                    elif yontem == "METİN" and yuzde < 60:
                        durum = "KONTROL ET"
                    else:
                        durum = "OK"
                    kw.append([_xl(v) for v in (test["test"], ders, n, es["b_no"],
                                                yuzde, yontem, ca_, cb_, kod, durum)])
                else:
                    kw.append([_xl(v) for v in (test["test"], ders, n, "", "", "",
                                                ca_, "", kod, "EŞLEŞMEDİ")])
    _kn = wb["Konular"] if "Konular" in wb.sheetnames else None
    if sablon_verildi and (_kn is None or not any(
            any(c not in (None, "") for c in r_)
            for r_ in _kn.iter_rows(min_row=2, values_only=True))):
        uyarilar.append("Konular sekmesi boş kaldı — yüklenen kazanım "
                        "dosyasını kontrol edin.")
    kw.append([])
    kw.append(["UYARILAR"])
    for u in uyarilar + farklar:
        kw.append([_xl(u)])
    return wb, farklar


def okunurluk_denetle(*pdfler):
    """PDF'teki yazılar okunabiliyor mu? Yazı tipinin harf eşlemesi (ToUnicode)
    yoksa görüntü doğru olsa da metin anlamsız harflere döner ("Bu testte" →
    "%X WHVWWH"); bu durumda karşılaştırma yapılamaz. Sağlam PDF'lerde bu tür
    karakterler binde birin altında, bozuklarda yüzde 15 civarındadır."""
    for pdf in pdfler:
        harf = bozuk = 0
        with fitz.open(pdf) as doc:
            for sayfa in doc:
                for c in sayfa.get_text():
                    if c.isspace():
                        continue
                    harf += 1
                    bozuk += ord(c) < 32
        if harf and bozuk / harf > 0.02:
            raise RuntimeError(
                f"'{Path(pdf).name}' içindeki yazılar okunamıyor (metnin yaklaşık "
                f"%{bozuk / harf * 100:.0f} kadarı anlamsız karakter). PDF görüntüsü doğru "
                f"olsa da yazı tiplerinin harf eşlemesi kaybolmuş; bu çoğunlukla PDF "
                f"küçültülürken ya da Distiller / 'PDF olarak yazdır' ile yeniden "
                f"kaydedilirken olur. Dizgiden alınan özgün PDF'i seçin. (Küçültmek "
                f"gerekirse Acrobat'ta Dosya > Farklı Kaydet > Küçültülmüş Boyutlu "
                f"PDF yazıları korur.)")


def calistir(a_pdf, b_pdf, sinav, sablon, anah_a, anah_b, cikti, log,
             yol_sor=None):
    log(f"Motor sürümü: {SURUM}")
    okunurluk_denetle(a_pdf, b_pdf)
    if sinav in ("OTOMATİK", "ORTAOKUL"):
        log("Yapı A kitapçığından çıkarılıyor...")
        yapi = yapi_cikar(a_pdf)
        log("  Bulunan yapı: " + " | ".join(
            f"{t['test']} ({sum(int(a) for _d, a in t['dersler'])} soru)"
            for t in yapi))
    else:
        yapi = YAPILAR[sinav]
    log("A kitapçığı okunuyor...")
    dok_a, dok_b = {}, {}     # okunan satırlar içerik doğrulamasında yeniden kullanılır
    a_s, u1 = sorulari_ayikla(a_pdf, yapi, "A", log, dokum=dok_a)
    log("B kitapçığı okunuyor...")
    b_s, u2 = sorulari_ayikla(b_pdf, yapi, "B", log, dokum=dok_b)

    # Kitapçık doğrulaması: soru içermeyen dosya (ör. cevap anahtarı) seçildiyse
    # kullanıcıyı bekletmeden, dosya adıyla açıkça söyle
    def _kitapcik_dogrula(ad, sorular, yol):
        if not sorular:
            raise RuntimeError(
                f"'{Path(yol).name}' içinde hiç soru bulunamadı. Bu dosya bir "
                f"deneme kitapçığı değil gibi görünüyor (örneğin cevap "
                f"anahtarı / 'CA' belgesi seçilmiş olabilir). {ad} Kitapçığı "
                f"alanına soruların olduğu kitapçık PDF'ini seçin.")
    _kitapcik_dogrula("A", a_s, a_pdf)
    _kitapcik_dogrula("B", b_s, b_pdf)

    # Yapı doğrulaması: seçilen tür dosyaya uymuyorsa doğru yapıyı kendisi bulur
    if sinav not in ("OTOMATİK", "ORTAOKUL") and len(u1) + len(u2) > 5:
        for aday_sinav, aday_yapi in YAPILAR.items():
            if aday_sinav == sinav:
                continue
            log(f"Seçilen '{sinav}' yapısına {len(u1)+len(u2)} soru uymadı; "
                f"'{aday_sinav}' yapısı deneniyor...")
            dok_a2, dok_b2 = {}, {}
            a2, v1 = sorulari_ayikla(a_pdf, aday_yapi, "A", log, dokum=dok_a2)
            b2, v2 = sorulari_ayikla(b_pdf, aday_yapi, "B", log, dokum=dok_b2)
            if len(v1) + len(v2) < len(u1) + len(u2):
                log(f"⚠ DİKKAT: Dosyalar '{aday_sinav}' düzeninde; sınav türü "
                    f"otomatik '{aday_sinav}' olarak düzeltildi.")
                sinav, yapi = aday_sinav, aday_yapi
                a_s, b_s, u1, u2 = a2, b2, v1, v2
                dok_a, dok_b = dok_a2, dok_b2
                break
    # TYT de AYT de uymuyorsa (ör. 10. sınıf kitapçığı TYT seçiliyken) yapı
    # kitapçığın kendisinden çıkarılır
    if sinav not in ("OTOMATİK", "ORTAOKUL") and len(u1) + len(u2) > 5:
        try:
            oto = yapi_cikar(a_pdf)
            dok_a3, dok_b3 = {}, {}
            a3, w1 = sorulari_ayikla(a_pdf, oto, "A", log, dokum=dok_a3)
            b3, w2 = sorulari_ayikla(b_pdf, oto, "B", log, dokum=dok_b3)
            if len(w1) + len(w2) < len(u1) + len(u2):
                log("⚠ DİKKAT: Dosyalar TYT/AYT düzeninde değil; yapı kitapçıktan "
                    "çıkarıldı (Lise/Ortaokul seçeneği gibi): " + " | ".join(
                        f"{t['test']} ({sum(int(a) for _d, a in t['dersler'])})"
                        for t in oto))
                sinav, yapi = "OTOMATİK", oto
                a_s, b_s, u1, u2 = a3, b3, w1, w2
                dok_a, dok_b = dok_a3, dok_b3
        except RuntimeError as h:
            logla(f"Otomatik yapı denenemedi: {h}")
    # Tek cevap motoru: sorular sayfa düzeninden ayrılır (A–B KONTROL ile aynı),
    # cevap harfi sorunun oluğundan (numaranın altı / şıkların solu) okunur ve
    # her soru denetlenir. Motor çalışmazsa eski okuma aynen kalır.
    yapilar, cevap_uyarilari = None, []
    try:
        yapilar = _kitapciklari_hazirla(a_pdf, b_pdf, yapi, (a_s, u1, dok_a),
                                        (b_s, u2, dok_b))
        cevap_uyarilari = cevaplari_belirle(yapilar, a_s, b_s, yapi, log)
    except Exception as h:
        logla("Cevap motoru çalışmadı:\n" + traceback.format_exc())
        log(f"  (cevap motoru atlandı: {h})")
    log("Sorular eşleştiriliyor...")
    esl, u3 = eslestir(a_s, b_s, yapi)
    try:                      # bağımsız içerik doğrulaması (kesin yanlışı düzeltir)
        log("Eşleşmeler içerikle doğrulanıyor...")
        eslestirme_dogrula(a_pdf, b_pdf, yapi, esl, u3, log,
                           hazir_a=(a_s, u1, dok_a), hazir_b=(b_s, u2, dok_b),
                           hazir_yapi=yapilar)
    except Exception as h:    # doğrulama yapılamazsa EŞLEŞTİR sonucu aynen kalır
        logla("İçerik doğrulaması yapılamadı:\n" + traceback.format_exc())
        log(f"  (içerik doğrulaması atlandı: {h})")
    pdf_a = {k: s["cevap"] for k, s in a_s.items() if s.get("cevap")}
    pdf_b = {k: s["cevap"] for k, s in b_s.items() if s.get("cevap")}
    log(f"PDF içinden okunan renkli cevaplar: A kitapçığı {len(pdf_a)}/{len(a_s)}, "
        f"B kitapçığı {len(pdf_b)}/{len(b_s)}")
    ca = {**pdf_a, **anahtar_oku(anah_a, yapi)}
    cb = {**pdf_b, **anahtar_oku(anah_b, yapi)}
    uyarilar = u1 + u2 + u3 + cevap_uyarilari
    if a_s and not pdf_a and not anah_a:
        uyarilar.append("A kitapçığında renkli cevap işareti bulunamadı; "
                        "Cevap sütunu boş kalacak.")
    beklenen_toplam = sum(int(a) for t_ in yapi for _d, a in t_["dersler"])
    if a_s and len(a_s) < beklenen_toplam * 0.5:
        log("")
        log("⚠⚠ DİKKAT: A kitapçığında beklenenin çok altında soru bulundu "
            f"({len(a_s)}/{beklenen_toplam}). Seçtiğiniz dosya soru kitapçığı "
            "olmayabilir — örneğin '- CA' ile biten CEVAP ANAHTARI dosyası "
            "seçilmiş olabilir. A satırından 'DENEME A' kitapçığını seçip "
            "tekrar deneyin.")
        log("")
    wb, farklar = excel_yaz(sablon, sinav, yapi, esl, a_s, ca, cb, uyarilar)

    hedef = cikti
    while True:
        try:
            wb.save(hedef)
            break
        except OSError as h:
            log(f"\n⚠ '{hedef}' konumuna yazılamadı ({h}).")
            log("  macOS gizlilik koruması: Masaüstü/Belgeler gibi klasörlere yazmak")
            log("  için konumun kayıt penceresinden seçilmesi gerekir.")
            yeni = yol_sor(hedef) if yol_sor else None
            if not yeni:
                raise RuntimeError("Kayıt tamamlanamadı. Çıktı konumunu penceredeki "
                                   "Seç... düğmesiyle belirleyip tekrar deneyin.") from h
            hedef = yeni

    log(f"\n✓ Tablo yazıldı: {hedef}")
    log(f"  Eşleşen soru: {len(esl)}  |  Uyarı: {len(uyarilar) + len(farklar)}")
    for u in uyarilar + farklar:
        log("  ⚠ " + u)
    if not (uyarilar or farklar):
        log("  Tüm eşleşmeler eşik üstünde, uyarı yok.")
    return hedef


# ----------------------------------------------------------------------------
# A–B KONTROL: B kitapçığını A'ya göre denetler (her A sorusu B'de var mı,
# metin/şık/cevap/görsel aynı mı, metne bağlı gruplar ve numaralar düzgün mü).
# EŞLEŞTİR akışından bağımsızdır; yalnızca PDF okuma ve soru bulma motorunu
# (satirlari_al, sorulari_ayikla) ortak kullanır, onları değiştirmez.
# ----------------------------------------------------------------------------
SECENEK_SATIRI_RE = re.compile(r"^\s*([A-E])\s*(?:\)|\.(?=\s*\S))")
KONTROL_DUR_RE = re.compile(r"^\s*Bu\s+testte\s+\d+\s+soru|"
                            r"^\s*Cevaplar\S*\s+cevap\s+k|TEST[İI]NE\s+GEÇ|"
                            r"TEST[İI]?\s+B[İI]TT[İI]|^\s*CEVAPLARINIZI\s+KONTROL",
                            re.IGNORECASE)
GRUP_NUMARA_RE = re.compile(
    r"(?<![\d.])(\d{1,3}(?:\s*\.?\s*(?:ve|ile|-|–|—|,)\s*\d{1,3})*)"
    r"\s*\.?\s*soru(?:lar|yu)", re.IGNORECASE)
KONTROL_METIN_YOK = 0.50       # metin benzerliği bunun altındaysa "B'de yok"
KONTROL_SEKIL_ESIK = 28.0      # şekil kutusunda küçük bölge gri ton farkı (0-255)


def _grup_basligi_coz(metin, devam=""):
    """'14 ve 15. soruları', '1-4. soruları', '5, 6 ve 7. soruları',
    '12. ve 13. soruları', '9. soruyu' ... → [numaralar]; başlık değilse None."""
    # Satır sonunda tireyle bölünmüş kelimeler ("cevaplayı-/nız") birleştirilir
    ilk = _kiyas_metni(metin)
    tum = _kiyas_metni(metin + "\n" + devam) if devam else ilk
    m = GRUP_NUMARA_RE.search(tum)
    if not m or m.start() > min(80, len(ilk)):
        return None
    tum = tum.casefold()
    if "cevaplay" not in tum:
        return None
    # "16-20. soruları Din Kültürü ... öğrenciler cevaplayacaktır" gibi seçmeli
    # ders yönergesi metin grubu değildir: başlık bir metne/parçaya göndermeli
    if not re.search(r"aşağıdaki|yukarıdaki|verilen|göre", tum):
        return None
    nums, tire = [], False
    for t in re.findall(r"\d{1,3}|[-–—]", m.group(1)):
        if t in "-–—":
            tire = True
            continue
        n = int(t)
        if tire and nums and n > nums[-1]:
            nums.extend(range(nums[-1] + 1, n + 1))
        else:
            nums.append(n)
        tire = False
    if not nums or len(nums) > 12 or nums != sorted(set(nums)) or nums[-1] > 60:
        return None
    return nums


def _renkli_isaret_mi(metin, renk):
    """Magenta cevap harfi ya da soru kodu parçası mı (metinden çıkarılır)."""
    t = metin.strip()
    if not t or not renkli_mi(renk):
        return False
    if re.fullmatch(r"[A-E]", t):
        return True
    t = t.replace(" ", "")
    return bool(re.fullmatch(r"[A-Za-z0-9ÇĞİÖŞÜçğıöşü]+", t)
                and re.search(r"\d", t))


def _sik_harfi(satir):
    m = SECENEK_SATIRI_RE.match(satir["metin"])
    return m.group(1) if m else None


# Görünüşü aynı, kodu farklı karakterler: satır sonu tiresi dizgi programına
# göre "-", yumuşak tire ya da Unicode tire olabilir; "ü" tek harf ya da "u"+"¨"
TIRE_RE = re.compile("[\u2010\u2011\u2012\u2043\ufe63\uff0d]")
TIRELER = "-\u00ad\u2010\u2011\u2012\u2043\ufe63\uff0d"
GORUNMEZ_RE = re.compile("[\u00ad\u200b\u200c\u200d\u2060\ufeff]")
BITISIK_HARF = {"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi",
                "\ufb04": "ffl"}


def _kiyas_metni(metin):
    """Karşılaştırılan metin: satır sonunda bölünmüş kelime birleşir; tire
    çeşitleri, görünmez karakterler, bitişik harfler (ﬁ) ve Unicode yazım
    farkları eşitlenir; boşluklar teke iner."""
    m = unicodedata.normalize("NFC", metin)
    m = "".join(BITISIK_HARF.get(c, c) for c in m)
    m = TIRE_RE.sub("-", m)
    m = re.sub(r"(\w)[-\u00ad][ \t]*\n[ \t]*(\w)", r"\1\2", m)
    m = GORUNMEZ_RE.sub("", m)
    return re.sub(r"\s+", " ", m).strip()


def _gorunur_ayni(a, b):
    """İki karşılaştırma metni baskıda aynı mı? Boşluk ve tire farkı sayılmaz:
    satır kırılması, iki yana yaslamadaki harf aralığı ya da tirenin kodu
    A ile B'de farklı olabilir; okuyan için metin aynıdır."""
    return re.sub(r"[\s\-]", "", a) == re.sub(r"[\s\-]", "", b)


def _temiz_satir(satir, ilk=False):
    m = "".join(t for t, rk in satir.get("spanlar", [])
                if not _renkli_isaret_mi(t, rk)).strip()
    if ilk:
        mm = SORU_BASI_RE.match(m)
        if mm:
            m = mm.group(2).strip()
    return m


def _mobilya_anahtari(satir):
    m = unicodedata.normalize("NFC", satir["metin"]).casefold().strip()
    # Sayfa numarası/kitapçık kodu tek anahtarda toplanır ("12." soru numarası
    # noktalı olduğu için buraya girmez)
    return "#" if re.fullmatch(r"[\d\s/–-]+", m) else re.sub(r"\s+", " ", m)


def _renkli_harf_satiri(satir):
    """Yalnızca renkli tek harften (A–E) oluşan satır: cevap işareti olabilir.
    Sayfa üst bilgisi sayılması için her sayfada aynı yerde durmalıdır (kitapçık
    harfi "A A A A A" gibi); iki sayfada aynı köşede duran cevap harfi —
    şıkları sayfa başına taşan sorularda olur — üst bilgi sayılmaz."""
    spanlar = [(t.strip(), c) for t, c in satir.get("spanlar", []) if t.strip()]
    return (len(spanlar) == 1 and re.fullmatch(r"[A-E]", spanlar[0][0]) is not None
            and renkli_mi(spanlar[0][1]))


def _kenar_bolgesi(satir, sayfa_yuk, sayfa_gen):
    return (satir["y1"] < 0.12 * sayfa_yuk or satir["y0"] > 0.90 * sayfa_yuk
            or satir["x1"] < 0.08 * sayfa_gen or satir["x0"] > 0.92 * sayfa_gen)


def _mobilya_kumesi(kitapciklar):
    """Sayfa üst/alt bilgisi, kenar kodu gibi tekrar eden satırlar
    (karşılaştırma metninden çıkarılır; A ve B'de simetrik). Yalnızca konuma
    bakılır — soru numarası ya da kalıp soru kökü gibi sık geçen içerik
    satırları sayfa bilgisi sayılmaz:
      1) kenar boşluğunda (üst %12, alt %10, yanlar %8) en az iki sayfada aynı
         yerde (yükseklik ve sol kenar) duran satır (rakamları farklı olabilir:
         sayfa no vb.; "9." gibi soru numarası satırları hariç)
      2) sayfanın her yerinde, sayfaların yarısında birebir aynı yerde duran satır."""
    kume = set()
    for k in kitapciklar:
        satirlar, yuk, gen = k["satirlar"], k["yukseklikler"], k["genislikler"]
        sayfa_say = 1 + max((r["sayfa"] for r in satirlar), default=0)
        kenar, sabit = {}, {}
        for r in satirlar:
            if SECENEK_SATIRI_RE.match(r["metin"]):
                continue                     # şıklar asla sayfa bilgisi değildir
            m = _mobilya_anahtari(r)
            if (_kenar_bolgesi(r, yuk[r["sayfa"]], gen[r["sayfa"]])
                    and not SORU_BASI_RE.match(r["metin"])
                    and not _renkli_harf_satiri(r)):
                for anahtar in _kenar_anahtarlari(r, m):
                    kenar.setdefault(anahtar, set()).add(r["sayfa"])
            sabit.setdefault((m, round(r["x0"] / 3), round(r["y0"] / 3)),
                             set()).add(r["sayfa"])
        kume |= {("kenar", k_) for k_, s_ in kenar.items() if len(s_) >= 2}
        kume |= {("sabit", k_) for k_, s_ in sabit.items()
                 if len(s_) >= max(3, 0.5 * sayfa_say)}
    return kume


def _mobilya_mi(satir, kume, sayfa_yuk, sayfa_gen):
    if SECENEK_SATIRI_RE.match(satir["metin"]):
        return False
    m = _mobilya_anahtari(satir)
    if ("sabit", (m, round(satir["x0"] / 3), round(satir["y0"] / 3))) in kume:
        return True
    if SORU_BASI_RE.match(satir["metin"]) or _renkli_harf_satiri(satir) or \
            not _kenar_bolgesi(satir, sayfa_yuk, sayfa_gen):
        return False
    # Komşu konumlara da bakılır (yuvarlama sınırında kalan satırlar için)
    for kaba, ky, tur, kx in _kenar_anahtarlari(satir, m):
        for dy in (-1, 0, 1):
            for dx in (-1, 0, 1):
                if ("kenar", (kaba, ky + dy, tur, kx + dx)) in kume:
                    return True
    return False


def _kenar_anahtarlari(satir, m):
    """Kenar satırının konum anahtarları: aynı yükseklik + aynı sol kenar ya da
    aynı orta nokta (ortalı sayfa numarası "1" ile "12"nin sol kenarı farklıdır)."""
    kaba = re.sub(r"\d+", "#", m)
    ky = round(satir["y0"] / 3)
    return [(kaba, ky, "x0", round(satir["x0"] / 6)),
            (kaba, ky, "xo", round((satir["x0"] + satir["x1"]) / 12))]


def _satir_sinifi(r, orta):
    """F: orta çizginin iki yanına da belirgin taşan (tam genişlik) satır;
    L/R: sol/sağ yarıdaki satır (sütun kenarını birkaç punto aşan satır
    kendi sütununda kalır)."""
    if orta - r["x0"] > 40 and r["x1"] - orta > 40:
        return "F"
    return "L" if (r["x0"] + r["x1"]) / 2 < orta else "R"


def _satir_dizisi(satirlar):
    """Satırları yukarıdan aşağı, aynı yükseklikteki (±3 pt) satırları
    soldan sağa dizer (yan yana şıklar A B C D sırasıyla okunur)."""
    gruplar = []
    for r in sorted(satirlar, key=lambda r: r["y0"]):
        if gruplar and abs(r["y0"] - gruplar[-1][0]) <= 3:
            gruplar[-1][1].append(r)
        else:
            gruplar.append([r["y0"], [r]])
    return [r for _y, g in gruplar for r in sorted(g, key=lambda r: r["x0"])]


def _bant_sirasi(sat, sinif):
    """Soru numarası olmayan sayfa için: tam genişlik satırları bant ayırır,
    her bantta önce sol sonra sağ yarı okunur."""
    sira, bant = [], []

    def bosalt():
        for taraf in ("L", "R"):
            sira.extend(_satir_dizisi([r for r in bant if sinif[id(r)] == taraf]))
        bant.clear()
    for r in sorted(sat, key=lambda r: (round(r["y0"], 1), r["x0"])):
        if sinif[id(r)] == "F":
            bosalt()
            sira.append(r)
        else:
            bant.append(r)
    bosalt()
    return sira


def _okuma_sirasi(satirlar, genislikler, capa_sirasi, yukseklikler=None):
    """Soru numaralarına dayalı okuma sırası. Her soru numarasından bir bölge
    başlar: soru gövdesi tam genişlikse sayfa genişliğinde, değilse kendi
    sütununda; bölge aşağıdaki ilk soru numarasında (sütun sorusu için aynı
    sütundaki ya da tam genişlikteki) biter. Satır, içinde bulunduğu bölgenin
    sorusuna yazılır; böylece aynı sayfada tam genişlik soru (yan yana şıklarıyla)
    ve iki sütun bir arada olsa da her satır kendi sorusunda kalır. Hiçbir
    bölgeye düşmeyen satırlar (sütun başına taşan devam, sayfa başındaki metin)
    altlarındaki ilk uygun sorudan hemen önce okunur.
    capa_sirasi: {id(soru numarası satırı): beklenen sıra}.
    Dönen: (sıra, satır sınıfı, soru sınıfı {id: F/L/R}, iki sütunlu sayfalar,
    motor sırasına dönülen sayfalar)."""
    sayfalar = {}
    for r in satirlar:
        sayfalar.setdefault(r["sayfa"], []).append(r)
    sira, sinif, capa_sinif, iki_sutun, geri_donulen = [], {}, {}, {}, []
    for pno in sorted(sayfalar):
        sat = sayfalar[pno]
        orta = genislikler[pno] / 2
        for r in sat:
            sinif[id(r)] = _satir_sinifi(r, orta)
        capalar = [r for r in sat if id(r) in capa_sirasi]
        if not capalar:
            iki_sutun[pno] = False
            sira.extend(_bant_sirasi(sat, sinif))
            continue
        yuk = yukseklikler[pno] if yukseklikler else 1e9
        govde_sat = [r for r in sat if 0.12 * yuk < r["y0"] < 0.90 * yuk]

        def yan(r):
            return "L" if (r["x0"] + r["x1"]) / 2 < orta else "R"

        def tam_genislik_mi(c, karsi_dolu):
            """Soru tam genişlik mi? Numaranın yanındaki ilk gövde satırları orta
            çizgiyi aşıyorsa evet. Aşmıyorsa (numara tek başına, gövde şeklin
            altında olabilir) karşı sütunda üstte süren bir soru yoksa ve bir
            sonraki numaraya kadar tam genişlik satır ya da iki yarıya yayılan
            ardışık şık satırı (A) … B) …) varsa yine evet."""
            if any(x is not c and yan(x) != yan(c) and abs(x["y0"] - c["y0"]) < 30
                   for x in capalar):
                return False
            govde = sorted((r for r in sat if r is not c
                            and c["x0"] - 2 <= r["x0"] <= c["x0"] + 60
                            and c["y0"] - 4 <= r["y0"] <= c["y0"] + 40),
                           key=lambda r: r["y0"])[:3]
            if any(sinif[id(r)] == "F" for r in govde + [c]):
                return True
            if karsi_dolu:
                return False
            alt = min([x["y0"] for x in capalar if x["y0"] > c["y0"] + 2] + [1e9])
            aralik = [r for r in govde_sat if c["y0"] - 4 <= r["y0"] < alt]
            if any(sinif[id(r)] == "F" for r in aralik):
                return True
            siralar = {}
            for r in aralik:
                h = SECENEK_SATIRI_RE.match(r["metin"])
                if h:
                    siralar.setdefault(round(r["y0"] / 3), []).append(
                        (r["x0"], h.group(1), yan(r)))
            for satir in siralar.values():
                harfler = "".join(h for _x, h, _y in sorted(satir))
                if (len({y for _x, _h, y in satir}) == 2 and len(harfler) >= 2
                        and harfler in "ABCDE"):
                    return True
            return False

        # Yukarıdan aşağı: karşı sütunda, son tam genişlik sorudan sonra başlamış
        # bir sütun sorusu varsa o yarı "dolu"dur
        son_tam_y = -1e9
        for c in sorted(capalar, key=lambda r: r["y0"]):
            karsi_dolu = any(capa_sinif.get(id(x)) in ("L", "R") and yan(x) != yan(c)
                             and son_tam_y < x["y0"] < c["y0"] - 2 for x in capalar)
            capa_sinif[id(c)] = "F" if tam_genislik_mi(c, karsi_dolu) else yan(c)
            if capa_sinif[id(c)] == "F":
                son_tam_y = c["y0"]
        iki_sutun[pno] = any(capa_sinif[id(c)] != "F" for c in capalar)

        # Soruların sayfa üzerindeki okuma sırası (tam genişlik soru bant ayırır)
        geo, bant = [], []
        for c in sorted(capalar, key=lambda r: (r["y0"], r["x0"])):
            if capa_sinif[id(c)] == "F":
                geo.extend(sorted((x for x in bant if capa_sinif[id(x)] == "L"),
                                  key=lambda r: r["y0"]))
                geo.extend(sorted((x for x in bant if capa_sinif[id(x)] == "R"),
                                  key=lambda r: r["y0"]))
                bant = []
                geo.append(c)
            else:
                bant.append(c)
        geo.extend(sorted((x for x in bant if capa_sinif[id(x)] == "L"),
                          key=lambda r: r["y0"]))
        geo.extend(sorted((x for x in bant if capa_sinif[id(x)] == "R"),
                          key=lambda r: r["y0"]))
        dizi = [capa_sirasi[id(c)] for c in geo]
        if dizi != sorted(dizi):
            sira.extend(sat)                     # motorun sayfa içi sırası
            geri_donulen.append(pno + 1)
            continue

        def taraf(r):          # tam genişlik satır da ortasına göre bir yana
            return "L" if (r["x0"] + r["x1"]) / 2 < orta else "R"

        def sutununda(c, r):   # r satırı c sorusunun sütununda mı
            return capa_sinif[id(c)] in ("F", taraf(r))

        bolgeler = {}
        for c in capalar:
            t = capa_sinif[id(c)]
            alt = min([x["y0"] for x in capalar if x["y0"] > c["y0"] + 2
                       and (t == "F" or capa_sinif[id(x)] in ("F", t))]
                      + [1e9]) - 2
            bolgeler[id(c)] = (c["y0"] - 4, alt)
        def numara_hizasinda(c, r):
            """r, c numarasıyla aynı satırda mı? Üst üste kesir (315/2) ya da üs
            satırı yukarı taşırsa satırın üst kenarı numaranın üstünde kalır;
            yine de numaranın sağında ve onun hizasındaysa o sorunun ilk
            satırıdır (bir önceki soruya yazılmamalı)."""
            return (c["y0"] - 30 < r["y0"] < c["y0"] - 4
                    and r["y1"] > c["y0"] + 2
                    and r["x0"] >= c["x1"] - 2 and sutununda(c, r))

        atanan = {id(c): [] for c in capalar}
        yetim = []
        for r in sat:
            if id(r) in atanan:
                continue
            hizada = [c for c in capalar if numara_hizasinda(c, r)]
            if hizada:
                atanan[id(min(hizada, key=lambda c: r["x0"] - c["x1"]))].append(r)
                continue
            uygun = [c for c in capalar
                     if bolgeler[id(c)][0] <= r["y0"] < bolgeler[id(c)][1]
                     and sutununda(c, r)]
            if uygun:
                atanan[id(max(uygun, key=lambda c: c["y0"]))].append(r)
            else:
                yetim.append(r)
        once = {id(c): [] for c in capalar}
        kuyruk = []
        for r in yetim:
            hedef = next((c for c in geo if c["y0"] > r["y0"] and
                          (sinif[id(r)] == "F" or sutununda(c, r))), None)
            (once[id(hedef)] if hedef is not None else kuyruk).append(r)
        for c in geo:
            sira.extend(_satir_dizisi(once[id(c)]))
            sira.append(c)
            sira.extend(_satir_dizisi(atanan[id(c)]))
        sira.extend(_satir_dizisi(kuyruk))
    return sira, sinif, capa_sinif, iki_sutun, geri_donulen


def _kitapcik_oku(pdf, yapi, etiket, log, hazir=None):
    """hazir: daha önce okunmuş (sorular, uyarılar, döküm) — PDF yeniden okunmaz."""
    if hazir is not None:
        sorular, uyarilar, dokum = hazir[0], hazir[1], dict(hazir[2])
    else:
        dokum = {}
        sorular, uyarilar = sorulari_ayikla(pdf, yapi, etiket, log, dokum=dokum)
    doc = fitz.open(pdf)
    genislikler = [p.rect.width for p in doc]
    yukseklikler = [p.rect.height for p in doc]
    doc.close()
    dokum.update(pdf=pdf, etiket=etiket, sorular=sorular, uyarilar=uyarilar,
                 genislikler=genislikler, yukseklikler=yukseklikler)
    return dokum


def _kitapcik_bolumle(k, mobilya, log):
    """Soruları okuma sırasına göre böler; her sorunun temiz metnini (gövde +
    şıklar), bölgesini, metne bağlı grup başlıklarını ve numarası beklenen
    sıraya uymayan şıklı soru başlangıçlarını çıkarır."""
    satirlar, beklenen = k["satirlar"], k["beklenen"]
    capa_sirasi = {id(satirlar[s]): b for b, s in k["eslesen"].items()}
    sira, sinif, capa_sinif, iki_sutun, geri = _okuma_sirasi(
        satirlar, k["genislikler"], capa_sirasi, k["yukseklikler"])
    if geri:
        logla(f"{k['etiket']}: şu sayfalarda motorun satır sırası kullanıldı: "
              f"{geri}")

    def mob(r):
        return _mobilya_mi(r, mobilya, k["yukseklikler"][r["sayfa"]],
                           k["genislikler"][r["sayfa"]])
    poz = {id(r): i for i, r in enumerate(sira)}
    capalar = sorted((poz[id(satirlar[s])], b) for b, s in k["eslesen"].items())

    def secenekli(bas, son):
        return any(SECENEK_SATIRI_RE.match(sira[j]["metin"])
                   for j in range(bas, min(son, len(sira))))

    # --- metne bağlı grup başlıkları -----------------------------------------
    basliklar = []
    for i, r in enumerate(sira):
        devam = "\n".join(x["metin"] for x in sira[i + 1:i + 3]
                          if x["sayfa"] == r["sayfa"])
        nums = _grup_basligi_coz(r["metin"], devam)
        if nums:
            bitis = i            # başlık cümlesi alt satıra taştıysa onu da al
            for j in range(i, min(i + 3, len(sira))):
                if sira[j]["sayfa"] != r["sayfa"]:
                    break
                if "cevaplay" in _kiyas_metni("\n".join(
                        x["metin"] for x in sira[i:j + 1])).casefold():
                    bitis = j
                    break
            # Başlığın son satırı tireyle bölünmüşse ("cevaplayı-") devamı da başlık
            while (bitis + 1 < min(i + 4, len(sira))
                   and sira[bitis]["metin"].rstrip()[-1:] in TIRELER
                   and sira[bitis + 1]["sayfa"] == r["sayfa"]):
                bitis += 1
            basliklar.append({"poz": i, "bitis": bitis, "satir": r,
                              "nolar": nums})
    baslik_poz = {b["poz"] for b in basliklar}
    aday_pozlar = sorted({poz[id(satirlar[s])]: n for s, n in
                          k["adaylar"] + k["yedekler"]}.items())
    # Motor "12. ve 13. soruları ..." başlık satırını 12. sorunun çapası
    # saymışsa gerçek "12." satırına kaydır (başlık + parça soruya karışmasın)
    duzeltilmis = []
    for i, (p, b_i) in enumerate(capalar):
        if p in baslik_poz:
            no = beklenen[b_i][2]
            sinir = capalar[i + 1][0] if i + 1 < len(capalar) else len(sira)
            gercek = next((q for q, n in aday_pozlar
                           if p < q < sinir and n == no), None)
            if gercek is not None:
                p = gercek
        duzeltilmis.append((p, b_i))
    capalar = duzeltilmis

    # --- numarası beklenene uymayan şıklı soru başlangıçları ------------------
    capa_poz = [p for p, _b in capalar]
    kullanilan = {id(sira[p]) for p in capa_poz}
    tum_aday = [p for p, _n in aday_pozlar]
    # Sırasız numara yalnızca sütun kenarına hizalı numaralarda aranır (sorunun
    # içindeki numaralı madde ya da alt satıra kayan şık parçası sayılmaz)
    kenar_adaylar = sorted({poz[id(satirlar[s])]: n for s, n in k["adaylar"]}.items())
    fazlalar = []
    for p, n in kenar_adaylar:
        r = sira[p]
        if id(r) in kullanilan or p in baslik_poz or mob(r):
            continue
        onceki = max((q for q in capa_poz if q < p), default=None)
        if onceki is None:
            continue
        sonraki = min([q for q in tum_aday if q > p] + [len(sira)])
        if secenekli(onceki, p) and secenekli(p, sonraki):
            b_i = next(b for q, b in capalar if q == onceki)
            fazlalar.append({"poz": p, "no_yazan": n, "ti": beklenen[b_i][0],
                             "ders": beklenen[b_i][1]})

    # Mükerrer numarada hangisi yanlış? Sıra 5, 7, 7, 8 ise (6 eksik) 6'nın
    # yerinde duran ilk "7." yanlış numaralıdır; motor ilkini 7 saymışsa
    # rolleri değiştir
    capali = {(beklenen[b_][0], beklenen[b_][2]) for _p, b_ in capalar}
    for f in fazlalar:
        onceki = max(((p_, b_) for p_, b_ in capalar if p_ < f["poz"]), default=None)
        if onceki is None:
            continue
        ti_, _d, no_ = beklenen[onceki[1]]
        if ti_ != f["ti"] or no_ != f["no_yazan"] or no_ < 2 or \
                (ti_, no_ - 1) in capali:
            continue
        capalar.remove(onceki)
        capalar.append((f["poz"], onceki[1]))
        f["poz"] = onceki[0]
    capalar.sort()

    # --- soru bölümleri --------------------------------------------------------
    sinirlar = sorted([(p, ("S", b)) for p, b in capalar] +
                      [(f["poz"], ("F", j)) for j, f in enumerate(fazlalar)])
    ogeler = []
    for s_i, (p, (tur, deg)) in enumerate(sinirlar):
        son = sinirlar[s_i + 1][0] if s_i + 1 < len(sinirlar) else len(sira)
        bolum = sira[p:son]
        kes = len(bolum)
        for j in range(1, len(bolum)):
            if p + j in baslik_poz or KONTROL_DUR_RE.search(bolum[j]["metin"]):
                kes = j
                break
        bolum = bolum[:kes]
        sec_idx = [j for j, r in enumerate(bolum)
                   if SECENEK_SATIRI_RE.match(r["metin"])]
        if sec_idx:
            son_sec = sec_idx[-1]
            adimlar = sorted(bolum[a + 1]["y0"] - bolum[a]["y0"]
                             for a in range(sec_idx[0], son_sec)
                             if bolum[a + 1]["y0"] > bolum[a]["y0"])
            tipik = adimlar[len(adimlar) // 2] if adimlar else 14
            j = son_sec + 1
            while (j < len(bolum) and bolum[j]["sayfa"] == bolum[j - 1]["sayfa"]
                   and -4 <= bolum[j]["y0"] - bolum[j - 1]["y0"] < 1.8 * tipik):
                j += 1                   # son şıkkın devam satırları (aynı satırdaki
                #                          birkaç puntoluk taban farkı dahil)
            bolum = bolum[:j]
        tutulan = [(j, r) for j, r in enumerate(bolum)
                   if j == 0 or not mob(r)]
        ham_bolum = [r for _j, r in tutulan]   # açık cevap taraması (sayfa
        #                                        üstündeki renkli harfler hariç)
        parcalar = [_temiz_satir(r, ilk=(j == 0)) for j, r in tutulan]
        metin = "\n".join(x for x in parcalar if x)
        cmp = _kiyas_metni(metin)
        norm = normalize(metin)
        bas = sira[p]
        if tur == "S":
            ti, ders, no = beklenen[deg]
            etiket_no = str(no)
        else:
            f = fazlalar[deg]
            ti, ders, no = f["ti"], f["ders"], None
            etiket_no = f"'{f['no_yazan']}.' (numara hatalı)"
        ogeler.append({
            "ti": ti, "ders": ders, "no": no, "etiket_no": etiket_no,
            "poz": p, "sayfa": bas["sayfa"], "metin": metin, "cmp": cmp,
            "norm": norm, "tri": {norm[i:i + 3] for i in range(len(norm) - 2)},
            "a_sayisi": sum(1 for _j, r in tutulan if _sik_harfi(r) == "A"),
            "bolge": _kontrol_bolge(bas, [r for _j, r in tutulan], sinif,
                                    iki_sutun, satirlar,
                                    taraf=capa_sinif.get(id(bas))),
            "nesne_bolgeleri": _nesne_bolgeleri(
                [r for _j, r in tutulan], sinif, k["genislikler"],
                tam_sayfa=bas["sayfa"] if capa_sinif.get(id(bas)) == "F" else None),
            "ham_satirlar": ham_bolum,
            "cevap": (k["sorular"].get((ti, no)) or {}).get("cevap") if no else None,
            "kod": (k["sorular"].get((ti, no)) or {}).get("kod") if no else None,
        })

    # --- başlıkların altındaki sorular ve parça metni ---------------------------
    for b in basliklar:
        sonrakiler = [o for o in ogeler if o["poz"] > b["poz"]][:len(b["nolar"])]
        b["sorular"] = sonrakiler
        ilk_poz = sonrakiler[0]["poz"] if sonrakiler else len(sira)
        parca = [r for r in sira[b["bitis"] + 1:ilk_poz]
                 if not mob(r)
                 and not KONTROL_DUR_RE.search(r["metin"])]
        b["metin"] = "\n".join(_temiz_satir(r) for r in parca)
        b["cmp"] = _kiyas_metni(b["metin"])
        b["ti"] = sonrakiler[0]["ti"] if sonrakiler else None
        b["sayfa"] = b["satir"]["sayfa"]
        bas_sat = [x for x in sira[b["poz"]:b["bitis"] + 1]
                   if x["sayfa"] == b["sayfa"]]
        b["baslik_bolge"] = (b["sayfa"], min(x["x0"] for x in bas_sat) - 4,
                             min(x["y0"] for x in bas_sat) - 3,
                             max(x["x1"] for x in bas_sat) + 4,
                             max(x["y1"] for x in bas_sat) + 3)
        ilk_soru = sira[ilk_poz] if ilk_poz < len(sira) else None
        # Başlık satırı kısa olsa da parça tam genişlikse bölge de tam genişlik
        tam = any(sinif.get(id(r)) == "F" for r in parca
                  if r["sayfa"] == b["sayfa"])
        # Görsel bölge başlığın altından başlar (başlıktaki numaralar A ile
        # B'de doğal olarak farklıdır; onlar ayrıca denetleniyor)
        ilk_parca = next((r for r in parca if r["sayfa"] == b["sayfa"]), None)
        b["nesne_bolgeleri"] = _nesne_bolgeleri(
            parca, sinif, k["genislikler"], tam_sayfa=b["sayfa"] if tam else None)
        b["bolge"] = _kontrol_bolge(ilk_parca or b["satir"], parca, sinif,
                                    iki_sutun, satirlar, bitis=ilk_soru,
                                    taraf=("F" if tam else
                                           sinif.get(id(b["satir"]), "F")))
    k.update(sira=sira, ogeler=ogeler, basliklar=basliklar, fazlalar=fazlalar)
    return k


def _kontrol_bolge(bas, tutulan, sinif, iki_sutun, satirlar, bitis=None,
                   taraf=None):
    """Sorunun (ya da parçanın) ilk sayfadaki dikdörtgeni: (sayfa,x0,y0,x1,y1)."""
    pno = bas["sayfa"]
    sayfa_sat = [r for r in satirlar if r["sayfa"] == pno]
    taraf = taraf or sinif.get(id(bas), "F")
    if iki_sutun.get(pno) and taraf in ("L", "R"):
        ayni = [r for r in tutulan if r["sayfa"] == pno
                and sinif.get(id(r)) == taraf]
        # Bölge sütun metninin sol kenarından başlar, genişliği iki sütunda
        # aynıdır: A'da solda, B'de sağda duran aynı soru üst üste oturur
        # (kısa satırlar — sayfa no, cevap harfi — ve kenarı taşan tek tük
        # satırlar sütun sınırını bozmasın diye yüzdelik kullanılır)
        kenarlar = {}
        for r in sayfa_sat:
            t_ = sinif.get(id(r))
            if t_ in ("L", "R") and len(r["metin"]) >= 12:
                kenarlar.setdefault(t_, ([], []))
                kenarlar[t_][0].append(r["x0"])
                kenarlar[t_][1].append(r["x1"])

        def yuzdelik(d, q):
            d = sorted(d)
            return d[min(len(d) - 1, int(q * len(d)))]
        sinirlar = {t_: (yuzdelik(a, 0.1), yuzdelik(b, 0.9))
                    for t_, (a, b) in kenarlar.items()}
        if taraf not in sinirlar:
            sinirlar[taraf] = (bas["x0"], max(r["x1"] for r in ayni or [bas]))
        gen = max(b - a for a, b in sinirlar.values())
        x0 = sinirlar[taraf][0] - 2
        x1 = x0 + gen + 4
    else:
        ayni = [r for r in tutulan if r["sayfa"] == pno]
        x0 = min(r["x0"] for r in sayfa_sat) - 2
        x1 = max(r["x1"] for r in sayfa_sat) + 2
    # Yatay hiza şıkların girintisine göre: soru numarası ve magenta harf
    # dışarıda kalır, A'da solda B'de sağda duran soru birebir üst üste biner
    sik = next((r for r in ayni if _sik_harfi(r)), None)
    if sik is not None and sik["x0"] - 2 > x0:
        x0 = sik["x0"] - 2
    y1 = max([r["y1"] for r in ayni] + [bas["y1"]]) + 3
    if (bitis is not None and bitis["sayfa"] == pno and bitis["y0"] > bas["y1"]
            and (taraf == "F" or sinif.get(id(bitis)) == taraf
                 or not iki_sutun.get(pno))):
        y1 = bitis["y0"] - 3             # parça: ilk soruya kadar (tablo/şekil dahil)
    return (pno, x0, bas["y0"] - 3, x1, y1)


def _kontrol_kupur(doc, bolge, onbellek, anahtar, yukseklik=None, olcek=0.8):
    """Bölgenin gri tonlamalı görüntüsü; magenta cevap harfi/kodu beyazlatılır
    (baskı PDF'inde olmadıkları için karşılaştırmayı bozmasınlar)."""
    if anahtar in onbellek:
        return onbellek[anahtar]
    pno, x0, y0, x1, y1 = bolge
    if yukseklik is not None:
        y1 = min(y1, y0 + yukseklik)
    y1 = min(y1, y0 + 500)
    sayfa = doc[pno]
    r = fitz.Rect(x0, y0, x1, y1) & sayfa.rect
    sonuc = None
    if not r.is_empty and r.width >= 20 and r.height >= 12:
        anahtar_d = ("dl", id(doc), pno)       # sayfa çizim listesi bir kez kurulur
        if anahtar_d not in onbellek:
            onbellek[anahtar_d] = sayfa.get_displaylist()
        pix = onbellek[anahtar_d].get_pixmap(
            matrix=fitz.Matrix(olcek, olcek), colorspace=fitz.csGRAY,
            alpha=False, clip=r)
        a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width).astype(np.float32)
        anahtar_i = ("isaret", id(doc), pno)
        if anahtar_i not in onbellek:          # (kontrol_et önceden doldurur)
            onbellek[anahtar_i] = [
                fitz.Rect(sp["bbox"])
                for blk in sayfa.get_text("dict").get("blocks", [])
                for ln in blk.get("lines", []) for sp in ln.get("spans", [])
                if _renkli_isaret_mi(sp.get("text", ""), sp.get("color", 0))]
        isaretler = onbellek[anahtar_i]
        for ir in isaretler:
            kes = ir & r
            if kes.is_empty:
                continue
            ya, yb = int((kes.y0 - r.y0) * olcek), int((kes.y1 - r.y0) * olcek) + 1
            xa, xb = int((kes.x0 - r.x0) * olcek), int((kes.x1 - r.x0) * olcek) + 1
            a[max(0, ya):yb, max(0, xa):xb] = 255.0
        if a.shape[0] > 4 and a.shape[1] > 4:
            b = a.copy()
            b[1:-1, 1:-1] = (a[:-2, :-2] + a[:-2, 1:-1] + a[:-2, 2:] +
                             a[1:-1, :-2] + a[1:-1, 1:-1] + a[1:-1, 2:] +
                             a[2:, :-2] + a[2:, 1:-1] + a[2:, 2:]) / 9.0
            a = b
        sonuc = a
    onbellek[anahtar] = sonuc
    return sonuc


def _kupur_cifti(docA, docB, bolgeA, bolgeB, onbellek):
    """İki bölgeyi aynı yükseklik ve genişlikte (kısa/dar olanınki) keser."""
    h = round(min(bolgeA[4] - bolgeA[2], bolgeB[4] - bolgeB[2]), 1)
    ga, gb = bolgeA[3] - bolgeA[1], bolgeB[3] - bolgeB[1]
    if h < 12 or abs(ga - gb) > 0.25 * max(ga, gb):
        return None, None            # düzen farklı (sütun/tam genişlik): kıyas yok
    ka = _kontrol_kupur(docA, bolgeA, onbellek, ("A", bolgeA))
    kb = _kontrol_kupur(docB, bolgeB, onbellek, ("B", bolgeB))
    if ka is None or kb is None:
        return None, None
    # Aynı ölçekte (0.8 px/pt) ve aynı sol üst köşeden: kısa/dar olana kırp
    n_h, n_w = int(h * 0.8), int(min(ga, gb) * 0.8)
    return ka[:n_h, :n_w], kb[:n_h, :n_w]


def _gorsel_puan(docA, docB, bolgeA, bolgeB, onbellek):
    """Genel örtüşme puanı (0-1)."""
    return _gorsel_benzerlik(*_kupur_cifti(docA, docB, bolgeA, bolgeB, onbellek))


def _sayfa_nesneleri(doc, pno, onbellek):
    """Sayfadaki çizim ve resim parçalarının GÖRÜNEN dikdörtgenleri: her çizim
    parçası bağlı olduğu kırpma alanıyla kesilir (kırpılmış vektör resmin
    görünmeyen parçaları düşer); sayfa/sütun ayırıcı çizgiler alınmaz."""
    anahtar = ("nesne", id(doc), pno)
    if anahtar not in onbellek:
        sayfa = doc[pno]
        W = sayfa.rect.width
        dikdortgenler = []
        yigin = []                             # (düzey, kırpma alanı)
        for d in sayfa.get_drawings(extended=True):
            duzey = d.get("level", 0)
            while yigin and yigin[-1][0] >= duzey:
                yigin.pop()
            ust = yigin[-1][1] if yigin else sayfa.rect
            if d.get("type") == "clip":
                yigin.append((duzey, fitz.Rect(d["scissor"]) & ust))
                continue
            if d.get("type") not in ("f", "s", "fs"):
                continue
            r = fitz.Rect(d["rect"])
            if r.width < 0.5 and r.height < 0.5:
                continue
            gorunen = fitz.Rect(max(r.x0, ust.x0), max(r.y0, ust.y0),
                                min(r.x1, ust.x1), min(r.y1, ust.y1))
            if gorunen.x1 < gorunen.x0 or gorunen.y1 < gorunen.y0:
                continue                       # tamamen kırpılmış
            dikdortgenler.append(gorunen)
        # Resimler (çizim listesinde yer almaz): çizim kaydından, hızlı
        for tur, kutu in sayfa.get_bboxlog():
            if tur in ("fill-image", "fill-imgmask"):
                dikdortgenler.append(fitz.Rect(kutu))
        dikdortgenler = [
            r for r in dikdortgenler
            if not (r.width > 0.6 * W and r.height < 3)              # yatay çizgi
            and not (r.width < 3 and abs((r.x0 + r.x1) / 2 - W / 2) < 8)]  # sütun
        onbellek[anahtar] = dikdortgenler
    return onbellek[anahtar]


def _murekkep_anahtari(doc, pno, kutu):
    return ("murekkep_say", id(doc), pno, tuple(round(x, 1) for x in kutu))


def _murekkep_alani(doc, pno, kutu, onbellek, olcek=1.0):
    """Kutunun görünür (mürekkepli) kısmının dikdörtgeni; tamamen boşsa None.
    Mürekkepli piksel sayısı onbellek[("murekkep_say", ...)] içine yazılır."""
    anahtar = ("murekkep", id(doc), pno, tuple(round(x, 1) for x in kutu))
    if anahtar in onbellek:
        return onbellek[anahtar]
    anahtar_d = ("dl", id(doc), pno)
    if anahtar_d not in onbellek:
        onbellek[anahtar_d] = doc[pno].get_displaylist()
    r = fitz.Rect(kutu) & doc[pno].rect
    sonuc = None
    if not r.is_empty:
        pix = onbellek[anahtar_d].get_pixmap(matrix=fitz.Matrix(olcek, olcek),
                                             colorspace=fitz.csGRAY, alpha=False,
                                             clip=r)
        a = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            pix.height, pix.width).copy()
        # Sütun ayırıcı çizgi mürekkep sayılmaz (kutuyu yana genişletmesin)
        orta = doc[pno].rect.width / 2
        for m_ in (fitz.Rect(orta - 4, r.y0, orta + 4, r.y1),):
            kes = m_ & r
            if kes.is_empty:
                continue
            a[max(0, int((kes.y0 - r.y0) * olcek)):int((kes.y1 - r.y0) * olcek) + 1,
              max(0, int((kes.x0 - r.x0) * olcek)):int((kes.x1 - r.x0) * olcek) + 1] = 255
        satir = np.where((a < 240).any(axis=1))[0]
        sutun = np.where((a < 240).any(axis=0))[0]
        if len(satir) and len(sutun):
            sonuc = fitz.Rect(r.x0 + sutun[0] / olcek, r.y0 + satir[0] / olcek,
                              r.x0 + (sutun[-1] + 1) / olcek,
                              r.y0 + (satir[-1] + 1) / olcek)
            onbellek[_murekkep_anahtari(doc, pno, sonuc)] = int((a < 240).sum())
    onbellek[anahtar] = sonuc
    return sonuc


def _sekil_kutulari(doc, bolgeler, onbellek):
    """Soru bölgelerindeki resim/çizim parçalarını birbirine değen ya da çok
    yakın olanları birleştirerek "şekil kutuları"na toplar. Karo karo bölünmüş
    resim ya da yüzlerce parçalı çizim tek kutu olur. Dönen: [(sayfa, Rect)]."""
    sonuc = []
    for pno, x0, y0, x1, y1 in bolgeler:
        parcalar = [r for r in _sayfa_nesneleri(doc, pno, onbellek)
                    if r.x0 >= x0 - 2 and r.x1 <= x1 + 2 and r.y0 >= y0 - 2
                    and r.y1 <= y1 + 2]
        kutular = []
        for r in parcalar:
            genis = fitz.Rect(r.x0 - 6, r.y0 - 6, r.x1 + 6, r.y1 + 6)
            birlesik = fitz.Rect(r)
            kalan = []
            for k_ in kutular:
                if k_.intersects(genis):
                    birlesik |= k_
                else:
                    kalan.append(k_)
            kalan.append(birlesik)
            kutular = kalan
        degisti = True                          # zincirleme birleşmeler
        while degisti:
            degisti = False
            for i in range(len(kutular)):
                for j in range(i + 1, len(kutular)):
                    a_, b_ = kutular[i], kutular[j]
                    if fitz.Rect(a_.x0 - 6, a_.y0 - 6, a_.x1 + 6, a_.y1 + 6).intersects(b_):
                        kutular[i] = a_ | b_
                        del kutular[j]
                        degisti = True
                        break
                if degisti:
                    break
        # Kırpılmış çizim/resimlerin görünmeyen parçaları da çizim kaydında yer
        # alır (kenar boşluğuna, sütun arasına taşar). Her kutu çizilip yalnızca
        # mürekkep olan alana daraltılır; tamamen boş kutu atılır.
        kutular = [_murekkep_alani(doc, pno, k_, onbellek) for k_ in kutular
                   if min(k_.width, k_.height) >= 3 and k_.width * k_.height >= 150]
        kutular = [k_ for k_ in kutular if k_ is not None
                   and min(k_.width, k_.height) >= 3 and k_.width * k_.height >= 150]
        # Sıra: satır satır (üst kenarları ±4 pt olanlar aynı satır), soldan sağa
        satirlar_ = []
        for k_ in sorted(kutular, key=lambda r: r.y0):
            if satirlar_ and abs(k_.y0 - satirlar_[-1][0]) <= 4:
                satirlar_[-1][1].append(k_)
            else:
                satirlar_.append([k_.y0, [k_]])
        sonuc.extend((pno, k_) for _y, sat in satirlar_
                     for k_ in sorted(sat, key=lambda r: r.x0))
    return sonuc


def _en_farkli_kare(a, b, kare=12):
    """En iyi hizada (en farklı karesi en az farklı olan kaydırmada) en farklı
    karenin farkı ve A ile B'deki sol üst pikseli: (fark, ya, xa, yb, xb)."""
    h, w = min(a.shape[0], b.shape[0]), min(a.shape[1], b.shape[1])
    en_iyi = None
    for dy in range(-2, 3):
        for dx in (-1, 0, 1):
            aa = a[max(0, dy):h, max(0, dx):w]
            bb = b[max(0, -dy):h, max(0, -dx):w]
            hh = min(aa.shape[0], bb.shape[0]) // kare * kare
            ww = min(aa.shape[1], bb.shape[1]) // kare * kare
            if hh < kare or ww < kare:
                continue
            d = np.abs(aa[:hh, :ww] - bb[:hh, :ww]).reshape(
                hh // kare, kare, ww // kare, kare).mean(axis=(1, 3))
            i, j = np.unravel_index(int(d.argmax()), d.shape)
            v = float(d[i, j])
            if en_iyi is None or v < en_iyi[0]:
                en_iyi = (v, max(0, dy) + i * kare, max(0, dx) + j * kare,
                          max(0, -dy) + i * kare, max(0, -dx) + j * kare)
    return en_iyi


def _yerel_fark(a, b, kare=12):
    """En iyi hizada, 12x12 piksellik karelerin en büyük ortalama gri farkı:
    küçük bir bölgedeki değişikliği (şekil, etiket) gösterir."""
    if a is None or b is None:
        return None
    sonuc = _en_farkli_kare(a, b, kare)
    return sonuc[0] if sonuc else None


def _fark_karesi(ia, ib, pa, ra, pb, rb, kare=12, olcek=0.8):
    """Rapor için: farklı bulunan şeklin en farklı küçük bölgesinin A ve B
    sayfasındaki dikdörtgeni (kıyastaki iki hizadan iyi olanına göre)."""
    adaylar = []
    s = _en_farkli_kare(ia, ib, kare)
    if s:
        adaylar.append((s, 0, 0, 0, 0))
    if ia.shape != ib.shape:
        h_, w_ = min(ia.shape[0], ib.shape[0]), min(ia.shape[1], ib.shape[1])
        s = _en_farkli_kare(ia[-h_:, -w_:], ib[-h_:, -w_:], kare)
        if s:
            adaylar.append((s, ia.shape[0] - h_, ia.shape[1] - w_,
                            ib.shape[0] - h_, ib.shape[1] - w_))
    if not adaylar:
        return []
    (_v, ya, xa, yb, xb), oya, oxa, oyb, oxb = min(adaylar, key=lambda t: t[0][0])
    boy = kare / olcek
    return [("A", (pa, ra.x0 + (xa + oxa) / olcek - 3, ra.y0 + (ya + oya) / olcek - 3,
                   ra.x0 + (xa + oxa) / olcek + boy + 3,
                   ra.y0 + (ya + oya) / olcek + boy + 3)),
            ("B", (pb, rb.x0 + (xb + oxb) / olcek - 3, rb.y0 + (yb + oyb) / olcek - 3,
                   rb.x0 + (xb + oxb) / olcek + boy + 3,
                   rb.y0 + (yb + oyb) / olcek + boy + 3))]


def _sekil_denetle(docA, docB, bolgelerA, bolgelerB, onbellek, durum, bildir):
    """A ve B'deki şekil kutularını sırasıyla eşleyip her birini kendi
    kutusunda görüntü olarak kıyaslar (konum ve satır aralığından bağımsız).
    Dönen: rapordaki "Şekil/Resim" hücresi."""
    ka = _sekil_kutulari(docA, bolgelerA, onbellek)
    kb = _sekil_kutulari(docB, bolgelerB, onbellek)
    if not ka and not kb:
        return "—"

    def boyut_uyar(ra, rb):
        return (abs(ra.width - rb.width) <= max(8, 0.12 * max(ra.width, rb.width))
                and abs(ra.height - rb.height) <= max(8, 0.12 * max(ra.height,
                                                                    rb.height)))
    if len(ka) != len(kb) or not all(boyut_uyar(ra, rb)
                                     for (_pa, ra), (_pb, rb) in zip(ka, kb)):
        # Aynı şekiller satır aralığı farkı yüzünden farklı birleşmiş olabilir:
        # sorudaki toplam şekil mürekkebi aynıysa fark yok sayılır; belirgin
        # farklıysa eksik/fazla şekil olabilir
        ta = sum(onbellek.get(_murekkep_anahtari(docA, p, r), 0) for p, r in ka)
        tb = sum(onbellek.get(_murekkep_anahtari(docB, p, r), 0) for p, r in kb)
        if abs(ta - tb) <= max(60, 0.06 * max(ta, tb)):
            return f"aynı ({len(ka)}↔{len(kb)} parça)"
        durum.append("ŞEKİL FARKI")
        oran = (tb - ta) / max(ta, 1) * 100
        bildir("UYARI", f"Şekil/resim içeriği farklı: B'de şekil mürekkebi A'ya göre "
               f"%{abs(oran):.0f} {'fazla' if oran > 0 else 'az'} (A'da {len(ka)}, "
               f"B'de {len(kb)} şekil parçası) — eksik ya da fazla şekil olabilir, "
               f"elle bakın.",
               [("A", (p, *r)) for p, r in ka] + [("B", (p, *r)) for p, r in kb])
        return f"FARKLI ({len(ka)}↔{len(kb)})"
    for i, ((pa, ra), (pb, rb)) in enumerate(zip(ka, kb), 1):
        ia = _kontrol_kupur(docA, (pa, ra.x0, ra.y0, ra.x1, ra.y1), onbellek,
                            ("sekilA", pa, tuple(ra)))
        ib = _kontrol_kupur(docB, (pb, rb.x0, rb.y0, rb.x1, rb.y1), onbellek,
                            ("sekilB", pb, tuple(rb)))
        # Ölçü: en iyi hizada küçük bölgelerdeki gri ton farkı. (Genel örtüşme
        # puanı küçük/boş kutularda ve desenli dolgularda yanıltıcı olduğu için
        # kullanılmaz; gerçek kitapçıklardaki 222 aynı şekilde en yüksek 20 çıktı.)
        yerel = _yerel_fark(ia, ib)
        if (yerel is not None and ia is not None and ib is not None
                and ia.shape != ib.shape):
            # Boyutlar biraz farklıysa fazlalık hangi kenardaysa ortak alan oraya
            # göre hizalanır: üst-sol ve alt-sağ hizadan iyi olanı alınır
            h_, w_ = min(ia.shape[0], ib.shape[0]), min(ia.shape[1], ib.shape[1])
            alt = _yerel_fark(ia[-h_:, -w_:], ib[-h_:, -w_:])
            if alt is not None:
                yerel = min(yerel, alt)
        if yerel is not None and yerel > KONTROL_SEKIL_ESIK:
            durum.append("ŞEKİL FARKI")
            bildir("HATA", f"{i}. şekil/resim farklı görünüyor (en farklı bölge "
                   f"{yerel:.0f}/255) — elle bakın.",
                   [("A", (pa, *ra)), ("B", (pb, *rb))]
                   + _fark_karesi(ia, ib, pa, ra, pb, rb))
            return "FARKLI"
    return f"aynı ({len(ka)} şekil)"


def _nesne_bolgeleri(satir_listesi, sinif, genislikler, tam_sayfa=None):
    """Satırların bulunduğu her sayfa/sütun için (sayfa, x0, y0, x1, y1).
    tam_sayfa: bu sayfadaki satırlar tam genişlik sayılır."""
    gruplar = {}
    for r in satir_listesi:
        pno = r["sayfa"]
        orta = genislikler[pno] / 2
        if pno == tam_sayfa or sinif.get(id(r)) == "F":
            yan = "F"
        else:
            yan = "L" if (r["x0"] + r["x1"]) / 2 < orta else "R"
        gruplar.setdefault((pno, yan), []).append(r)
    sonuc = []
    for (pno, yan), grup in gruplar.items():
        W = genislikler[pno]
        x0, x1 = {"F": (0, W), "L": (0, W / 2), "R": (W / 2, W)}[yan]
        sonuc.append((pno, x0, min(r["y0"] for r in grup) - 3, x1,
                      max(r["y1"] for r in grup) + 3))
    return sonuc


def _test_adlari(k, yapi):
    """Rapor için test adları: testin ilk sorusunun sayfasındaki üst başlık
    kutusundan ("SÖZEL BÖLÜM - TÜRKÇE" → "TÜRKÇE"); bulunamazsa yapıdaki ad."""
    adlar = []
    for ti, t in enumerate(yapi):
        ad = t["test"]
        ilk = min((o for o in k["ogeler"] if o["ti"] == ti and o["no"]),
                  key=lambda o: o["no"], default=None)
        if ilk is not None:
            bulunan = _sayfa_basligi(k["satirlar"], ilk["sayfa"],
                                     k["yukseklikler"][ilk["sayfa"]], set(adlar))
            if bulunan and ("BÖLÜM" in " ".join(
                    r["metin"] for r in k["satirlar"] if r["sayfa"] == ilk["sayfa"]
                    and r["y1"] < 0.2 * k["yukseklikler"][ilk["sayfa"]])):
                ad = bulunan       # TYT/AYT'de yapıdaki ad (Türkçe vb.) kalır
        adlar.append(ad)
    return adlar


def _sik_harfleri(o):
    """Sorudaki şık harfleri; yan yana dizilmiş şıklar (A) 12  B) 15 …) dahil."""
    harfler = set()
    for r in o["ham_satirlar"]:
        h = _sik_harfi(r)
        if h:
            harfler.add(h)
        harfler.update(re.findall(r"(?:^|\s)([A-E])\s*\)", r["metin"]))
    return harfler


def _cevap_adaylari(o, genislikler=None):
    """Bir sorunun cevap işareti adayları (EŞLEŞTİR ve A–B KONTROL ortak kural).
    Aday: soru alanındaki renkli tek harf (A–E); ayrıca rengi ne olursa olsun
    (siyaha çevrilmiş, yanlış stil almış) tek başına duran ve soru numarasıyla
    sol hizada olan tek harf ("C", "C.", "C)") — renksiz aday. "Oluk" adayı: harf,
    yanındaki metnin (soru kökü ya da şıklar) girintisinin solundaki boşlukta
    durur — numaranın altı ya da şıkların solu; denemede cevap hep bu boşluktadır.
    Şeklin içindeki renkli etiketler ve girintiye dayalı tablo harfleri oluk adayı
    olmaz. Ölçü her işaretin kendi yerinde alınır: yanındaki metin (aynı sütun,
    ±30 pt) işaretin sağından başlamalı; aynı sayfa yarısında (sütunda) ya da
    işaretin üstünden geçen, işaretin sağ kenarından önce başlayan metin satırı
    olmamalı (yan yana kısa şıklar dahil). Böylece şıklar başka sütuna/sayfaya
    taşsa da çalışır. Dönen: [{harf, satir, renk, kutu, oluk, renkli}]."""
    satirlar_ = o["ham_satirlar"]
    if not satirlar_:
        return []
    bas = satirlar_[0]                       # soru numarası satırı
    bas_kutu = (bas.get("kutular") or [(bas["x0"], bas["y0"], bas["x1"], bas["y1"])])[0]
    metin_bas, adaylar = [], []
    for r in satirlar_:
        spanlar = r.get("spanlar", [])
        kutular = r.get("kutular") or [(r["x0"], r["y0"], r["x1"], r["y1"])] * len(spanlar)
        dolu = [(t.strip(), renk, tuple(kutu)) for (t, renk), kutu in zip(spanlar, kutular)
                if t.strip()]
        ilk_metin = None
        for i, (tt, renk, kutu) in enumerate(dolu):
            sonraki = dolu[i + 1][2][0] if i + 1 < len(dolu) else None
            tek_basina = sonraki is None or sonraki - kutu[2] >= 6
            if renkli_mi(renk):
                if re.fullmatch(r"[A-E]", tt):
                    adaylar.append({"harf": tt, "satir": r, "renk": renk,
                                    "kutu": kutu, "renkli": True})
                continue
            if (re.fullmatch(r"[A-E][.)]?", tt) and tek_basina and r is not bas
                    and r["sayfa"] == bas["sayfa"]
                    and -6 <= kutu[0] - bas_kutu[0] <= 14 and kutu[1] > bas["y0"]):
                adaylar.append({"harf": tt[0], "satir": r, "renk": renk,
                                "kutu": kutu, "renkli": False})
                continue                     # renksiz aday girinti sayılmaz
            if ilk_metin is None and not re.fullmatch(r"\d{1,3}\s*[.)]?", tt):
                ilk_metin = kutu[0]          # soru numarası girinti sayılmaz
        if ilk_metin is not None and len(r["metin"].strip()) >= 2:
            metin_bas.append((r["sayfa"], ilk_metin, r["x1"], r["y0"], r["y1"]))
    for a in adaylar:
        x0, y0, x1, y1 = a["kutu"]
        pno = a["satir"]["sayfa"]
        orta = (genislikler[pno] if genislikler else 612) / 2
        komsu = [(mx0, mx1) for p, mx0, mx1, my0, my1 in metin_bas
                 if p == pno and my1 > y0 - 30 and my0 < y1 + 30]
        sagda = [mx0 for mx0, _mx1 in komsu if x1 - 1 <= mx0 < x1 + 80]
        solda = [mx0 for mx0, mx1 in komsu if mx0 < x1 - 1
                 and (mx1 > x0 or (mx0 >= orta) == (x0 >= orta))]
        a["oluk"] = bool(sagda) and not solda
    return adaylar


def cevap_motoru(k):
    """Kitapçıktaki her sorunun cevap harfi ve teşhisi (EŞLEŞTİR ile A–B KONTROL
    aynı motoru kullanır). Sorular sayfa düzeninden ayrılır (tam genişlik, iki
    sütun, sütun/sayfa devri); cevap sorunun oluğundaki renkli harftir. Kitapçığın
    cevap rengi (çoğunlukla magenta) kendiliğinden öğrenilir. Kendini denetler:
    cevap yoksa, birden fazla işaret varsa, harf beklenen yerde değilse ya da
    sorunun şıklarında yoksa teşhis yazılır.
    Dönen: (cevaplar {(ti, no): harf|None}, teşhisler [(ti, no, sayfa, mesaj)],
            kitapçık cevaplı mı)."""
    from collections import Counter
    sorular = [(o, _cevap_adaylari(o, k.get("genislikler"))) for o in k["ogeler"]
               if o["no"]]
    oluklu = [ad for _o, ad in sorular if any(a["oluk"] and a["renkli"] for a in ad)]
    if not sorular or len(oluklu) < 0.5 * len(sorular):
        return {}, [], False                 # baskı kitapçığı: cevap işareti yok
    baskin = Counter(a["renk"] for ad in oluklu for a in ad
                     if a["oluk"] and a["renkli"]).most_common(1)[0][0]
    cevaplar, teshis = {}, []
    for o, ad in sorular:
        anahtar = (o["ti"], o["no"])
        renksiz = sorted({a["harf"] for a in ad if a["oluk"] and not a["renkli"]})
        ad = [a for a in ad if a["renkli"]]
        oluk = [a for a in ad if a["oluk"]]
        oluk = [a for a in oluk if a["renk"] == baskin] or oluk
        harfler = sorted({a["harf"] for a in oluk})
        harf = None
        if len(harfler) == 1:
            harf = harfler[0]
        elif len(harfler) > 1:
            # Birden fazla işaret: A şıkkı satırına en yakın olanı al, uyar
            sik_a = [r for r in o["ham_satirlar"] if _sik_harfi(r) == "A"]
            def uzaklik(a):
                if not sik_a:
                    return 0
                return min(abs(a["kutu"][1] - r["y0"]) + 1000 * (a["satir"]["sayfa"] != r["sayfa"])
                           for r in sik_a)
            harf = min(oluk, key=uzaklik)["harf"]
            teshis.append((o["ti"], o["no"], o["sayfa"],
                           f"birden fazla cevap işareti ({', '.join(harfler)}); "
                           f"A şıkkına en yakın olan '{harf}' alındı"))
        elif len(renksiz) == 1:
            harf = renksiz[0]
            teshis.append((o["ti"], o["no"], o["sayfa"],
                           f"cevap harfi '{harf}' cevap renginde değil (siyah ya da "
                           f"başka renk basılmış); rengi elle değiştirilmiş olabilir"))
        else:
            diger = sorted({a["harf"] for a in ad if a["renk"] == baskin})
            if len(diger) == 1:
                harf = diger[0]
                teshis.append((o["ti"], o["no"], o["sayfa"],
                               f"cevap harfi '{harf}' beklenen yerde (numaranın altı / "
                               f"şıkların solu) değil"))
            else:
                teshis.append((o["ti"], o["no"], o["sayfa"],
                               "cevap işareti okunamadı"
                               + (f" (adaylar: {', '.join(diger)})" if diger else "")))
        siklar = _sik_harfleri(o)
        if harf and "A" in siklar and len(siklar) >= 3 and harf not in siklar:
            teshis.append((o["ti"], o["no"], o["sayfa"],
                           f"cevap '{harf}' ama soruda {''.join(sorted(siklar))} "
                           f"şıkları var"))
        cevaplar[anahtar] = harf
    return cevaplar, teshis, True


def _oluk_piksel_izleri(k, doc, dpi=110):
    """Metin katmanında görünmeyen cevap izleri (eğriye çevrilmiş ya da resme
    gömülü cevap harfi). Her sorunun numara hizasındaki kök–şık arası şeridi
    (numaranın solundan 4 pt, sağından 16 pt; numaranın alt kenarından ilk şıkka
    kadar)
    görüntü olarak taranır. Metin olmayan, harf boyutunda ve tek başına duran
    (sağında, üstünde, altında mürekkep sürmeyen) her iz döner — gerçek
    unutulmuş işaret tek başına durur; şerit kenarına dayanmış şekil, tablo,
    çizgi sayılmaz. Dönen: [(öğe, sayfa, Rect, cevap renginde mi)]."""
    from scipy import ndimage
    olcek = 72.0 / dpi
    metin = {}
    for r in k["satirlar"]:
        for kutu in r.get("kutular") or [(r["x0"], r["y0"], r["x1"], r["y1"])]:
            metin.setdefault(r["sayfa"], []).append(fitz.Rect(kutu))
    bulunan = []
    for o in k["ogeler"]:
        if not o["no"] or not o["ham_satirlar"]:
            continue
        bas = o["ham_satirlar"][0]
        pno = bas["sayfa"]
        sx = (bas.get("kutular") or [(bas["x0"],)])[0][0]
        W = k["genislikler"][pno]
        sol = sx < W / 2
        siklar = [r["y0"] for r in o["ham_satirlar"]
                  if r["sayfa"] == pno and _sik_harfi(r) and (r["x0"] < W / 2) == sol
                  and r["y0"] > bas["y0"] + 4]
        if not siklar:
            continue
        y_ust, y_alt = bas["y1"] + 1, min(siklar) - 2      # numaranın altından
        if y_alt - y_ust < 8:
            continue
        serit = fitz.Rect(sx - 4, y_ust, sx + 16, y_alt)
        sayfa = doc[pno]
        kirp = fitz.Rect(serit.x0, serit.y0 - 4, serit.x1 + 8, serit.y1 + 4) & sayfa.rect
        if kirp.is_empty:
            continue
        pix = sayfa.get_pixmap(dpi=dpi, clip=kirp, colorspace=fitz.csRGB,
                               alpha=False, annots=False)
        a = np.frombuffer(pix.samples, np.uint8).reshape(
            pix.height, pix.width, 3).astype(np.int16)
        ink = (a <= 200).any(axis=2)
        if not ink.any():
            continue
        etiket, _n = ndimage.label(ndimage.binary_dilation(ink),
                                   structure=np.ones((3, 3)))
        for sl in ndimage.find_objects(etiket):
            if sl is None:
                continue
            py0, py1, px0, px1 = sl[0].start, sl[0].stop, sl[1].start, sl[1].stop
            r = fitz.Rect(kirp.x0 + px0 * olcek, kirp.y0 + py0 * olcek,
                          kirp.x0 + px1 * olcek, kirp.y0 + py1 * olcek)
            if not r.intersects(serit) or r.x1 > serit.x1 + 1:
                continue
            if not (3 <= r.width <= 22 and 3 <= r.height <= 22) or \
                    max(r.width, r.height) < 7 or r.width >= 19:
                continue
            if r.y0 <= bas["y1"] + 0.5:          # numara glifinin eteği
                continue
            sag = int(6 / olcek) + 1
            dik = int(3 / olcek) + 1
            if (ink[py0:py1, px1:px1 + sag].any() or
                    ink[max(py0 - dik, 0):py0, px0:px1].any() or
                    ink[py1:py1 + dik, px0:px1].any()):
                continue                         # tek başına değil: içerik parçası
            if any(m.intersects(r) for m in metin.get(pno, [])):
                continue                         # metin katmanında: ayrıca denetlenir
            parca = a[py0:py1, px0:px1][ink[py0:py1, px0:px1]]
            ort = parca.mean(axis=0) if len(parca) else (0, 0, 0)
            cevap_rengi = ort[0] > 150 and ort[1] < 110 and ort[2] > 90
            bulunan.append((o, pno, r, bool(cevap_rengi)))
    return bulunan


def _magenta_mi(renk):
    r, g, b = (renk >> 16) & 255, (renk >> 8) & 255, renk & 255
    return r > 170 and g < 110 and b > 90


def _acik_cevaplar(k):
    """Soru numarasının altındaki boşlukta (şıkların solunda) renkli A–E harfi,
    her yerde magenta A–E harfi ya da renkli soru kodu: açık kalmış cevap/kod.
    Şekil içindeki renkli etiketler (haritadaki A, B, C ...) sayılmaz."""
    bulunan = []
    for o in k["ogeler"]:
        if not o["ham_satirlar"]:
            continue
        for a in _cevap_adaylari(o, k.get("genislikler")):   # EŞLEŞTİR ile aynı kural
            if a["renkli"] and (a["oluk"] or _magenta_mi(a["renk"])):
                bulunan.append((o, a["satir"], a["harf"], f"#{a['renk']:06X}"))
            elif not a["renkli"] and a["oluk"]:
                bulunan.append((o, a["satir"], a["harf"],
                                f"#{a['renk']:06X}; siyah/renksiz basılmış — cevap "
                                f"rengi elle değiştirilmiş olabilir"))
        for r in o["ham_satirlar"]:
            for t, renk in r.get("spanlar", []):
                tt = t.strip()
                if not tt or not renkli_mi(renk) or re.fullmatch(r"[A-E]", tt):
                    continue
                kod = tt.replace(" ", "")
                if not (re.fullmatch(r"[A-Za-z0-9ÇĞİÖŞÜçğıöşü.\-_/]{4,20}", tt)
                        and re.search(r"\d", kod) and re.search(r"[A-Za-z]", kod)):
                    continue     # soru kodu boşluksuz tek parçadır ("81 ilde" değil)
                bulunan.append((o, r, tt, f"#{renk:06X}"))
    return bulunan


def _isaret_onbellegi(A, B, docA, docB, onbellek):
    """Görüntüde beyazlatılacak magenta cevap/kod satırları: motorun okuduğu
    satırlardan (sayfa metnini yeniden okumamak için)."""
    for kk, doc in ((A, docA), (B, docB)):
        for pno in range(len(doc)):
            onbellek[("isaret", id(doc), pno)] = []
            onbellek[("metinler", id(doc), pno)] = []
        for r in kk["satirlar"]:
            onbellek[("metinler", id(doc), r["sayfa"])].append(
                fitz.Rect(r["x0"], r["y0"], r["x1"], r["y1"]))
            dolu = [(t, rk) for t, rk in r.get("spanlar", []) if t.strip()]
            if dolu and all(_renkli_isaret_mi(t, rk) for t, rk in dolu):
                onbellek[("isaret", id(doc), r["sayfa"])].append(
                    fitz.Rect(r["x0"], r["y0"], r["x1"], r["y1"]))


def _icerik_eslestir(A, B, docA, docB, onbellek, cevap_a, cevap_b):
    """Her testin içinde (dersler arası) A ve B sorularını içerikle birebir
    eşler: tekil soru kodu, gövde + şık metni (kısa metinde görüntü), cevap.
    Dönen: ([(oa, ob, metin oranı)], [(ti, oa, en yakın ob)] A'da kalan,
    [(ti, ob)] B'de kalan)."""
    eslesmeler, eksik_a, eksik_b = [], [], []
    testler = sorted({o["ti"] for o in A["ogeler"]} | {o["ti"] for o in B["ogeler"]})
    for ti in testler:
        la = [o for o in A["ogeler"] if o["ti"] == ti]
        lb = [o for o in B["ogeler"] if o["ti"] == ti]
        if not la or not lb:
            continue
        kod_a = {}
        for o in la:
            if o["kod"]:
                kod_a.setdefault(o["kod"], []).append(o)
        kod_b = {}
        for o in lb:
            if o["kod"]:
                kod_b.setdefault(o["kod"], []).append(o)
        M = np.zeros((len(la), len(lb)))
        for i, oa in enumerate(la):
            for j, ob in enumerate(lb):
                if len(oa["norm"]) < 40 or len(ob["norm"]) < 40:
                    p = _gorsel_puan(docA, docB, oa["bolge"], ob["bolge"],
                                     onbellek)
                    s_ = p if p is not None else 0.0
                else:
                    u = oa["tri"] | ob["tri"]
                    s_ = len(oa["tri"] & ob["tri"]) / len(u) if u else 0.0
                if (oa["kod"] and oa["kod"] == ob["kod"]
                        and len(kod_a[oa["kod"]]) == 1
                        and len(kod_b[ob["kod"]]) == 1):
                    s_ += 1.0                       # tekil soru kodu: kesin eş
                ca_ = cevap_a.get((ti, oa["no"])) or oa["cevap"]
                cb_ = cevap_b.get((ti, ob["no"])) or ob["cevap"]
                if ca_ and cb_:
                    s_ += 0.03 if ca_ == cb_ else -0.03
                M[i, j] = s_
        sat, sut = linear_sum_assignment(-M)
        alinan_a, alinan_b = set(), set()
        for i, j in zip(sat, sut):
            oa, ob = la[i], lb[j]
            oran = _benzerlik_orani(oa["cmp"], ob["cmp"])
            if oran < KONTROL_METIN_YOK and max(len(oa["norm"]), len(ob["norm"])) >= 40:
                continue                     # gerçek eş değil; aşağıda raporlanır
            alinan_a.add(i)
            alinan_b.add(j)
            eslesmeler.append((oa, ob, oran))
        for i, oa in enumerate(la):
            if i not in alinan_a:
                eksik_a.append((ti, oa, max(lb, key=lambda ob: len(oa["tri"] & ob["tri"])
                                            / (len(oa["tri"] | ob["tri"]) or 1))))
        eksik_b.extend((ti, ob) for j, ob in enumerate(lb) if j not in alinan_b)
    return eslesmeler, eksik_a, eksik_b


def _kitapciklari_hazirla(a_pdf, b_pdf, yapi, hazir_a, hazir_b):
    """A ve B'yi A–B KONTROL'ün soru ayırma motoruyla böler (okunmuş satırlarla)."""
    sessiz = lambda _m: None
    A = _kitapcik_oku(a_pdf, yapi, "A", sessiz, hazir=hazir_a)
    B = _kitapcik_oku(b_pdf, yapi, "B", sessiz, hazir=hazir_b)
    mobilya = _mobilya_kumesi([A, B])
    _kitapcik_bolumle(A, mobilya, sessiz)
    _kitapcik_bolumle(B, mobilya, sessiz)
    return A, B


def cevaplari_belirle(yapilar, a_s, b_s, yapi, log):
    """EŞLEŞTİR'in cevap sütunu ortak cevap motorundan gelir. Cevaplı kitapçıkta
    motorun bulamadığı ya da şüphelendiği her soru uyarı olarak yazılır (sessiz
    yanlış yerine görünür uyarı). Cevapsız (baskı) kitapçıkta şekildeki renkli
    harflerden cevap üretilmez. Dönen: uyarı satırları."""
    uyarilar = []
    for kk, sorular in zip(yapilar, (a_s, b_s)):
        cevaplar, teshis, cevapli = cevap_motoru(kk)
        for anahtar, s_ in sorular.items():
            s_["cevap"] = cevaplar.get(anahtar) if cevapli else None
        for o in kk["ogeler"]:              # içerik doğrulaması da aynı cevabı görsün
            o["cevap"] = cevaplar.get((o["ti"], o["no"])) if cevapli and o["no"] else None
        if not cevapli:
            continue
        okunan = sum(1 for k_ in sorular if cevaplar.get(k_))
        log(f"  {kk['etiket']} kitapçığı: cevap harfi {okunan}/{len(sorular)} soruda "
            f"okundu" + (f", {len(teshis)} soruda uyarı" if teshis else ""))
        for ti, no, pno, mesaj in teshis:
            test = yapi[ti]["test"] if ti < len(yapi) else ""
            uyarilar.append(f"CEVAP KONTROL ET: {kk['etiket']} kitapçığı {test} "
                            f"{no}. soru (s.{pno + 1}) — {mesaj}.")
    return uyarilar


def eslestirme_dogrula(a_pdf, b_pdf, yapi, esl, uyarilar, log, hazir_a=None,
                       hazir_b=None, hazir_yapi=None):
    """EŞLEŞTİR sonucunu bağımsız içerik karşılaştırmasıyla doğrular ve yalnızca
    kesin yanlış eşleri düzeltir: A sorusunun metni B'deki başka bir soruyla
    birebir (≥ %98) aynıyken EŞLEŞTİR'in seçtiği sorunun metni belirgin farklıysa
    (< %90). Her düzeltme uyarılara "DÜZELTİLDİ" diye yazılır. Karışık sayfa
    düzeninde (tam genişlik soru + iki sütun) EŞLEŞTİR'in metni karışabiliyor."""
    from collections import Counter
    if hazir_yapi is not None:          # cevap motoru için zaten bölünmüş
        A, B = hazir_yapi
    else:
        A, B = _kitapciklari_hazirla(a_pdf, b_pdf, yapi, hazir_a, hazir_b)
    docA, docB = fitz.open(a_pdf), fitz.open(b_pdf)
    onbellek = {}
    _isaret_onbellegi(A, B, docA, docB, onbellek)
    ciftler, _ea, _eb = _icerik_eslestir(A, B, docA, docB, onbellek, {}, {})
    docA.close()
    docB.close()
    b_soru = {(o["ti"], o["no"]): o for o in B["ogeler"] if o["no"]}
    adaylar = {}
    for oa, ob, oran in ciftler:
        if not oa["no"] or not ob["no"] or oran < 0.98:
            continue
        k = (oa["ti"], oa["no"])
        mevcut = esl.get(k)
        if mevcut is None:
            continue
        if mevcut.get("b_no") == ob["no"]:
            # Eş doğru ve metin birebir aynı: ilk eşleştirmenin (karışık düzende
            # bozulan metinden çıkan) düşük benzerlik uyarısı yersizdir
            if mevcut.get("benzerlik", 1) < oran:
                test = yapi[k[0]]["test"]
                uyarilar[:] = [u for u in uyarilar
                               if not (u.startswith("DÜŞÜK BENZERLİK") and test in u
                                       and f"A-{k[1]} → B-{ob['no']} " in u)]
                mevcut["benzerlik"] = oran
            continue
        eski_ob = b_soru.get((oa["ti"], mevcut.get("b_no")))
        eski_oran = _benzerlik_orani(oa["cmp"], eski_ob["cmp"]) if eski_ob else 0.0
        if eski_oran >= 0.90:
            continue                 # EŞLEŞTİR'in eşi de çok benziyor: dokunma
        adaylar[k] = (mevcut.get("b_no"), ob["no"], oran, eski_oran)
    yeni = {k: v.get("b_no") for k, v in esl.items()}
    for k, (_e, y, _o, _eo) in adaylar.items():
        yeni[k] = y
    say = Counter((k[0], b) for k, b in yeni.items())
    duzeltilen = 0
    for k, (eski, y, oran, eski_oran) in sorted(adaylar.items()):
        if say[(k[0], y)] > 1:
            continue                 # B numarası çakışıyor: dokunma
        esl[k] = {"b_no": y, "benzerlik": oran, "yontem": "İÇERİK"}
        test = yapi[k[0]]["test"]
        uyarilar[:] = [u for u in uyarilar
                       if not (u.startswith("DÜŞÜK BENZERLİK") and test in u
                               and f"A-{k[1]} → B-{eski} " in u)]
        mesaj = (f"DÜZELTİLDİ: {test} A-{k[1]} → B-{y} (ilk eşleştirme B-{eski} "
                 f"demişti; B-{y} ile metin %{oran * 100:.0f} aynı, B-{eski} ile "
                 f"%{eski_oran * 100:.0f}).")
        uyarilar.append(mesaj)
        log("  ✎ " + mesaj)
        duzeltilen += 1
    log(f"İçerik doğrulaması: {len(esl) - duzeltilen}/{len(esl)} eşleşme "
        f"doğrulandı" + (f", {duzeltilen} düzeltildi." if duzeltilen else "."))


def _fark_ozeti(a, b, azami=3):
    """İki metin arasındaki ilk farkları kelime düzeyinde özetler."""
    ka, kb = a.split(), b.split()
    sm = difflib.SequenceMatcher(None, ka, kb, autojunk=False)
    parcalar = []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        sa = " ".join(ka[i1:i2])[:60] or "—"
        sb = " ".join(kb[j1:j2])[:60] or "—"
        parcalar.append(f"A: «{sa}» → B: «{sb}»")
        if len(parcalar) >= azami:
            break
    return "; ".join(parcalar)


def _yuzde(oran):
    """Benzerlik yüzdesi; farklı metin hiçbir zaman "%100" görünmesin."""
    y = oran * 100
    return f"%{y:.0f}" if y < 99.5 else f"%{min(y, 99.9):.1f}".replace(".", ",")


def _benzerlik_orani(a, b):
    if not a and not b:
        return 1.0
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


def _bolge_kelimeleri(doc, bolgeler):
    """Bölgelerdeki kelimeler görünüş sırasıyla: [(metin, [(sayfa, Rect)])].
    Renkli cevap harfi/kodu ve baştaki soru numarası alınmaz, satır sonunda
    tirelenmiş kelime birleştirilir."""
    kelimeler = []
    for pno, x0, y0, x1, y1 in bolgeler:
        sayfa = doc[pno]
        r = fitz.Rect(x0, y0, x1, y1) & sayfa.rect
        if r.is_empty:
            continue
        satirlar_ = []
        for blk in sayfa.get_text("rawdict", clip=r).get("blocks", []):
            for ln in blk.get("lines", []):
                kel, cur, rr = [], "", None
                for sp in ln.get("spans", []):
                    harfler = sp.get("chars", [])
                    if _renkli_isaret_mi("".join(c["c"] for c in harfler),
                                         sp.get("color", 0)):
                        continue
                    for c in harfler:
                        if c["c"].isspace():
                            if cur:
                                kel.append((cur, rr))
                            cur, rr = "", None
                            continue
                        cr = fitz.Rect(c["bbox"])
                        cur += c["c"]
                        rr = cr if rr is None else rr | cr
                if cur:
                    kel.append((cur, rr))
                if kel:
                    bb = ln["bbox"]
                    satirlar_.append(((bb[1] + bb[3]) / 2, bb[0], kel))
        # Satırlar yukarıdan aşağı (orta noktaları 3 pt içinde olanlar aynı
        # satır, soldan sağa): A ile B'de aynı düzen aynı sırayı verir
        satirlar_.sort(key=lambda t: (t[0], t[1]))
        sira_, grup = [], []
        for t in satirlar_:
            if grup and t[0] - grup[0][0] > 3:
                sira_.append(grup)
                grup = []
            grup.append(t)
        if grup:
            sira_.append(grup)
        for grup in sira_:
            satir_kel = [(m, [(pno, rc)]) for _y, _x, kel in
                         sorted(grup, key=lambda t: t[1]) for m, rc in kel]
            if (kelimeler and satir_kel and kelimeler[-1][0][-1:] in TIRELER
                    and len(kelimeler[-1][0]) > 1 and satir_kel[0][0][:1].islower()):
                onceki_m, onceki_r = kelimeler.pop()
                ilk_m, ilk_r = satir_kel.pop(0)
                kelimeler.append((onceki_m[:-1] + ilk_m, onceki_r + ilk_r))
            kelimeler.extend(satir_kel)
    if kelimeler and re.fullmatch(r"\d{1,3}[.)]?", kelimeler[0][0]):
        kelimeler = kelimeler[1:]
    return kelimeler


def _fark_kelimeleri(docs, gorunum):
    """A ile B arasında farklı olan kelimelerin yerleri (raporda sarıyla
    gösterilir): [(kitapçık, sayfa, (x0, y0, x1, y1))]. Metinler büyük ölçüde
    farklıysa (yanlış soru, başka düzen) boş döner."""
    ka = _bolge_kelimeleri(docs["A"], [b for et, b in gorunum if et == "A"])
    kb = _bolge_kelimeleri(docs["B"], [b for et, b in gorunum if et == "B"])
    if not ka or not kb:
        return []
    anahtar = lambda m: re.sub(r"-", "", _kiyas_metni(m))
    sm = difflib.SequenceMatcher(None, [anahtar(m) for m, _r in ka],
                                 [anahtar(m) for m, _r in kb], autojunk=False)
    sonuc, farkli = [], 0
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        farkli += (i2 - i1) + (j2 - j1)
        for et, liste in (("A", ka[i1:i2]), ("B", kb[j1:j2])):
            for _m, yerler in liste:
                sonuc.extend((et, pno, tuple(rc)) for pno, rc in yerler)
    if farkli > max(8, 0.3 * (len(ka) + len(kb))):
        return []
    return sonuc


def _sorunlari_sirala(sorunlar):
    """Önce hatalar, sonra uyarılar; her biri kitapçıktaki yerine göre (A'nın
    sayfa sırası, sonra B'ninki). Her soruna rapordaki sıra numarası verilir."""
    def anahtar(s):
        yerler = list(s["yerler"]) + list(s["gorunum"] or [])
        for i, et in enumerate(("A", "B")):
            b = next((b for e, b in yerler if e == et), None)
            if b is not None:
                return ({"HATA": 0}.get(s["onem"], 1), i, b[0], b[2], b[1])
        return ({"HATA": 0}.get(s["onem"], 1), 2, 0, 0, 0)
    sorunlar.sort(key=anahtar)
    for i, s in enumerate(sorunlar, 1):
        s["sira"] = i
    return sorunlar


def _sorun_yeri(s):
    """'A s.3, B s.5' biçiminde sorunun bulunduğu sayfalar."""
    sayfalar = {}
    for et, b in list(s["yerler"]) or list(s.get("gorunum") or []):
        sayfalar.setdefault(et, set()).add(b[0] + 1)
    return ", ".join(f"{et} s.{'/'.join(map(str, sorted(p)))}"
                     for et, p in sayfalar.items())


def kontrol_et(a_pdf, b_pdf, yapi, log, cevap_a=None, cevap_b=None,
               yapi_b=None):
    """A ve B kitapçığını denetler. Dönen: (sonuc sözlüğü, sorunlar listesi)."""
    cevap_a, cevap_b = cevap_a or {}, cevap_b or {}
    okunurluk_denetle(a_pdf, b_pdf)
    sorunlar = []

    def sorun(onem, kitapcik, aciklama, test="", ders="", soru="", yerler=(),
              gorunum=None, kutular=(), vurgula=False):
        # yerler: işaretlenecek bölgeler; gorunum: raporda gösterilecek bölgeler
        # (yoksa yerlerdeki sorunun tamamı); kutular: ayrıca çerçevelenecek yerler
        sorunlar.append({"onem": onem, "kitapcik": kitapcik, "test": test,
                         "ders": ders, "soru": soru, "aciklama": aciklama,
                         "yerler": list(yerler), "gorunum": gorunum,
                         "kutular": list(kutular), "vurgula": vurgula, "vurgu": []})

    log("A kitapçığı okunuyor...")
    A = _kitapcik_oku(a_pdf, yapi, "A", log)
    log("B kitapçığı okunuyor...")
    B = _kitapcik_oku(b_pdf, yapi, "B", log)
    for kk in (A, B):
        if not kk["sorular"]:
            raise RuntimeError(
                f"'{Path(kk['pdf']).name}' içinde hiç soru bulunamadı. Soru "
                f"kitapçığı PDF'ini seçtiğinizden emin olun.")
    mobilya = _mobilya_kumesi([A, B])
    log("Sayfa düzeni çözümleniyor...")
    _kitapcik_bolumle(A, mobilya, log)
    _kitapcik_bolumle(B, mobilya, log)
    # Test adları sayfa başlığından (yapıdaki ad yalnızca yedek); tek dersli
    # testlerde ders adı da aynı adla gösterilir
    adlar = _test_adlari(A, yapi)
    for kk in (A, B):
        for o in kk["ogeler"]:
            if o["ti"] < len(yapi) and o["ders"] == yapi[o["ti"]]["test"]:
                o["ders"] = adlar[o["ti"]]
    test_adi = lambda ti: adlar[ti] if ti is not None and ti < len(adlar) else ""
    tam_bolge = {}                  # (kitapçık, ilk bölge) → sorunun tüm parçaları
    for kk in (A, B):
        for o in kk["ogeler"]:
            tam_bolge[(kk["etiket"], o["bolge"])] = o["nesne_bolgeleri"]
        for b in kk["basliklar"]:
            tam_bolge.setdefault((kk["etiket"], b["bolge"]), b["nesne_bolgeleri"])
    yer = lambda kk, o: (kk["etiket"], o["bolge"])

    # --- 0) Açık kalan cevap (baskıda cevap/kod görünmemeli) ---------------------
    # Cevap kıyası da yalnızca bu yolla bulunan gerçek cevap işaretlerini kullanır
    # (şekildeki renkli harf etiketleri cevap sayılmaz).
    acik_say = {}
    for kk in (A, B):
        # Cevap yerinde tek başına duran renksiz harf (rengi siyaha çevrilmiş
        # cevap) soru metninden çıkarılır: açık cevap olarak ayrıca bildirilir,
        # "metin farklı" diye bir kez daha yazılmaz
        for o in kk["ogeler"]:
            cikar = {id(a["satir"]) for a in _cevap_adaylari(o, kk.get("genislikler"))
                     if not a["renkli"] and a["oluk"]
                     and a["satir"]["metin"].strip().rstrip(".)") == a["harf"]}
            if cikar:
                parcalar = [_temiz_satir(r, ilk=(i == 0))
                            for i, r in enumerate(o["ham_satirlar"]) if id(r) not in cikar]
                o["metin"] = "\n".join(x for x in parcalar if x)
                o["cmp"] = _kiyas_metni(o["metin"])
                o["norm"] = normalize(o["metin"])
                o["tri"] = {o["norm"][i:i + 3] for i in range(len(o["norm"]) - 2)}
        motor, _teshis, _cevapli = cevap_motoru(kk)   # EŞLEŞTİR ile aynı motor
        for o in kk["ogeler"]:
            o["cevap"] = motor.get((o["ti"], o["no"])) if o["no"] else None
        acik = _acik_cevaplar(kk)
        for o, r, metin, renk in acik:
            if re.fullmatch(r"[A-E]", metin) and not o["cevap"]:
                o["cevap"] = metin
        soru_say = len(kk["ogeler"])
        cevapli = {id(o) for o, _r, _m, _c in acik}
        acik_say[kk["etiket"]] = len(acik)
        if soru_say and len(cevapli) >= 0.5 * soru_say:
            sorun("UYARI", kk["etiket"], f"Bu kitapçıkta cevaplar açık "
                  f"({len(cevapli)}/{soru_say} soruda renkli cevap/kod var) — "
                  f"baskı PDF'i değil, cevaplı dizgi PDF'i gibi görünüyor.")
        else:
            # Metin katmanında olmayan izler (eğri / resim): yalnız baskıda
            try:
                with fitz.open(kk["pdf"]) as d_:
                    izler = _oluk_piksel_izleri(kk, d_)
            except Exception as h:
                logla(f"Piksel katmanı taranamadı: {h}")
                izler = []
            for o, pno, r, cevap_rengi in izler:
                sorun("HATA" if cevap_rengi else "UYARI", kk["etiket"],
                      ("Cevap yerinde cevap renginde iz" if cevap_rengi else
                       "Cevap yerinde metin olmayan iz")
                      + f" ({r.width:.0f}×{r.height:.0f} pt): metin katmanında yok — "
                      f"eğriye çevrilmiş ya da resme gömülü cevap harfi olabilir"
                      + ("." if cevap_rengi else "; renkten bağımsız, elle bakın."),
                      test_adi(o["ti"]), o["ders"], o["etiket_no"],
                      [(kk["etiket"], (pno, r.x0 - 3, r.y0 - 3, r.x1 + 3, r.y1 + 3))],
                      gorunum=[(kk["etiket"], b_) for b_ in
                               tam_bolge.get((kk["etiket"], o["bolge"])) or [o["bolge"]]],
                      kutular=[(kk["etiket"], (pno, r.x0 - 3, r.y0 - 3, r.x1 + 3, r.y1 + 3))])
            for o, r, metin, renk in acik:
                tur = "cevap harfi" if len(metin) == 1 else "soru kodu"
                isaret = (kk["etiket"], (r["sayfa"], r["x0"] - 3, r["y0"] - 3,
                                         r["x1"] + 3, r["y1"] + 3))
                sorun("HATA", kk["etiket"], f"Açık kalan {tur}: '{metin}' "
                      f"(renk {renk}) — baskıda görünmemeli.", test_adi(o["ti"]),
                      o["ders"], o["etiket_no"], [isaret],
                      gorunum=[(kk["etiket"], b_) for b_ in
                               tam_bolge.get((kk["etiket"], o["bolge"])) or [o["bolge"]]],
                      kutular=[isaret])

    # --- 1) Yapı karşılaştırması (otomatik yapıda) -------------------------------
    if yapi_b is not None:
        ozet = lambda y: [(normalize(t["test"]),
                           sum(int(a) for _d, a in t["dersler"])) for t in y]
        if ozet(yapi) != ozet(yapi_b):
            sorun("HATA", "A-B", "Test yapısı farklı. A: " + " | ".join(
                f"{t['test']} ({sum(int(a) for _d, a in t['dersler'])})"
                for t in yapi) + "  —  B: " + " | ".join(
                f"{t['test']} ({sum(int(a) for _d, a in t['dersler'])})"
                for t in yapi_b))

    # --- 2) Numaralandırma (her kitapçık kendi içinde) ------------------------
    numara_ozet = {}
    for kk in (A, B):
        bulunan = {(o["ti"], o["no"]) for o in kk["ogeler"] if o["no"]}
        eksik_say = 0
        for ti, ders, no in kk["beklenen"]:
            if (ti, no) not in bulunan:
                eksik_say += 1
                # Raporda yeri: bir önceki ve bir sonraki soru (arada olmalı)
                komsu = sorted((o for o in kk["ogeler"] if o["ti"] == ti
                                and o["no"] in (no - 1, no + 1)),
                               key=lambda o: o["poz"])
                sorun("HATA", kk["etiket"], f"{no}. soru bulunamadı "
                      f"(numara atlanmış, yanlış yazılmış ya da soru basılmamış).",
                      test_adi(ti), ders, str(no),
                      gorunum=[(kk["etiket"], b_) for o in komsu for b_ in
                               tam_bolge.get((kk["etiket"], o["bolge"])) or [o["bolge"]]])
        for f in kk["fazlalar"]:
            o = next(x for x in kk["ogeler"] if x["poz"] == f["poz"])
            onceki = max((x for x in kk["ogeler"] if x["no"] and x["poz"] < o["poz"]
                          and x["ti"] == o["ti"]), key=lambda x: x["poz"],
                         default=None)
            tahmin = ""
            if onceki and (o["ti"], onceki["no"] + 1) not in bulunan:
                tahmin = f" Beklenen numara büyük olasılıkla {onceki['no'] + 1}."
            sorun("HATA", kk["etiket"], f"Numarası sıraya uymayan şıklı soru: "
                  f"'{f['no_yazan']}.' yazıyor (mükerrer/yanlış numara ya da "
                  f"fazladan soru).{tahmin}", test_adi(f["ti"]), f["ders"],
                  o["etiket_no"], [yer(kk, o)])
            o["_sorun"] = sorunlar[-1]
        for o in kk["ogeler"]:
            if o["no"] and o["a_sayisi"] >= 2:
                sorun("HATA", kk["etiket"], f"{o['no']}. sorunun içinde ikinci "
                      f"bir şık takımı var — numarasız ya da numarası okunamayan "
                      f"bir soru olabilir.", test_adi(o["ti"]), o["ders"],
                      str(o["no"]), [yer(kk, o)])
            elif o["no"] and o["a_sayisi"] == 0 and o["norm"]:
                sorun("UYARI", kk["etiket"], f"{o['no']}. soruda şık (A) …) "
                      f"bulunamadı — şıklar görsel olabilir, elle bakın.",
                      test_adi(o["ti"]), o["ders"], str(o["no"]), [yer(kk, o)])
        numara_ozet[kk["etiket"]] = (len(bulunan), len(kk["beklenen"]),
                                     eksik_say, len(kk["fazlalar"]))
        # Grup başlığı numaraları altındaki sorularla aynı mı?
        for b in kk["basliklar"]:
            alt = [o["no"] if o["no"] else o["etiket_no"] for o in b["sorular"]]
            if alt != b["nolar"]:
                sorun("HATA", kk["etiket"], "Metne bağlı grup başlığı "
                      f"{_nolar_yazi(b['nolar'])}. soruları gösteriyor ama "
                      f"altındaki sorular: {', '.join(map(str, alt)) or 'yok'}.",
                      test_adi(b["ti"]), "", _nolar_yazi(b["nolar"]),
                      [(kk["etiket"], b["baslik_bolge"])])

    # --- 3) A ↔ B soru eşleştirmesi (her test kendi içinde, dersler arası) -----
    docA, docB = fitz.open(a_pdf), fitz.open(b_pdf)
    onbellek = {}
    _isaret_onbellegi(A, B, docA, docB, onbellek)
    eslesmeler, eksik_a, eksik_b = _icerik_eslestir(A, B, docA, docB, onbellek,
                                                    cevap_a, cevap_b)
    eslesmeler = [(oa, ob, oran, None) for oa, ob, oran in eslesmeler]
    for ti, oa, en_iyi in eksik_a:
        sorun("HATA", "A-B", f"A-{oa['etiket_no']} sorusu B kitapçığında "
              f"bulunamadı (en yakın: B-{en_iyi['etiket_no']}, metin "
              f"%{_benzerlik_orani(oa['cmp'], en_iyi['cmp']) * 100:.0f}).",
              test_adi(ti), oa["ders"], f"A-{oa['etiket_no']}", [yer(A, oa)])
    for ti, ob in eksik_b:
        sorun("HATA", "A-B", f"B-{ob['etiket_no']} sorusunun A kitapçığında "
              f"karşılığı yok.", test_adi(ti), ob["ders"],
              f"B-{ob['etiket_no']}", [yer(B, ob)])

    # Eşleşen çiftlerin denetimi
    a_b, cift_satirlari = {}, []
    for oa, ob, oran, gorsel in eslesmeler:
        ti = oa["ti"]
        etiket = f"A-{oa['etiket_no']} / B-{ob['etiket_no']}"
        yerler = [yer(A, oa), yer(B, ob)]
        durum = []
        if oa["no"]:
            a_b[(ti, oa["no"])] = ob
        for o_, karsi in ((oa, f"B-{ob['etiket_no']}"), (ob, f"A-{oa['etiket_no']}")):
            if o_.get("_sorun"):         # numarası hatalı sorunun karşılığını yaz
                o_["_sorun"]["aciklama"] += f" İçerik {karsi} ile aynı soru."
        if _gorunur_ayni(oa["cmp"], ob["cmp"]):
            pass
        elif oa["norm"] == ob["norm"]:
            durum.append("NOKTALAMA")
            sorun("UYARI", "A-B", "Noktalama farkı: "
                  + _fark_ozeti(oa["cmp"], ob["cmp"]), test_adi(ti),
                  oa["ders"], etiket, yerler, vurgula=True)
        else:
            durum.append("METİN FARKI")
            sorun("HATA", "A-B", f"Metin farklı ({_yuzde(oran)}): "
                  + _fark_ozeti(oa["cmp"], ob["cmp"]), test_adi(ti),
                  oa["ders"], etiket, yerler, vurgula=True)
        if oa["ders"] != ob["ders"]:
            durum.append("DERS FARKLI")
            sorun("HATA", "A-B", f"Soru B'de başka dersin aralığında: A'da "
                  f"{oa['ders']}, B'de {ob['ders']}.", test_adi(ti), oa["ders"],
                  etiket, yerler)
        ca_ = cevap_a.get((ti, oa["no"])) or oa["cevap"] or ""
        cb_ = cevap_b.get((ti, ob["no"])) or ob["cevap"] or ""
        if ca_ and cb_ and ca_ != cb_:
            durum.append("CEVAP FARKLI")
            sorun("HATA", "A-B", f"Cevap farklı: A'da {ca_}, B'de {cb_}.",
                  test_adi(ti), oa["ders"], etiket, yerler)
        if oa["kod"] and ob["kod"] and oa["kod"] != ob["kod"]:
            durum.append("KOD FARKLI")
            sorun("UYARI", "A-B", f"Soru kodu farklı: A {oa['kod']}, "
                  f"B {ob['kod']}.", test_adi(ti), oa["ders"], etiket, yerler)
        sekil = _sekil_denetle(
            docA, docB, oa["nesne_bolgeleri"], ob["nesne_bolgeleri"], onbellek,
            durum, lambda onem, mesaj, kutular=(): sorun(
                onem, "A-B", mesaj, test_adi(ti), oa["ders"], etiket, yerler,
                kutular=kutular))
        cift_satirlari.append((oa, ob, oran, sekil, ca_, cb_, durum))

    # Aynı soru bir kitapçıkta iki kez basılmış mı?
    for kk in (A, B):
        og = [o for o in kk["ogeler"] if len(o["norm"]) >= 40]
        for i in range(len(og)):
            for j in range(i + 1, len(og)):
                x, y = og[i], og[j]
                if x["ti"] != y["ti"]:
                    continue
                u = x["tri"] | y["tri"]
                if u and len(x["tri"] & y["tri"]) / len(u) >= 0.8 and \
                        _benzerlik_orani(x["cmp"], y["cmp"]) >= 0.9:
                    sorun("HATA", kk["etiket"], f"Aynı soru iki kez basılmış: "
                          f"{x['etiket_no']}. ve {y['etiket_no']}. sorular.",
                          test_adi(x["ti"]), x["ders"],
                          f"{x['etiket_no']} / {y['etiket_no']}",
                          [yer(kk, x), yer(kk, y)])

    # --- 4) Metne bağlı gruplar A ↔ B -------------------------------------------
    grup_satirlari = []
    b_baslik_bul = {}
    for b in B["basliklar"]:
        for o in b["sorular"]:
            b_baslik_bul.setdefault(o["poz"], b)
    kullanilan_b_baslik = set()
    for ga in A["basliklar"]:
        ti = ga["ti"]
        a_nolar = [o["no"] for o in ga["sorular"] if o["no"]]
        karsilik = [a_b.get((ti, n)) for n in a_nolar]
        b_nolar = [ob["no"] if ob and ob["no"] else None for ob in karsilik]
        durum, parca_oran, parca_gorsel, gb = [], None, None, None
        etiket = f"A {_nolar_yazi(ga['nolar'])}"
        if not a_nolar or None in karsilik:
            durum.append("SORU EKSİK")
        else:
            gb = next((b_baslik_bul[ob["poz"]] for ob in karsilik
                       if ob["poz"] in b_baslik_bul), None)
            gecerli = [n for n in b_nolar if n is not None]
            if len(gecerli) != len(b_nolar):
                durum.append("B'DE NUMARA HATALI")   # ayrıntı numaralandırmada
            if len(gecerli) == len(b_nolar):
                if gecerli != list(range(gecerli[0], gecerli[0] + len(gecerli))):
                    # Grup içinde sıranın değişmesi (ör. B'de ters dizilmesi) hata
                    # değildir; yan yana ve aynı metnin altında olmaları yeter
                    if sorted(gecerli) != list(range(min(gecerli),
                                                     min(gecerli) + len(gecerli))):
                        durum.append("GRUP DAĞILMIŞ")
                        sorun("HATA", "A-B", f"Metne bağlı grup B'de dağılmış: A "
                              f"{_nolar_yazi(a_nolar)} → B "
                              f"{', '.join(map(str, gecerli))} (yan yana değil).",
                              test_adi(ti), "", etiket,
                              [(A["etiket"], ga["bolge"])] +
                              [yer(B, ob) for ob in karsilik])
            if gb is None or not ({id(ob) for ob in karsilik} <=
                                  {id(o) for o in gb["sorular"]}):
                durum.append("B'DE METİN/BAŞLIK YOK")
                sorun("HATA", "A-B", f"A {_nolar_yazi(a_nolar)} metne bağlı "
                      f"grubunun karşılığı B'de bir grup başlığı/metin altında "
                      f"değil (B {', '.join(map(str, b_nolar))}).", test_adi(ti),
                      "", etiket, [yer(B, ob) for ob in karsilik])
                gb = None
            if gb is not None:
                kullanilan_b_baslik.add(id(gb))
                if None not in b_nolar and gb["nolar"] != sorted(b_nolar):
                    durum.append("B BAŞLIĞI YANLIŞ")   # ayrıntı numaralandırmada
                parca_oran = _benzerlik_orani(ga["cmp"], gb["cmp"])
                if not _gorunur_ayni(ga["cmp"], gb["cmp"]):
                    if normalize(ga["cmp"]) == normalize(gb["cmp"]):
                        durum.append("NOKTALAMA")
                        sorun("UYARI", "A-B", "Parça metninde noktalama farkı: "
                              + _fark_ozeti(ga["cmp"], gb["cmp"]),
                              test_adi(ti), "", etiket,
                              [(A["etiket"], ga["bolge"]), (B["etiket"], gb["bolge"])],
                              vurgula=True)
                    else:
                        durum.append("PARÇA FARKLI")
                        # B'deki metin başka bir A grubunun metni mi?
                        # (metin ile altındaki sorular karışmış)
                        baska = max((x for x in A["basliklar"] if x is not ga),
                                    key=lambda x: _benzerlik_orani(x["cmp"], gb["cmp"]),
                                    default=None)
                        if baska is not None and \
                                _benzerlik_orani(baska["cmp"], gb["cmp"]) >= 0.9:
                            mesaj = (f"B'de metin ile altındaki sorular uyumsuz: "
                                     f"B {_nolar_yazi(gb['nolar'])}. soruların "
                                     f"üstündeki metin, A'da "
                                     f"{_nolar_yazi(baska['nolar'])}. soruların "
                                     f"metni (bu sorular A'da "
                                     f"{_nolar_yazi(a_nolar)} metnine bağlı).")
                        else:
                            mesaj = (f"Parça metni farklı ({_yuzde(parca_oran)}): "
                                     + _fark_ozeti(ga["cmp"], gb["cmp"]))
                        sorun("HATA", "A-B", mesaj, test_adi(ti), "", etiket,
                              [(A["etiket"], ga["bolge"]), (B["etiket"], gb["bolge"])],
                              vurgula=baska is None or
                              _benzerlik_orani(baska["cmp"], gb["cmp"]) < 0.9)
                parca_gorsel = _sekil_denetle(
                    docA, docB, ga["nesne_bolgeleri"], gb["nesne_bolgeleri"],
                    onbellek, durum, lambda onem, mesaj, kutular=(): sorun(
                        onem, "A-B", "Parçada: " + mesaj, test_adi(ti), "", etiket,
                        [(A["etiket"], ga["bolge"]), (B["etiket"], gb["bolge"])],
                        kutular=kutular))
        grup_satirlari.append((ti, ga, gb, a_nolar, b_nolar, parca_oran,
                               parca_gorsel, durum))
    for gb in B["basliklar"]:
        if id(gb) not in kullanilan_b_baslik:
            nolar = [o["no"] for o in gb["sorular"] if o["no"]]
            karsi = [k for k, ob in a_b.items()
                     if ob["no"] in nolar and ob["ti"] == gb["ti"]]
            sorun("HATA", "B", f"B'deki {_nolar_yazi(gb['nolar'])}. soruların "
                  f"metne bağlı grup başlığının A'da karşılığı yok "
                  f"(A karşılıkları: {', '.join(str(k[1]) for k in sorted(karsi)) or '—'}).",
                  test_adi(gb["ti"]), "", f"B {_nolar_yazi(gb['nolar'])}",
                  [(B["etiket"], gb["bolge"])])

    # --- 5) Rapor için: gösterilecek bölgeler, farklı kelimeler, sıra ------------
    for s in sorunlar:
        if s["gorunum"] is None:
            s["gorunum"] = [(et, b_) for et, b in s["yerler"]
                            for b_ in tam_bolge.get((et, b)) or [b]]
        if s.pop("vurgula"):
            try:
                s["vurgu"] = _fark_kelimeleri({"A": docA, "B": docB}, s["gorunum"])
            except Exception as h:       # yalnızca görsel yardım; kontrolü bozmasın
                logla(f"Fark vurgusu çıkarılamadı: {h}")
    _sorunlari_sirala(sorunlar)
    docA.close()
    docB.close()

    sonuc = {"A": A, "B": B, "yapi": yapi, "eslesmeler": cift_satirlari,
             "gruplar": grup_satirlari, "numara": numara_ozet,
             "cevap_var": (sum(1 for o in A["ogeler"] if o["cevap"]),
                           sum(1 for o in B["ogeler"] if o["cevap"])),
             "acik_cevap": acik_say, "test_adlari": adlar}
    return sonuc, sorunlar


def _nolar_yazi(nolar):
    nolar = [n for n in nolar if isinstance(n, int)]
    if not nolar:
        return "?"
    if len(nolar) > 2 and nolar == list(range(nolar[0], nolar[-1] + 1)):
        return f"{nolar[0]}-{nolar[-1]}"
    return ", ".join(map(str, nolar))


def kart_yolu(rapor):
    """Sorun kartları PDF'inin yolu (raporun yanında)."""
    rapor = Path(rapor)
    return rapor.with_name(rapor.stem + "_sorunlar.pdf")


def kontrol_raporu_yaz(sonuc, sorunlar, cikti, kartlar=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    KIRMIZI = PatternFill("solid", fgColor="E53935")    # hata: açık kırmızı
    SARI = PatternFill("solid", fgColor="FFE699")
    YESIL = PatternFill("solid", fgColor="C6EFCE")
    KALIN = Font(name="Tahoma", bold=True)
    BEYAZ_KALIN = Font(name="Tahoma", bold=True, color="FFFFFF")
    KOYU_KIRMIZI = Font(name="Tahoma", bold=True, color="B71C1C")

    def ekle(sayfa, satir):
        sayfa.append([_xl(v) for v in satir])

    def boya(hucre, tur):
        """tur: HATA (kırmızı, beyaz kalın yazı), UYARI/KONTROL ET (sarı), TEMİZ."""
        if tur == "HATA":
            hucre.fill, hucre.font = KIRMIZI, BEYAZ_KALIN
        elif tur == "TEMİZ":
            hucre.fill = YESIL
        else:
            hucre.fill = SARI
    adlar = sonuc["test_adlari"]
    test_adi = lambda ti: adlar[ti] if ti is not None and ti < len(adlar) else ""

    wb = openpyxl.Workbook()
    oz = wb.active
    oz.title = "Özet"
    hata = sum(1 for s in sorunlar if s["onem"] == "HATA")
    uyari = sum(1 for s in sorunlar if s["onem"] == "UYARI")
    es = sonuc["eslesmeler"]
    ayni = sum(1 for e in es if not e[6])
    toplam_a = sonuc["numara"]["A"][1]
    ekle(oz, ["A–B KİTAPÇIK KONTROL RAPORU", f"Motor sürümü {SURUM}"])
    ekle(oz, ["A kitapçığı", Path(sonuc["A"]["pdf"]).name])
    ekle(oz, ["B kitapçığı", Path(sonuc["B"]["pdf"]).name])
    ekle(oz, [])
    ekle(oz, ["Kontrol", "Sonuç", "Durum"])
    for c in oz[5]:
        c.font = KALIN
    satirlar = [
        ("Genel sonuç", f"{hata} hata, {uyari} uyarı"
         + (" — hangi soru, hangi sayfa: aşağıda" if sorunlar else ""),
         "TEMİZ" if not hata and not uyari else ("HATA" if hata else "UYARI")),
    ]
    for et in ("A", "B"):
        bul, bek, eks, faz = sonuc["numara"][et]
        satirlar.append((f"{et} kitapçığı numaralandırma",
                         f"{bul}/{bek} soru yerinde, {eks} eksik, {faz} sırasız numara",
                         "TEMİZ" if not eks and not faz else "HATA"))
    eslesen_a = sum(1 for e in es if e[0]["no"])
    satirlar.append(("A'daki her soru B'de var mı?",
                     f"{eslesen_a}/{toplam_a} soru B'de bulundu",
                     "TEMİZ" if eslesen_a == toplam_a else "HATA"))
    satirlar.append(("Soru metni + şıklar birebir",
                     f"{ayni}/{len(es)} çift tamamen aynı",
                     "TEMİZ" if ayni == len(es) else "KONTROL ET"))
    ca, cb = sonuc["cevap_var"]
    cevap_fark = sum(1 for e in es if "CEVAP FARKLI" in e[6])
    satirlar.append(("Cevap harfleri", f"{cevap_fark} farklı (A'da {ca}, B'de {cb} "
                     f"renkli cevap okundu)" + ("" if cb else " — B'de cevap yok, "
                                               "kıyas yapılamadı"),
                     "TEMİZ" if not cevap_fark else "HATA"))
    gr = sonuc["gruplar"]
    gr_ok = sum(1 for g in gr if not g[7])
    satirlar.append(("Metne bağlı gruplar", f"{gr_ok}/{len(gr)} grup B'de aynı",
                     "TEMİZ" if gr_ok == len(gr) else "HATA"))
    sekil_fark = sum(1 for e in es if "ŞEKİL FARKI" in e[6]) + \
        sum(1 for g in gr if "ŞEKİL FARKI" in g[7])
    satirlar.append(("Şekil, resim, tablo çizimleri",
                     f"{sekil_fark} soru/metinde fark",
                     "TEMİZ" if not sekil_fark else "KONTROL ET"))
    for et in ("A", "B"):
        acik = sonuc["acik_cevap"][et]
        satirlar.append((f"Açık kalan cevap ({et})",
                         "yok" if not acik else f"{acik} işaret bulundu",
                         "TEMİZ" if not acik else "HATA"))
    for s in satirlar:
        ekle(oz, list(s))
        boya(oz.cell(row=oz.max_row, column=3), s[2])
    oz.column_dimensions["A"].width = 40
    oz.column_dimensions["B"].width = 70
    oz.column_dimensions["C"].width = 16

    # Nereye bakmalı: her sorun, sayfasıyla birlikte doğrudan özette
    ekle(oz, [])
    if sorunlar:
        ekle(oz, ["NEREYE BAKMALI?", f"{len(sorunlar)} sorun — kitapçıktaki "
                   f"sırasıyla (önce hatalar)", "Sayfa"])
        for c in oz[oz.max_row]:
            c.font = KALIN
        if kartlar:
            ekle(oz, ["Görüntüler", f"Her sorunun A ve B görüntüsü yan yana, "
                       f"farklı kelimeler sarı: {Path(kartlar).name}", ""])
            c = oz.cell(row=oz.max_row, column=2)
            c.hyperlink = Path(kartlar).name
            c.font = Font(name="Tahoma", color="0563C1", underline="single")
        for s in sorunlar:
            ekle(oz, [f"{s['sira']}. {s['onem']} · {s['test'] or s['kitapcik']}"
                       + (f" · {s['soru']}" if s["soru"] else ""),
                       s["aciklama"], _sorun_yeri(s)])
            boya(oz.cell(row=oz.max_row, column=1), s["onem"])
            if s["onem"] == "HATA":           # açıklama da kırmızı okunsun
                oz.cell(row=oz.max_row, column=2).font = KOYU_KIRMIZI
            oz.cell(row=oz.max_row, column=2).alignment = Alignment(
                wrap_text=True, vertical="top")
            oz.cell(row=oz.max_row, column=1).alignment = Alignment(vertical="top")
            oz.cell(row=oz.max_row, column=3).alignment = Alignment(vertical="top")
    else:
        ekle(oz, ["NEREYE BAKMALI?", "Bakılacak soru yok — A ile B kitapçığında "
                   "sorun bulunamadı.", ""])
        oz.cell(row=oz.max_row, column=1).font = KALIN

    ws = wb.create_sheet("Sorunlar")
    ekle(ws, ["Sıra", "Önem", "Kitapçık", "Test", "Ders", "Soru", "Sayfa",
               "Açıklama"])
    for s in sorunlar:
        ekle(ws, [s.get("sira", ""), s["onem"], s["kitapcik"], s["test"],
                   s["ders"], s["soru"], _sorun_yeri(s), s["aciklama"]])
        boya(ws.cell(row=ws.max_row, column=2), s["onem"])
        if s["onem"] == "HATA":
            ws.cell(row=ws.max_row, column=8).font = KOYU_KIRMIZI
    if not sorunlar:
        ekle(ws, ["", "", "", "", "", "", "", "Sorun bulunamadı."])
    for col, gen in zip("ABCDEFGH", (6, 8, 9, 26, 18, 22, 14, 110)):
        ws.column_dimensions[col].width = gen

    se = wb.create_sheet("Soru Eşleşmesi")
    ekle(se, ["Test", "Ders (A)", "A No", "B No", "Ders (B)", "Metin %",
               "Şekil/Resim", "A Cevap", "B Cevap", "A Sayfa", "B Sayfa", "Durum"])
    for oa, ob, oran, gorsel, ca_, cb_, durum in sorted(
            es, key=lambda e: (e[0]["ti"], e[0]["no"] or 999)):
        ekle(se, [test_adi(oa["ti"]), oa["ders"], oa["etiket_no"], ob["etiket_no"],
                   ob["ders"], round(oran * 100, 1),
                   gorsel,
                   ca_, cb_, oa["sayfa"] + 1, ob["sayfa"] + 1,
                   ", ".join(durum) or "OK"])
        boya(se.cell(row=se.max_row, column=12),
             "TEMİZ" if not durum else
             ("UYARI" if set(durum) <= {"NOKTALAMA", "ŞEKİL FARKI", "KOD FARKLI"}
              else "HATA"))

    gs = wb.create_sheet("Metne Bağlı Gruplar")
    ekle(gs, ["Test", "A Başlığı", "A Sorular", "B Başlığı", "B Karşılıkları",
               "Parça Metni %", "Parça Şekil/Resim", "Durum"])
    for ti, ga, gb, a_nolar, b_nolar, p_oran, p_gorsel, durum in gr:
        ekle(gs, [test_adi(ti), _nolar_yazi(ga["nolar"]), _nolar_yazi(a_nolar),
                   _nolar_yazi(gb["nolar"]) if gb else "—",
                   ", ".join(str(n) if n else "?" for n in b_nolar),
                   round(p_oran * 100, 1) if p_oran is not None else "",
                   p_gorsel or "",
                   ", ".join(durum) or "OK"])
        boya(gs.cell(row=gs.max_row, column=8), "TEMİZ" if not durum else
             ("UYARI" if set(durum) <= {"NOKTALAMA"} else "HATA"))
    for sayfa in (se, gs):
        for c in sayfa[1]:
            c.font = KALIN
        sayfa.freeze_panes = "A2"
    for c in ws[1]:
        c.font = KALIN
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(cikti)


CIKTI_ISARETI = "Kitapcik Eslestirici kontrol ciktisi"


def eski_ciktiyi_sil(yol):
    """Önceki kontrolden kalan çıktı PDF'ini siler (yalnızca bu uygulamanın
    yazdığı, işaretli dosyayı: aynı adlı başka bir dosyaya dokunulmaz)."""
    yol = Path(yol)
    try:
        if yol.exists():
            with fitz.open(str(yol)) as d:
                bizim = (d.metadata or {}).get("keywords") == CIKTI_ISARETI
            if bizim:
                yol.unlink()
    except Exception as h:
        logla(f"Eski çıktı silinemedi ({yol}): {h}")


def kontrol_pdf_isaretle(sonuc, sorunlar, klasor):
    """Sorunlu soruların üzerine kırmızı çerçeve + açıklama notu koyar; farklı
    kelimeleri sarıyla vurgular."""
    yazilan = []
    for et in ("A", "B"):
        kaynak = sonuc[et]["pdf"]
        hedef = Path(klasor) / (Path(kaynak).stem + "_kontrol.pdf")
        eski_ciktiyi_sil(hedef)
        isler = [(s, b) for s in sorunlar for k, b in s["yerler"] if k == et]
        if not isler:
            continue
        doc = fitz.open(kaynak)
        for s, (pno, x0, y0, x1, y1) in isler:
            sayfa = doc[pno]
            r = fitz.Rect(x0, y0, x1, y1) & sayfa.rect
            if r.is_empty:
                continue
            renk = (0.9, 0.1, 0.1) if s["onem"] == "HATA" else (0.95, 0.6, 0.0)
            a = sayfa.add_rect_annot(r)
            a.set_colors(stroke=renk)
            a.set_border(width=1.5)
            a.set_info(title=f"{s.get('sira', '')}. {s['onem']} — {s['soru']}",
                       content=s["aciklama"])
            a.update()
        for s in sorunlar:
            if not any(k == et for k, _b in s["yerler"]):
                continue
            for k, (pno, x0, y0, x1, y1) in s["kutular"]:
                sayfa = doc[pno]             # (sayfa nesnesi açıkken eklenmeli)
                r = fitz.Rect(x0, y0, x1, y1) & sayfa.rect
                if k == et and not r.is_empty:
                    a = sayfa.add_rect_annot(r)
                    a.set_colors(stroke=(0.9, 0.1, 0.1))
                    a.set_border(width=1, dashes=[2, 2])
                    a.update()
            for k, pno, rc in s["vurgu"]:
                if k == et:
                    sayfa = doc[pno]
                    a = sayfa.add_highlight_annot(fitz.Rect(rc))
                    a.set_colors(stroke=(1, 0.85, 0))
                    a.update()
        md = dict(doc.metadata or {})
        md["keywords"] = CIKTI_ISARETI
        doc.set_metadata({k: v for k, v in md.items()
                          if k not in ("format", "encryption")})
        doc.save(str(hedef), garbage=1, deflate=True)
        doc.close()
        yazilan.append(str(hedef))
    return yazilan


def _kart_yazi(sayfa, nokta, metin, font, boyut, renk=(0, 0, 0)):
    tw = fitz.TextWriter(sayfa.rect, color=renk)
    tw.append(nokta, metin, font=font, fontsize=boyut)
    tw.write_text(sayfa)


def _kart_kisalt(metin, font, boyut, genislik):
    if font.text_length(metin, fontsize=boyut) <= genislik:
        return metin
    while metin and font.text_length(metin + "…", fontsize=boyut) > genislik:
        metin = metin[:-1]
    return metin + "…"


def kontrol_kartlari_yaz(sonuc, sorunlar, hedef, azami=300):
    """Her sorun için bir sayfa: üstte ne bulunduğu, altta sorunun A ve B
    kitapçığındaki görüntüsü yan yana (ya da alt alta). Farklı kelimeler sarı,
    farklı şekil bölgesi ve açık cevap kırmızı çerçeveli. İlk sayfa(lar) tıklanır
    içindekiler listesidir."""
    W, H, M = 595, 842, 28
    font, kalin = fitz.Font("helv"), fitz.Font("hebo")
    KIRMIZI, TURUNCU, GRI = (0.8, 0.1, 0.1), (0.85, 0.5, 0.0), (0.35, 0.35, 0.35)
    kaynak = {et: fitz.open(sonuc[et]["pdf"]) for et in ("A", "B")}
    out = fitz.open()
    gosterilen = sorunlar[:azami]
    SATIR = 16
    ilk_satir = 4                       # içindekiler sayfasında başlık satırları
    sayfa_basi = int((H - 2 * M) / SATIR) - ilk_satir
    icindekiler = max(1, -(-len(gosterilen) // sayfa_basi))
    for _ in range(icindekiler):
        out.new_page(width=W, height=H)
    toc = []
    try:
        for s in gosterilen:
            sayfa = out.new_page(width=W, height=H)
            renk = KIRMIZI if s["onem"] == "HATA" else TURUNCU
            baslik = (f"{s['sira']}. {s['onem']} · {s['test'] or s['kitapcik']}"
                      + (f" · {s['soru']}" if s["soru"] else ""))
            _kart_yazi(sayfa, (M, M + 12), _kart_kisalt(baslik, kalin, 13, W - 2 * M),
                       kalin, 13, renk)
            _kart_yazi(sayfa, (M, M + 28), "Nerede: " + (_sorun_yeri(s) or "—"),
                       font, 10, GRI)
            tw = fitz.TextWriter(sayfa.rect)
            kutu = fitz.Rect(M, M + 36, W - M, M + 36 + 4 * 12.5)
            tw.fill_textbox(kutu, s["aciklama"], font=font, fontsize=10)
            tw.write_text(sayfa)
            ust = kutu.y1 + 10
            toc.append([1, _kart_kisalt(baslik, font, 9, 400), out.page_count])
            # Paneller: her kitapçık için gösterilecek bölgeler
            paneller = []
            for et in ("A", "B"):
                parcalar = []
                for e, (pno, x0, y0, x1, y1) in s.get("gorunum") or []:
                    if e != et:
                        continue
                    sr = kaynak[et][pno].rect
                    r = fitz.Rect(x0, y0, x1, y1)
                    if r.height < 40:          # tek satırlık yer: çevresiyle göster
                        r.y0, r.y1 = r.y0 - 30, r.y1 + 30
                    if r.width < 160:
                        orta = (r.x0 + r.x1) / 2
                        r.x0, r.x1 = orta - 80, orta + 80
                    r &= sr
                    if not r.is_empty and all((pno, r) != (p_, r_) for p_, r_ in parcalar):
                        parcalar.append((pno, r))
                if parcalar:
                    paneller.append((et, parcalar))
            if not paneller:
                _kart_yazi(sayfa, (M, ust + 14), "Bu sorun için gösterilecek "
                           "görüntü yok (ayrıntı yukarıda).", font, 10, GRI)
                continue
            def sigar(parcalar, alan):         # bu alana sığan en büyük ölçek
                uygun_h = alan.height - 16 - 6 * (len(parcalar) - 1)
                return min(1.6, alan.width / max(r.width for _p, r in parcalar),
                           uygun_h / sum(r.height for _p, r in parcalar))
            if len(paneller) == 2:
                # Yan yana ya da alt alta: büyük gösterene göre (yan yana öncelikli);
                # iki kitapçık aynı ölçekte gösterilir
                pw = (W - 2 * M - 14) / 2
                yan = [fitz.Rect(M, ust, M + pw, H - M),
                       fitz.Rect(M + pw + 14, ust, W - M, H - M)]
                ph = (H - M - ust - 12) / 2
                alt = [fitz.Rect(M, ust, W - M, ust + ph),
                       fitz.Rect(M, ust + ph + 12, W - M, H - M)]
                o_yan = min(sigar(ps, a_) for (_e, ps), a_ in zip(paneller, yan))
                o_alt = min(sigar(ps, a_) for (_e, ps), a_ in zip(paneller, alt))
                alanlar, olcek = (yan, o_yan) if o_yan >= 0.8 * o_alt else (alt, o_alt)
            else:
                alanlar = [fitz.Rect(M, ust, W - M, H - M)]
                olcek = sigar(paneller[0][1], alanlar[0])
            yerlesim = []
            for (et, parcalar), alan in zip(paneller, alanlar):
                sayfa_nolar = sorted({p_ + 1 for p_, _r in parcalar})
                _kart_yazi(sayfa, (alan.x0, alan.y0 + 10),
                           f"{et} KİTAPÇIĞI — sayfa {', '.join(map(str, sayfa_nolar))}",
                           kalin, 10, (0.1, 0.25, 0.55))
                y = alan.y0 + 16
                for pno, r in parcalar:
                    hr = fitz.Rect(alan.x0, y, alan.x0 + r.width * olcek,
                                   y + r.height * olcek)
                    sayfa.show_pdf_page(hr, kaynak[et], pno, clip=r)
                    sayfa.draw_rect(hr, color=(0.7, 0.7, 0.7), width=0.6)
                    yerlesim.append((et, pno, r, hr, olcek))
                    y = hr.y1 + 6

            def donustur(et, pno, rc):
                rc = fitz.Rect(rc)
                for e, p_, r, hr, o in yerlesim:
                    k = rc & r
                    if e == et and p_ == pno and not k.is_empty:
                        return fitz.Rect(hr.x0 + (k.x0 - r.x0) * o,
                                         hr.y0 + (k.y0 - r.y0) * o,
                                         hr.x0 + (k.x1 - r.x0) * o,
                                         hr.y0 + (k.y1 - r.y0) * o)
                return None
            for et, pno, rc in s["vurgu"]:
                hr = donustur(et, pno, rc)
                if hr is not None:
                    sayfa.draw_rect(hr, color=None, fill=(1, 0.85, 0),
                                    fill_opacity=0.45, overlay=True)
            for et, (pno, x0, y0, x1, y1) in s["kutular"]:
                hr = donustur(et, pno, (x0, y0, x1, y1))
                if hr is not None:
                    sayfa.draw_rect(hr, color=KIRMIZI, width=1.4)
        # İçindekiler (tıklanır)
        sorun_say = {o: sum(1 for s in sorunlar if s["onem"] == o)
                     for o in ("HATA", "UYARI")}
        for i, s in enumerate(gosterilen):
            sayfa = out[i // sayfa_basi]
            satir = i % sayfa_basi
            if satir == 0:
                _kart_yazi(sayfa, (M, M + 14), "A–B KONTROL — BAKILACAK SORULAR",
                           kalin, 14, (0.1, 0.25, 0.55))
                _kart_yazi(sayfa, (M, M + 32), _kart_kisalt(
                    f"{Path(sonuc['A']['pdf']).name}  ↔  {Path(sonuc['B']['pdf']).name}"
                    f"   ·   {sorun_say['HATA']} hata, {sorun_say['UYARI']} uyarı"
                    + (f" (ilk {azami} gösteriliyor)" if len(sorunlar) > azami else ""),
                    font, 9, W - 2 * M), font, 9, GRI)
            y = M + (ilk_satir + satir) * SATIR
            renk = KIRMIZI if s["onem"] == "HATA" else TURUNCU
            sol = (f"{s['sira']}. {s['onem']} · {s['test'] or s['kitapcik']}"
                   + (f" · {s['soru']}" if s["soru"] else ""))
            sol = _kart_kisalt(sol, kalin, 9, 230)
            _kart_yazi(sayfa, (M, y), sol, kalin, 9, renk)
            _kart_yazi(sayfa, (M + 236, y), _kart_kisalt(
                f"{_sorun_yeri(s) or '—'}  ·  {s['aciklama']}", font, 9,
                W - 2 * M - 236), font, 9)
            sayfa.insert_link({"kind": fitz.LINK_GOTO, "page": icindekiler + i,
                               "from": fitz.Rect(M, y - 10, W - M, y + 4),
                               "to": fitz.Point(0, 0)})
        out.set_toc(toc)
        out.set_metadata({"title": "A–B kontrol: bakılacak sorular",
                          "creator": f"Kitapçık Eşleştirici {SURUM}",
                          "keywords": CIKTI_ISARETI})
        out.save(str(hedef), garbage=3, deflate=True)
    finally:
        out.close()
        for d in kaynak.values():
            d.close()
    return str(hedef)


def kontrol_calistir(a_pdf, b_pdf, sinav, anah_a, anah_b, cikti, log,
                     yol_sor=None):
    log(f"Motor sürümü: {SURUM} — A–B KONTROL")
    okunurluk_denetle(a_pdf, b_pdf)
    yapi_b = None
    if sinav in ("OTOMATİK", "ORTAOKUL"):
        log("Yapı A ve B kitapçıklarından çıkarılıyor...")
        yapi = yapi_cikar(a_pdf)
        yapi_b = yapi_cikar(b_pdf)
        log(f"  {len(yapi)} test bulundu ("
            + ", ".join(str(sum(int(a) for _d, a in t["dersler"])) for t in yapi)
            + " soru)")
    else:
        yapi = YAPILAR[sinav]
        # Seçilen tür dosyaya uymuyorsa (EŞLEŞTİR'deki gibi) diğerini dene
        _s, u = sorulari_ayikla(a_pdf, yapi, "A", lambda _m: None)
        if len(u) > 5:
            en_iyi = len(u)
            for aday_sinav, aday_yapi in YAPILAR.items():
                if aday_sinav == sinav:
                    continue
                _s2, u2 = sorulari_ayikla(a_pdf, aday_yapi, "A", lambda _m: None)
                if len(u2) < len(u):
                    log(f"⚠ DİKKAT: Dosyalar '{aday_sinav}' düzeninde; sınav "
                        f"türü otomatik '{aday_sinav}' olarak düzeltildi.")
                    sinav, yapi, en_iyi = aday_sinav, aday_yapi, len(u2)
                    break
            if en_iyi > 5:            # TYT de AYT de uymuyor: kitapçıktan çıkar
                try:
                    oto = yapi_cikar(a_pdf)
                    _s3, u3 = sorulari_ayikla(a_pdf, oto, "A", lambda _m: None)
                    if len(u3) < en_iyi:
                        log("⚠ DİKKAT: Dosyalar TYT/AYT düzeninde değil; yapı "
                            "kitapçıktan çıkarıldı (Lise/Ortaokul seçeneği gibi).")
                        yapi, yapi_b = oto, yapi_cikar(b_pdf)
                except RuntimeError as h:
                    logla(f"Otomatik yapı denenemedi: {h}")
    sonuc, sorunlar = kontrol_et(a_pdf, b_pdf, yapi, log,
                                 anahtar_oku(anah_a, yapi),
                                 anahtar_oku(anah_b, yapi), yapi_b=yapi_b)
    hedef = cikti
    while True:
        try:
            kontrol_raporu_yaz(sonuc, sorunlar, hedef,
                               kartlar=kart_yolu(hedef) if sorunlar else None)
            break
        except OSError as h:
            log(f"\n⚠ '{hedef}' konumuna yazılamadı ({h}).")
            yeni = yol_sor(hedef) if yol_sor else None
            if not yeni:
                raise RuntimeError("Rapor kaydedilemedi. Çıktı konumunu "
                                   "Seç... düğmesiyle belirleyin.") from h
            hedef = yeni
    try:
        pdfler = kontrol_pdf_isaretle(sonuc, sorunlar, Path(hedef).parent)
    except Exception as h:              # işaretli PDF isteğe bağlıdır
        logla(f"İşaretli PDF yazılamadı: {h}")
        pdfler = []
    kartlar = None
    eski_ciktiyi_sil(kart_yolu(hedef))
    if sorunlar:
        try:
            kartlar = kontrol_kartlari_yaz(sonuc, sorunlar, kart_yolu(hedef))
        except Exception:               # görsel yardımcı; rapor yine geçerli
            logla("Sorun kartları yazılamadı:\n" + traceback.format_exc())
    hata = [s for s in sorunlar if s["onem"] == "HATA"]
    uyari = [s for s in sorunlar if s["onem"] == "UYARI"]
    log("  Testler: " + " | ".join(
        f"{ad} ({sum(int(a) for _d, a in t['dersler'])})"
        for ad, t in zip(sonuc["test_adlari"], yapi)))
    log(f"\n✓ Kontrol raporu: {hedef}")
    if kartlar:
        log(f"  Bakılacak soruların A/B görüntüleri: {kartlar}")
    for p in pdfler:
        log(f"  İşaretli PDF: {p}")
    for et in ("A", "B"):
        bul, bek, eks, faz = sonuc["numara"][et]
        log(f"  {et} kitapçığı: {bul}/{bek} soru yerinde"
            + (f", {eks} eksik" if eks else "") + (f", {faz} sırasız numara" if faz else ""))
    es = sonuc["eslesmeler"]
    log(f"  A→B eşleşen soru: {sum(1 for e in es if e[0]['no'])}/"
        f"{sonuc['numara']['A'][1]}  |  metin birebir: "
        f"{sum(1 for e in es if not e[6])}/{len(es)}")
    gr = sonuc["gruplar"]
    if gr:
        log(f"  Metne bağlı gruplar: {sum(1 for g in gr if not g[7])}/{len(gr)} sorunsuz")
    sekil = sum(1 for e in es if "ŞEKİL FARKI" in e[6])
    log(f"  Şekil/resim farkı: {sekil} soruda" if sekil else
        "  Şekil/resimler: farksız")
    ac = sonuc["acik_cevap"]
    log(f"  Açık kalan cevap işareti: A {ac['A']}, B {ac['B']}")
    if not sorunlar:
        log("  ✓ Hiç sorun bulunmadı: A'daki her soru (metni, şıkları ve "
            "şekilleriyle) B'de aynen var, numaralar düzgün, açık cevap yok.")
    else:
        log(f"  {len(hata)} HATA, {len(uyari)} UYARI — nereye bakılacağı:")
        for s in sorunlar:
            log(f"  {s['sira']}. {'✗' if s['onem'] == 'HATA' else '⚠'} "
                f"[{s['kitapcik']}] {s['test']} {s['soru']} "
                f"({_sorun_yeri(s) or 'sayfa yok'}): {s['aciklama']}")
    return hedef


# ----------------------------------------------------------------------------
# ARAYÜZ (tkinter)
# ----------------------------------------------------------------------------
def gui_baslat():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext, ttk

    # --- eksik kütüphane kontrolü ve otomatik kurulum (görünür pencereyle) ---
    eksik = eksikleri_bul()
    if eksik:
        kok = tk.Tk()
        kok.title("Kitapçık Eşleştirici")
        kok.geometry("400x110")
        kok.attributes("-topmost", True)
        durum = tk.Label(kok, text="Başlatılıyor...", pady=20)
        durum.pack(fill="both", expand=True)
        kok.update()
        if messagebox.askyesno("Kurulum gerekli",
                               "Şu kütüphaneler eksik:\n  " + ", ".join(eksik) +
                               "\n\nŞimdi otomatik kurulsun mu? (internet gerekir)",
                               parent=kok):
            durum.configure(text="Kütüphaneler kuruluyor...\n"
                                 "Bu birkaç dakika sürebilir, pencereyi kapatmayın.")
            kok.update()
            komut = [sys.executable, "-m", "pip", "install", "--user"] + eksik
            p = subprocess.run(komut, capture_output=True, text=True)
            if p.returncode != 0:
                p = subprocess.run(komut + ["--break-system-packages"],
                                   capture_output=True, text=True)
            logla("pip kurulum çıktısı:\n" + (p.stdout or "") + (p.stderr or ""))
            importlib.invalidate_caches()
            try:
                import site
                site.addsitedir(site.getusersitepackages())
            except Exception:
                pass
            kalan = eksikleri_bul()
            istege_bagli = {"tkinterdnd2", "python-docx"}
            kritik = [p for p in kalan if p not in istege_bagli]
            if kalan and not kritik:
                logla(f"İsteğe bağlı kurulamayan: {kalan} — uygulama tam "
                      f"çalışır (ilgili özellik kapalı kalır).")
            if kritik:
                messagebox.showerror(
                    "Kurulamadı",
                    "Otomatik kurulum başarısız oldu.\nTerminale şunu yapıştırın:\n\n"
                    f"python3 -m pip install --user {' '.join(eksik)}\n\n"
                    "Sonra uygulamayı yeniden açın.\n\nHata:\n" + p.stderr[-400:],
                    parent=kok)
                kok.destroy()
                return
            messagebox.showinfo("Tamam", "Kurulum bitti, uygulama açılıyor.",
                                parent=kok)
        else:
            kok.destroy()
            return
        kok.destroy()
    try:
        bagimliliklari_yukle()
    except ImportError as h:
        if "incompatible architecture" not in str(h):
            raise
        # Paketler yanlış işlemci mimarisiyle kurulmuş (Intel/Apple karışması)
        kok = tk.Tk()
        kok.withdraw()
        onar = messagebox.askyesno(
            "Onarım gerekli",
            "Kütüphaneler bu bilgisayarın işlemcisine uymayan biçimde "
            "kurulmuş görünüyor (Intel/Apple Silicon karışması).\n\n"
            "Şimdi otomatik onarılsın mı? Birkaç dakika sürebilir "
            "(internet gerekir).", parent=kok)
        kok.destroy()
        hata_ozeti = ""
        if onar:
            r = subprocess.run(
                [sys.executable, "-m", "pip", "install", "--user",
                 "--force-reinstall", "--no-cache-dir",
                 "pymupdf", "numpy", "scipy"],
                capture_output=True, text=True, timeout=1800)
            logla("Mimari onarım pip çıktısı:\n"
                  + (r.stdout or "")[-600:] + (r.stderr or "")[-600:])
            importlib.invalidate_caches()
            # Doğrulama TEMİZ bir süreçte yapılır (bu oturumun hafızası kirli)
            dogrulama = subprocess.run(
                [sys.executable, "-c", "import pymupdf; print('ONARIM_OK')"],
                capture_output=True, text=True, timeout=120)
            if "ONARIM_OK" in (dogrulama.stdout or ""):
                kok = tk.Tk()
                kok.withdraw()
                messagebox.showinfo(
                    "Onarım tamam",
                    "Kütüphaneler onarıldı. Uygulama şimdi kendini yeniden "
                    "başlatacak.", parent=kok)
                kok.destroy()
                subprocess.Popen([sys.executable,
                                  str(Path(__file__).resolve())])
                return
            hata_ozeti = ((r.stderr or "") + "\n"
                          + (dogrulama.stderr or "")).strip()[-500:]
        logla("Onarım doğrulanamadı.\n" + hata_ozeti)
        kok = tk.Tk()
        kok.withdraw()
        messagebox.showerror(
            "Onarım tamamlanamadı",
            "Şu komutu Terminale yapıştırıp çalıştırın, bitince uygulamayı "
            "yeniden açın:\n\n"
            f"{sys.executable} -m pip install --user --force-reinstall "
            f"--no-cache-dir pymupdf numpy scipy"
            + (f"\n\nHata özeti:\n{hata_ozeti}" if hata_ozeti else ""),
            parent=kok)
        kok.destroy()
        return

    # Sürükle-bırak desteği (kurulamazsa/çalışmazsa normal pencereyle devam edilir)
    dnd = None
    try:
        from tkinterdnd2 import DND_FILES, TkinterDnD
        pencere = TkinterDnD.Tk()
        dnd = DND_FILES
    except Exception as h:
        logla(f"Sürükle-bırak devre dışı: {h}")
        try:
            tk._default_root = None   # yarım kalan pencere varsayılan olmasın
        except Exception:
            pass
        pencere = tk.Tk()
    pencere.title(f"Kitapçık Eşleştirici v{SURUM} — Limit Yayınları")
    pencere.geometry("780x720")
    pencere.minsize(720, 600)

    sinav = tk.StringVar(master=pencere, value="TYT")
    yollar = {ad: tk.StringVar(master=pencere) for ad in
              ("a", "b", "sablon", "anah_a", "anah_b", "cikti")}

    def sec(ad, tur, kaydet=False):
        if kaydet:
            yol = filedialog.asksaveasfilename(
                defaultextension=".xlsx", initialfile="kazanim_tablosu.xlsx",
                filetypes=[("Excel", "*.xlsx")])
        else:
            yol = filedialog.askopenfilename(filetypes=tur)
        if yol:
            yollar[ad].set(yol)
            if ad == "a" and not yollar["cikti"].get():
                yollar["cikti"].set(str(Path(yol).with_name("kazanim_tablosu.xlsx")))

    ana = ttk.Frame(pencere, padding=14)
    ana.pack(fill="both", expand=True)

    ust = ttk.Frame(ana)
    ust.pack(fill="x", pady=(0, 8))
    ttk.Label(ust, text="Sınav türü:").pack(side="left")
    ttk.Radiobutton(ust, text="TYT", variable=sinav,
                    value="TYT").pack(side="left", padx=6)
    ttk.Radiobutton(ust, text="AYT", variable=sinav,
                    value="AYT").pack(side="left")
    ttk.Radiobutton(ust, text="Lise 9-11", variable=sinav,
                    value="OTOMATİK").pack(side="left", padx=6)
    ttk.Radiobutton(ust, text="Ortaokul 5-8", variable=sinav,
                    value="ORTAOKUL").pack(side="left")

    girisler = {}

    def satir(etiket, ad, tur, kaydet=False):
        cerceve = ttk.Frame(ana)
        cerceve.pack(fill="x", pady=3)
        ttk.Label(cerceve, text=etiket, width=26).pack(side="left")
        giris = ttk.Entry(cerceve, textvariable=yollar[ad])
        giris.pack(side="left", fill="x", expand=True, padx=(0, 6))
        girisler[ad] = giris
        ttk.Button(cerceve, text="Seç...",
                   command=lambda: sec(ad, tur, kaydet)).pack(side="left")

    HEPSI = ("Tüm Dosyalar", "*")
    PDF = [("PDF", "*.pdf"), HEPSI]
    XLSX = [("Kazanım dosyası", "*.xlsx *.xlsm *.csv *.docx *.pdf"),
            ("Excel", "*.xlsx"), HEPSI]
    TXT = [("Metin", "*.txt"), HEPSI]
    satir("A Kitapçığı PDF (zorunlu):", "a", PDF)
    satir("B Kitapçığı PDF (zorunlu):", "b", PDF)
    satir("Kazanım şablonu Excel:", "sablon", XLSX)
    satir("Cevap anahtarı A (txt):", "anah_a", TXT)
    satir("Cevap anahtarı B (txt):", "anah_b", TXT)

    satir("Çıktı Excel dosyası:", "cikti", XLSX, kaydet=True)

    def ciktiyi_doldur(a_yolu):
        if not yollar["cikti"].get():
            yollar["cikti"].set(str(Path(a_yolu).with_name("kazanim_tablosu.xlsx")))

    if dnd:
        def yollari_ayikla(veri):
            try:
                return [str(p) for p in pencere.tk.splitlist(veri)]
            except Exception:
                return [veri.strip("{}").strip()] if veri else []

        def alana_birak(ad):
            def isleyici(olay):
                dosyalar = yollari_ayikla(olay.data)
                if dosyalar:
                    yollar[ad].set(dosyalar[0])
                    if ad == "a":
                        ciktiyi_doldur(dosyalar[0])
                return getattr(olay, "action", None)
            return isleyici

        def pencereye_birak(olay):
            for yol in sorted(yollari_ayikla(olay.data)):
                u = yol.lower()
                if u.endswith(".pdf"):
                    hedef = "a" if not yollar["a"].get() else "b"
                elif u.endswith((".xlsx", ".xlsm", ".csv", ".docx")):
                    hedef = "sablon"
                elif u.endswith(".txt"):
                    hedef = "anah_a" if not yollar["anah_a"].get() else "anah_b"
                else:
                    continue
                yollar[hedef].set(yol)
                if hedef == "a":
                    ciktiyi_doldur(yol)
            return getattr(olay, "action", None)

        try:
            pencere.drop_target_register(dnd)
            pencere.dnd_bind("<<Drop>>", pencereye_birak)
            for ad_, giris_ in girisler.items():
                giris_.drop_target_register(dnd)
                giris_.dnd_bind("<<Drop>>", alana_birak(ad_))
            ttk.Label(ana, text="İpucu: dosyaları ilgili kutunun üzerine ya da "
                                "pencerenin boş bir yerine sürükleyip bırakabilirsiniz "
                                "(A ve B PDF'lerini birlikte bırakın, ikisi de dolar).",
                      foreground="#666666").pack(anchor="w", pady=(4, 0))
        except Exception as h:
            logla(f"Sürükle-bırak kaydı yapılamadı: {h}")
    if not dnd:
        ttk.Label(ana, text="Not: Sürükle-bırak bu bilgisayarda kullanılamıyor; "
                            "dosyaları Seç... düğmeleriyle ekleyin (tam işlevlidir).",
                  foreground="#8a6d00").pack(anchor="w", pady=(4, 0))

    alt_cerceve = ttk.Frame(ana)
    alt_cerceve.pack(side="bottom", pady=(6, 2))
    buton = ttk.Button(alt_cerceve, text="EŞLEŞTİR",
                       command=lambda: calistir_tikla())
    buton.pack(side="left", padx=6, ipadx=30, ipady=6)
    kontrol_buton = ttk.Button(alt_cerceve, text="A–B KONTROL",
                               command=lambda: kontrol_tikla())
    kontrol_buton.pack(side="left", padx=6, ipadx=18, ipady=6)
    ttk.Label(ana, text="A–B KONTROL (baskı kontrolü): A'daki her soru metni, "
                        "şıkları ve şekilleriyle B'de var mı, numaralar branş "
                        "branş düzgün mü, metne bağlı gruplar aynı mı, açık kalan "
                        "cevap var mı. Rapor Excel + işaretli PDF olarak çıktı "
                        "klasörüne yazılır.",
              foreground="#666666", wraplength=740,
              justify="left").pack(side="bottom", anchor="w")

    log_kutusu = scrolledtext.ScrolledText(ana, height=16, state="disabled",
                                           font=("Menlo", 11))
    log_kutusu.pack(fill="both", expand=True, pady=(10, 8))
    kuyruk = queue.Queue()
    kayit_kuyrugu = queue.Queue()

    def log(msg):
        kuyruk.put(str(msg))

    def yol_sor(varsayilan):
        kuyruk.put(("KAYIT", varsayilan))
        return kayit_kuyrugu.get()

    def kuyrugu_isle():
        while not kuyruk.empty():
            oge = kuyruk.get()
            if isinstance(oge, tuple) and oge and oge[0] == "BITTI":
                bitti(*oge[1:])
                continue
            if isinstance(oge, tuple) and oge and oge[0] == "KAYIT":
                messagebox.showinfo(
                    "Kayıt izni gerekli",
                    "macOS, çıktı klasörüne doğrudan yazmaya izin vermedi.\n"
                    "Şimdi açılacak pencerede kayıt konumunu seçin — seçtiğiniz "
                    "konuma yazma izni verilmiş olur.", parent=pencere)
                yol = filedialog.asksaveasfilename(
                    parent=pencere, defaultextension=".xlsx",
                    initialdir=os.path.dirname(oge[1]) or str(Path.home()),
                    initialfile=os.path.basename(oge[1]),
                    filetypes=[("Excel", "*.xlsx")])
                kayit_kuyrugu.put(yol or "")
                continue
            log_kutusu.configure(state="normal")
            log_kutusu.insert("end", str(oge) + "\n")
            log_kutusu.see("end")
            log_kutusu.configure(state="disabled")
        pencere.after(150, kuyrugu_isle)

    def bitti(cikti_yolu, mesaj="Tablo hazır.", ekler=()):
        buton.configure(state="normal")
        kontrol_buton.configure(state="normal")
        ekler = [e for e in ekler if e and Path(e).exists()]
        soru = ("Excel dosyası ve bakılacak soruların A/B görüntüleri (PDF) "
                "şimdi açılsın mı?" if ekler else "Excel dosyası şimdi açılsın mı?")
        if cikti_yolu and messagebox.askyesno(
                "Tamamlandı", f"{mesaj}\n{soru}", parent=pencere):
            subprocess.run(["open", cikti_yolu])
            for e in ekler:
                subprocess.run(["open", e])

    def calistir_tikla():
        a, b = yollar["a"].get().strip(), yollar["b"].get().strip()
        if not a or not b:
            messagebox.showwarning("Eksik", "A ve B kitapçık PDF'lerini seçin.",
                                   parent=pencere)
            return
        cikti = yollar["cikti"].get().strip() or str(
            Path(a).with_name("kazanim_tablosu.xlsx"))
        buton.configure(state="disabled")
        kontrol_buton.configure(state="disabled")
        log_kutusu.configure(state="normal")
        log_kutusu.delete("1.0", "end")
        log_kutusu.configure(state="disabled")

        def is_parcasi():
            try:
                yol = calistir(a, b, sinav.get(),
                               yollar["sablon"].get().strip() or None,
                               yollar["anah_a"].get().strip() or None,
                               yollar["anah_b"].get().strip() or None,
                               cikti, log, yol_sor=yol_sor)
                kuyruk.put(("BITTI", yol))
            except Exception as h:
                log(f"\nHATA: {h}")
                kuyruk.put(("BITTI", None))

        threading.Thread(target=is_parcasi, daemon=True).start()

    def kontrol_tikla():
        a, b = yollar["a"].get().strip(), yollar["b"].get().strip()
        if not a or not b:
            messagebox.showwarning("Eksik", "A ve B kitapçık PDF'lerini seçin.",
                                   parent=pencere)
            return
        klasor = Path(yollar["cikti"].get().strip() or a).parent
        cikti = str(klasor / "kitapcik_kontrol_raporu.xlsx")
        buton.configure(state="disabled")
        kontrol_buton.configure(state="disabled")
        log_kutusu.configure(state="normal")
        log_kutusu.delete("1.0", "end")
        log_kutusu.configure(state="disabled")

        def is_parcasi():
            try:
                yol = kontrol_calistir(a, b, sinav.get(),
                                       yollar["anah_a"].get().strip() or None,
                                       yollar["anah_b"].get().strip() or None,
                                       cikti, log, yol_sor=yol_sor)
                kuyruk.put(("BITTI", yol, "Kontrol raporu hazır.",
                            [str(kart_yolu(yol))]))
            except Exception as h:
                logla("KONTROL HATASI:\n" + traceback.format_exc())
                log(f"\nHATA: {h}")
                kuyruk.put(("BITTI", None))

        threading.Thread(target=is_parcasi, daemon=True).start()

    pencere.lift()
    pencere.attributes("-topmost", True)
    pencere.after(900, lambda: pencere.attributes("-topmost", False))
    pencere.focus_force()

    kuyrugu_isle()
    pencere.mainloop()


if __name__ == "__main__":
    try:
        logla(f"--- Uygulama başlatılıyor --- Python {sys.version.split()[0]} ---")
        gui_baslat()
    except Exception:
        detay = traceback.format_exc()
        logla("BAŞLATMA HATASI:\n" + detay)
        print(detay)
        hata_penceresi("Uygulama başlatılırken bir hata oluştu. Ayrıntı şu dosyaya "
                       "kaydedildi: Library/Logs/KitapcikEslestirici.log")
